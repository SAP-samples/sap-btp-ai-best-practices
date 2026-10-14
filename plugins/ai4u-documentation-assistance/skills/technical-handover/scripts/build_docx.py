#!/usr/bin/env python3
"""
Build an SAP-styled technical handover .docx from a handover Markdown file.

Usage:
    python build_docx.py handover.md -o Handover.docx [--brand ../assets/brand.json] [--pdf]

Markdown conventions (see references/markdown-conventions.md for the full list):
  * YAML front matter -> cover page + header/footer + document properties (`header` overrides the page-header text,
                         `repository`/`commit` add optional cover rows)
  * [[TOC]]                          -> table of contents (Word field, pre-populated)
  * # Title {.part tagline="..."}    -> part divider (not numbered, not in TOC)
  * # Title {.unnumbered}            -> chapter without number; "# Appendix A. ..." is never numbered
  * # Title {.newpage}               -> force page break before the chapter
  * ![Caption](figures/x.png "width=0.7")  alone in a paragraph -> numbered figure
  * Table: Caption                   -> paragraph right before a pipe table -> numbered table caption
  * > [!NOTE] Optional title         -> callout box. Types: NOTE TIP IMPORTANT WARNING CAUTION GAP OPEN EVIDENCE
  * \\newpage  or  <!-- pagebreak --> -> page break
Dependencies: python-docx, markdown-it-py, pyyaml (optional)
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu
from markdown_it import MarkdownIt
from markdown_it.tree import SyntaxTreeNode

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_BRAND = SKILL_DIR / "assets" / "brand.json"


# ----------------------------------------------------------------------------- helpers
def parse_front_matter(text: str):
    if not text.startswith("---"):
        return {}, text
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not m:
        return {}, text
    raw, body = m.group(1), text[m.end():]
    try:
        import yaml  # type: ignore
        fm = yaml.safe_load(raw) or {}
    except Exception:
        fm = {}
        key = None
        for line in raw.splitlines():
            if re.match(r"^\s*-\s+", line) and key:
                fm.setdefault(key, [])
                if isinstance(fm[key], list):
                    fm[key].append(line.split("-", 1)[1].strip().strip("\"'"))
                continue
            mm = re.match(r"^(\w[\w_-]*)\s*:\s*(.*)$", line)
            if mm:
                key, val = mm.group(1), mm.group(2).strip().strip("\"'")
                fm[key] = val if val else []
    return fm, body


def slug_to_bookmark(slug: str) -> str:
    s = re.sub(r"[^0-9A-Za-z_]", "_", slug.lstrip("#"))
    s = re.sub(r"_+", "_", s).strip("_") or "x"
    return ("h_" + s)[:40]


def github_slug(text: str) -> str:
    s = text.strip().lower()
    s = re.sub(r"[^\w\- ]", "", s)
    return s.replace(" ", "-")


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.upper())


def set_cell_fill(cell, hexfill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def set_cell_borders(cell, **edges):
    """edges: top/left/bottom/right = (sz_eighths, color) or None for nil"""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        if edge not in edges:
            continue
        el = OxmlElement(f"w:{edge}")
        spec = edges[edge]
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec[1])
        borders.append(el)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_table_borders(table, spec: dict):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        s = spec.get(edge)
        if s is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(s[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), s[1])
        borders.append(el)
    tblPr.append(borders)


def set_table_fixed_widths(table, widths_emu):
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(sum(widths_emu) / 635)))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_emu):
            gc.set(qn("w:w"), str(int(widths_emu[i] / 635)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_emu):
                cell.width = Emu(int(widths_emu[i]))


def mark_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def mark_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def para_border(paragraph, edge="bottom", sz=6, color="0070F2", space=4):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pbdr.append(el)


def add_field(paragraph, instr: str, placeholder: str = "1", size=None, color=None):
    def r():
        run = OxmlElement("w:r")
        if size or color:
            rPr = OxmlElement("w:rPr")
            if color:
                c = OxmlElement("w:color"); c.set(qn("w:val"), color); rPr.append(c)
            if size:
                s = OxmlElement("w:sz"); s.set(qn("w:val"), str(int(size * 2))); rPr.append(s)
            run.append(rPr)
        return run
    b = r(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); b.append(fc)
    i = r(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; i.append(it)
    s = r(); fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "separate"); s.append(fc2)
    t = r(); tt = OxmlElement("w:t"); tt.text = placeholder; t.append(tt)
    e = r(); fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end"); e.append(fc3)
    for el in (b, i, s, t, e):
        paragraph._p.append(el)


def set_style_font(style, name, size=None, bold=None, italic=None, color=None):
    f = style.font
    f.name = name
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = rgb(color)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if rFonts.get(qn(f"w:{a}")) is not None:
            del rFonts.attrib[qn(f"w:{a}")]
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{a}"), name)
    # drop theme colour so the rgb wins
    c = rPr.find(qn("w:color"))
    if c is not None:
        for a in ("themeColor", "themeShade", "themeTint"):
            if c.get(qn(f"w:{a}")) is not None:
                del c.attrib[qn(f"w:{a}")]


# ----------------------------------------------------------------------------- builder
class HandoverBuilder:
    def __init__(self, md_path: Path, brand: dict):
        self.md_path = md_path
        self.base = md_path.parent
        self.warnings = []
        self.brand = brand
        self.C = brand["colors"]
        self.font = brand.get("font", "Arial")
        self.mono = brand.get("font_mono", "Consolas")
        text = md_path.read_text(encoding="utf-8")
        self.fm, self.body = parse_front_matter(text)
        self.body = self.resolve_crossrefs(self.body)
        self.doc = Document()
        self.h1 = 0
        self.h2 = 0
        self.h1_numbered = False
        self.fig = 0
        self.tab = 0
        self.toc_entries = []  # (level, text, bookmark)
        self.bm_id = 100
        self.pending_table_caption = None
        self.toc_placeholder = None
        self.num_next = 9100
        self.abs_bullet = None
        self.abs_decimal = None
        self.stats = {"chapters": 0, "figures": 0, "tables": 0, "callouts": 0, "code_blocks": 0}

    def resolve_crossrefs(self, body: str) -> str:
        """Labels: ![Caption](x.png "id=fig-arch") and 'Table: Caption {#tbl-gaps}'. References: @fig-arch, @tbl-gaps."""
        labels, fig, tab = {}, 0, 0
        in_code = False
        lines = body.split("\n")
        for i, ln in enumerate(lines):
            if ln.strip().startswith("```"):
                in_code = not in_code
                continue
            if in_code:
                continue
            if re.match(r"^\s*!\[[^\]]*\]\([^)]*\)\s*$", ln):
                fig += 1
                m = re.search(r'id=([\w-]+)', ln)
                if m:
                    labels[m.group(1)] = f"Figure {fig}"
            elif ln.startswith("Table:"):
                tab += 1
                m = re.search(r"\{#([\w-]+)\}\s*$", ln)
                if m:
                    labels[m.group(1)] = f"Table {tab}"
                    lines[i] = ln[: m.start()].rstrip()
        out, in_code = [], False
        for ln in lines:
            if ln.strip().startswith("```"):
                in_code = not in_code
            if not in_code:
                def rep(m):
                    key = m.group(1)
                    if key in labels:
                        return labels[key]
                    self.warnings.append(f"unresolved cross-reference @{key}")
                    return m.group(0)
                ln = re.sub(r"(?<![\w@`])@((?:fig|tbl)-[\w-]+)", rep, ln)
            out.append(ln)
        return "\n".join(out)

    # ------------------------------------------------------------------ setup
    def setup_document(self):
        sec = self.doc.sections[0]
        if self.brand.get("page", "A4").upper() == "A4":
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        else:
            sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        sec.orientation = WD_ORIENT.PORTRAIT
        m = Cm(self.brand.get("margins_cm", 2.2))
        sec.left_margin = sec.right_margin = m
        sec.top_margin = Cm(2.4)
        sec.bottom_margin = Cm(2.2)
        sec.header_distance = Cm(1.0)
        sec.footer_distance = Cm(1.0)
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin

        st = self.doc.styles
        normal = st["Normal"]
        set_style_font(normal, self.font, self.brand.get("body_size_pt", 10.5), color=self.C["text"])
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        for name, size, color, before, after in (
            ("Heading 1", 18, self.C["navy"], 20, 8),
            ("Heading 2", 13.5, self.C["blue_dark"], 14, 5),
            ("Heading 3", 11.5, self.C["navy"], 10, 4),
        ):
            s = st[name]
            set_style_font(s, self.font, size, bold=True, italic=False, color=color)
            pf = s.paragraph_format
            pf.space_before, pf.space_after = Pt(before), Pt(after)
            pf.keep_with_next = True
            pf.line_spacing = 1.05

        cap = st["Caption"]
        set_style_font(cap, self.font, 8.5, bold=False, italic=True, color=self.C["muted"])
        cap.paragraph_format.space_after = Pt(12)
        cap.paragraph_format.space_before = Pt(4)

        from docx.enum.style import WD_STYLE_TYPE
        for nm, size, bold, indent in (("TOC 1", 10.5, True, 0), ("TOC 2", 10, False, 0.6)):
            try:
                s = st[nm]
            except KeyError:
                s = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
                s.base_style = normal
            set_style_font(s, self.font, size, bold=bold, color=self.C["text"])
            s.paragraph_format.left_indent = Cm(indent)
            s.paragraph_format.space_after = Pt(3 if bold else 1)
            s.paragraph_format.space_before = Pt(4 if bold else 0)
        try:
            th = st["TOC Heading"]
        except KeyError:
            th = st.add_style("TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
            th.base_style = normal
        set_style_font(th, self.font, 18, bold=True, color=self.C["navy"])
        th.paragraph_format.space_after = Pt(10)

        self._setup_numbering()

        cp = self.doc.core_properties
        cp.title = str(self.fm.get("title", "Technical Handover"))
        cp.subject = "Technical handover"
        authors = self.fm.get("authors") or self.fm.get("prepared_by") or ""
        cp.author = ", ".join(authors) if isinstance(authors, list) else str(authors)
        cp.keywords = "technical handover; architecture; API; AI; production readiness"
        import datetime as _dt
        now = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None, microsecond=0)
        cp.created = now
        cp.modified = now
        cp.last_modified_by = cp.author
        cp.revision = 1
        cp.comments = str(self.fm.get("subtitle", ""))
        cp.category = str(self.fm.get("classification", ""))

        if self.brand.get("update_fields_on_open", True):
            settings = self.doc.settings.element
            uf = OxmlElement("w:updateFields")
            uf.set(qn("w:val"), "true")
            settings.append(uf)

    def _setup_numbering(self):
        numbering = self.doc.part.numbering_part.element

        def abstract(aid, kind):
            a = OxmlElement("w:abstractNum")
            a.set(qn("w:abstractNumId"), str(aid))
            mlt = OxmlElement("w:multiLevelType"); mlt.set(qn("w:val"), "hybridMultilevel"); a.append(mlt)
            for lvl in range(4):
                l = OxmlElement("w:lvl"); l.set(qn("w:ilvl"), str(lvl))
                start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); l.append(start)
                fmt = OxmlElement("w:numFmt")
                txt = OxmlElement("w:lvlText")
                if kind == "bullet":
                    fmt.set(qn("w:val"), "bullet")
                    txt.set(qn("w:val"), ["•", "–", "▪", "•"][lvl])
                else:
                    fmt.set(qn("w:val"), ["decimal", "lowerLetter", "lowerRoman", "decimal"][lvl])
                    txt.set(qn("w:val"), [f"%1.", f"%2)", f"%3.", f"%4."][lvl])
                l.append(fmt); l.append(txt)
                jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); l.append(jc)
                pPr = OxmlElement("w:pPr")
                ind = OxmlElement("w:ind")
                ind.set(qn("w:left"), str(400 + 360 * lvl)); ind.set(qn("w:hanging"), "300")
                pPr.append(ind); l.append(pPr)
                rPr = OxmlElement("w:rPr")
                rf = OxmlElement("w:rFonts")
                for k in ("ascii", "hAnsi", "cs"):
                    rf.set(qn(f"w:{k}"), self.font)
                rPr.append(rf)
                col = OxmlElement("w:color"); col.set(qn("w:val"), self.C["blue"] if kind == "bullet" else self.C["navy"])
                rPr.append(col)
                if kind != "bullet":
                    b = OxmlElement("w:b"); rPr.append(b)
                l.append(rPr)
                a.append(l)
            return a

        self.abs_bullet, self.abs_decimal = 9001, 9002
        first_num = numbering.find(qn("w:num"))
        for aid, kind in ((self.abs_bullet, "bullet"), (self.abs_decimal, "decimal")):
            el = abstract(aid, kind)
            if first_num is not None:
                first_num.addprevious(el)
            else:
                numbering.append(el)

    def new_num(self, ordered: bool) -> int:
        numbering = self.doc.part.numbering_part.element
        nid = self.num_next
        self.num_next += 1
        num = OxmlElement("w:num"); num.set(qn("w:numId"), str(nid))
        an = OxmlElement("w:abstractNumId"); an.set(qn("w:val"), str(self.abs_decimal if ordered else self.abs_bullet))
        num.append(an)
        if ordered:
            for lvl in range(4):
                o = OxmlElement("w:lvlOverride"); o.set(qn("w:ilvl"), str(lvl))
                so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1")
                o.append(so); num.append(o)
        numbering.append(num)
        return nid

    # ------------------------------------------------------------------ cover, header, footer
    def build_cover(self):
        fm, C = self.fm, self.C
        doc = self.doc
        logo = self.brand.get("logo")
        if logo:
            lp = (SKILL_DIR / logo) if not Path(logo).is_absolute() else Path(logo)
            if lp.exists():
                doc.add_paragraph().add_run().add_picture(str(lp), height=Cm(1.2))
        sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(60 if not logo else 30)

        band = doc.add_table(rows=1, cols=1)
        band.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = band.rows[0].cells[0]
        set_cell_fill(cell, C["navy"])
        set_cell_margins(cell, top=500, bottom=520, left=500, right=500)
        set_table_borders(band, {})
        p = cell.paragraphs[0]
        r = p.add_run(self.brand.get("document_label", "TECHNICAL HANDOVER"))
        r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = rgb("89D1FF")
        self._letter_spacing(r, 40)
        p.paragraph_format.space_after = Pt(14)
        p = cell.add_paragraph()
        r = p.add_run(str(fm.get("title", "Technical Handover")))
        r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = rgb("FFFFFF")
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(10)
        if fm.get("subtitle"):
            p = cell.add_paragraph()
            r = p.add_run(str(fm["subtitle"]))
            r.font.size = Pt(13); r.font.color.rgb = rgb("D1EFFF")
            p.paragraph_format.space_after = Pt(0)
        set_table_fixed_widths(band, [self.text_width])

        bar = doc.add_paragraph()
        para_border(bar, "top", sz=36, color=C["blue"], space=0)
        bar.paragraph_format.space_after = Pt(28)

        rows = []
        def add(label, key):
            v = fm.get(key)
            if v:
                if isinstance(v, list):
                    v = ", ".join(str(x) for x in v)
                rows.append((label, str(v)))
        add("Customer", "customer"); add("Project", "project"); add("Document version", "version")
        add("Date", "date"); add("Status", "status"); add("Classification", "classification")
        add("Prepared by", "authors"); add("Prepared for", "audience")
        add("Source repository", "repository"); add("Source commit", "commit")  # optional, off by default
        if rows:
            t = doc.add_table(rows=len(rows), cols=2)
            set_table_borders(t, {"insideH": (4, C["border"]), "bottom": (4, C["border"]), "top": (4, C["border"])})
            for i, (k, v) in enumerate(rows):
                c0, c1 = t.rows[i].cells
                for c in (c0, c1):
                    set_cell_margins(c, top=70, bottom=70, left=60, right=60)
                r = c0.paragraphs[0].add_run(k.upper())
                r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = rgb(C["muted"])
                self._letter_spacing(r, 10)
                r = c1.paragraphs[0].add_run(v)
                r.font.size = Pt(10.5); r.font.color.rgb = rgb(C["text"])
                for c in (c0, c1):
                    c.paragraphs[0].paragraph_format.space_after = Pt(0)
            set_table_fixed_widths(t, [int(self.text_width * 0.3), int(self.text_width * 0.7)])
        if fm.get("cover_note"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(36)
            r = p.add_run(str(fm["cover_note"]))
            r.font.size = Pt(8.5); r.font.italic = True; r.font.color.rgb = rgb(C["muted"])
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _letter_spacing(self, run, twentieths):
        rPr = run._r.get_or_add_rPr()
        sp = OxmlElement("w:spacing"); sp.set(qn("w:val"), str(twentieths)); rPr.append(sp)

    def build_header_footer(self):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        self._fill_header_footer(sec)

    def _fill_header_footer(self, sec):
        fm, C = self.fm, self.C
        tw = sec.page_width - sec.left_margin - sec.right_margin
        for stname in ("Header", "Footer"):
            try:
                self.doc.styles[stname].paragraph_format.tab_stops.clear_all()
            except Exception:
                pass
        hp = sec.header.paragraphs[0]
        hp.paragraph_format.tab_stops.add_tab_stop(tw, alignment=WD_TAB_ALIGNMENT.RIGHT)
        left = f"{fm.get('header') or fm.get('project') or fm.get('title', '')}  |  Technical Handover"
        r = hp.add_run(left); r.font.size = Pt(8); r.font.color.rgb = rgb(C["muted"])
        right = str(fm.get("classification", ""))
        if right:
            r = hp.add_run("\t" + right); r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = rgb(C["blue"])
        para_border(hp, "bottom", sz=4, color=C["border"], space=4)

        fp = sec.footer.paragraphs[0]
        fp.paragraph_format.tab_stops.add_tab_stop(tw, alignment=WD_TAB_ALIGNMENT.RIGHT)
        bits = [b for b in (f"Version {fm['version']}" if fm.get("version") else "", str(fm.get("date", ""))) if b]
        r = fp.add_run("  ·  ".join(bits)); r.font.size = Pt(8); r.font.color.rgb = rgb(C["muted"])
        r = fp.add_run("\tPage "); r.font.size = Pt(8); r.font.color.rgb = rgb(C["muted"])
        add_field(fp, "PAGE", "1", size=8, color=C["muted"])
        r = fp.add_run(" of "); r.font.size = Pt(8); r.font.color.rgb = rgb(C["muted"])
        add_field(fp, "NUMPAGES", "1", size=8, color=C["muted"])

    # ------------------------------------------------------------------ inline rendering
    def _run_el(self, text, fmt, br=False):
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        if fmt.get("code"):
            rf = OxmlElement("w:rFonts")
            for k in ("ascii", "hAnsi", "cs"):
                rf.set(qn(f"w:{k}"), self.mono)
            rPr.append(rf)
        if fmt.get("bold"):
            rPr.append(OxmlElement("w:b"))
        if fmt.get("italic"):
            rPr.append(OxmlElement("w:i"))
        if fmt.get("strike"):
            rPr.append(OxmlElement("w:strike"))
        color = fmt.get("color")
        if fmt.get("link"):
            color = self.C["blue"]
        elif fmt.get("code"):
            color = color or self.C["blue_dark"]
        if color:
            c = OxmlElement("w:color"); c.set(qn("w:val"), color); rPr.append(c)
        size = fmt.get("size")
        if fmt.get("code") and not size:
            size = (fmt.get("base_size") or self.brand.get("body_size_pt", 10.5)) - 1
        elif fmt.get("base_size") and not size:
            size = fmt["base_size"]
        if size:
            s = OxmlElement("w:sz"); s.set(qn("w:val"), str(int(round(size * 2)))); rPr.append(s)
        if fmt.get("link"):
            u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
        if fmt.get("code"):
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EEF1F4")
            rPr.append(shd)
        if len(rPr):
            r.append(rPr)
        if br:
            r.append(OxmlElement("w:br"))
        else:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            r.append(t)
        return r

    def render_inline(self, parent_el, node, fmt, paragraph):
        """parent_el: w:p or w:hyperlink element"""
        for ch in node.children or []:
            t = ch.type
            if t == "text":
                parent_el.append(self._run_el(ch.content, fmt))
            elif t == "softbreak":
                parent_el.append(self._run_el(" ", fmt))
            elif t == "hardbreak":
                parent_el.append(self._run_el("", fmt, br=True))
            elif t == "html_inline":
                if re.match(r"<br\s*/?>", ch.content, re.I):
                    parent_el.append(self._run_el("", fmt, br=True))
            elif t == "strong":
                self.render_inline(parent_el, ch, {**fmt, "bold": True}, paragraph)
            elif t == "em":
                self.render_inline(parent_el, ch, {**fmt, "italic": True}, paragraph)
            elif t == "s":
                self.render_inline(parent_el, ch, {**fmt, "strike": True}, paragraph)
            elif t == "code_inline":
                parent_el.append(self._run_el(ch.content, {**fmt, "code": True}))
            elif t == "link":
                href = ch.attrs.get("href", "")
                hl = OxmlElement("w:hyperlink")
                if href.startswith("#"):
                    hl.set(qn("w:anchor"), slug_to_bookmark(href))
                else:
                    rid = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
                    hl.set(qn("r:id"), rid)
                self.render_inline(hl, ch, {**fmt, "link": True}, paragraph)
                parent_el.append(hl)
            elif t == "image":
                self.warnings.append(f"inline image ignored (put images alone in a paragraph): {ch.attrs.get('src')}")
            else:
                if ch.children:
                    self.render_inline(parent_el, ch, fmt, paragraph)
                elif getattr(ch, "content", None):
                    parent_el.append(self._run_el(ch.content, fmt))

    def inline_text(self, node) -> str:
        out = []
        def walk(n):
            for ch in n.children or []:
                if ch.type in ("text", "code_inline"):
                    out.append(ch.content)
                elif ch.type in ("softbreak", "hardbreak"):
                    out.append(" ")
                elif ch.children:
                    walk(ch)
        walk(node)
        return "".join(out)

    # ------------------------------------------------------------------ containers
    def new_par(self, container, style=None):
        if getattr(container, "_fresh", False):
            container._fresh = False
            p = container.paragraphs[0]
            if style:
                p.style = self.doc.styles[style]
            return p
        return container.add_paragraph(style=style) if style else container.add_paragraph()

    def cell_container(self, cell):
        cell._fresh = True
        return cell

    # ------------------------------------------------------------------ blocks
    def render_blocks(self, nodes, container, list_ctx=None, base_size=None):
        for node in nodes:
            self.render_block(node, container, list_ctx, base_size)

    def render_block(self, node, container, list_ctx=None, base_size=None):
        t = node.type
        if t == "heading":
            self.render_heading(node, container)
        elif t == "paragraph":
            self.render_paragraph(node, container, list_ctx, base_size)
        elif t in ("bullet_list", "ordered_list"):
            self.render_list(node, container, list_ctx, base_size)
        elif t == "table":
            self.render_table(node, container)
        elif t in ("fence", "code_block"):
            self.render_code(node, container)
        elif t == "blockquote":
            self.render_callout(node, container)
        elif t == "hr":
            p = self.new_par(container)
            para_border(p, "bottom", sz=4, color=self.C["border"])
        elif t == "html_block":
            if re.search(r"<!--\s*landscape\s*-->", node.content, re.I):
                self.switch_orientation(landscape=True)
            elif re.search(r"<!--\s*portrait\s*-->", node.content, re.I):
                self.switch_orientation(landscape=False)
            elif re.search(r"pagebreak|newpage", node.content, re.I):
                self.new_par(container).add_run().add_break(WD_BREAK.PAGE)
        else:
            if node.children:
                self.render_blocks(node.children, container, list_ctx, base_size)

    def switch_orientation(self, landscape: bool):
        prev = self.doc.sections[-1]
        want = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        if prev.orientation == want:
            return
        body = self.doc.element.body
        if getattr(self, "_body_len_at_switch", None) == len(body) and len(self.doc.sections) > 1:
            # nothing written since the last switch: flip the current section instead of adding an empty one
            w, h = prev.page_width, prev.page_height
            prev.orientation = want
            prev.page_width, prev.page_height = (max(w, h), min(w, h)) if landscape else (min(w, h), max(w, h))
            self.text_width = prev.page_width - prev.left_margin - prev.right_margin
            for part in (prev.header, prev.footer):
                for para in part.paragraphs:
                    para._p.getparent().remove(para._p)
            prev.header.add_paragraph(); prev.footer.add_paragraph()
            self._fill_header_footer(prev)
            self._body_len_at_switch = len(body)
            return
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        w, h = prev.page_width, prev.page_height
        sec.orientation = want
        sec.page_width, sec.page_height = (max(w, h), min(w, h)) if landscape else (min(w, h), max(w, h))
        for a in ("left_margin", "right_margin", "top_margin", "bottom_margin", "header_distance", "footer_distance"):
            setattr(sec, a, getattr(prev, a))
        sec.different_first_page_header_footer = False
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin
        # own header/footer so the right-aligned tab matches the new width
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        self._fill_header_footer(sec)
        self._body_len_at_switch = len(self.doc.element.body)

    def render_heading(self, node, container):
        level = int(node.tag[1])
        raw = self.inline_text(node.children[0]).strip()
        attrs = {}
        m = re.search(r"\{([^}]*)\}\s*$", raw)
        if m:
            spec = m.group(1)
            raw = raw[: m.start()].rstrip()
            for cls in re.findall(r"\.([\w-]+)", spec):
                attrs[cls] = True
            for k, v in re.findall(r'(\w+)="([^"]*)"', spec):
                attrs[k] = v
        slug_bm = slug_to_bookmark(github_slug(raw))

        if level == 1 and attrs.get("part"):
            self.render_part(raw, attrs.get("tagline"))
            return

        number = ""
        if level == 1:
            numbered = not attrs.get("unnumbered") and not re.match(r"^(appendix|annex)\b", raw, re.I)
            self.h2 = 0
            self.h1_numbered = numbered
            if numbered:
                self.h1 += 1
                number = f"{self.h1}."
            self.stats["chapters"] += 1
        elif level == 2:
            if self.h1_numbered and not attrs.get("unnumbered"):
                self.h2 += 1
                number = f"{self.h1}.{self.h2}"
        style = f"Heading {min(level, 3)}"
        p = self.new_par(container, style)
        if attrs.get("newpage") and level <= 2:
            p.paragraph_format.page_break_before = True
        bid = self.bm_id; self.bm_id += 1
        bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), str(bid)); bs.set(qn("w:name"), slug_bm)
        p._p.append(bs)
        if number:
            r = p.add_run(number + "  ")
            r.font.color.rgb = rgb(self.C["blue"])
        p.add_run(raw)
        be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), str(bid))
        p._p.append(be)
        if level <= 2:
            self.toc_entries.append((level, (number + "  " if number else "") + raw, slug_bm))

    def render_part(self, title, tagline):
        C = self.C
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        m = re.match(r"^(part\s+[ivxlcdm\d]+)\s*[:.\-–—]\s*(.*)$", title, re.I)
        label, main = (m.group(1).upper(), m.group(2)) if m else ("PART", title)
        r = p.add_run(label); r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = rgb(C["blue"])
        self._letter_spacing(r, 40)
        p = self.doc.add_paragraph()
        r = p.add_run(main); r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = rgb(C["navy"])
        p.paragraph_format.space_after = Pt(6)
        para_border(p, "bottom", sz=24, color=C["blue"], space=8)
        if tagline:
            p = self.doc.add_paragraph()
            r = p.add_run(tagline); r.font.size = Pt(11.5); r.font.italic = True; r.font.color.rgb = rgb(C["muted"])
            p.paragraph_format.space_after = Pt(18)

    def render_paragraph(self, node, container, list_ctx, base_size, numpr=None, indent_level=None):
        inline = node.children[0] if node.children else None
        if inline is None:
            return
        text = self.inline_text(inline).strip()
        kids = [c for c in (inline.children or []) if not (c.type == "text" and not c.content.strip())]
        # page break / TOC markers
        if text in ("\\newpage", "\\pagebreak"):
            self.new_par(container).add_run().add_break(WD_BREAK.PAGE)
            return
        if text == "[[TOC]]":
            self.render_toc_placeholder(container)
            return
        # figure
        if len(kids) == 1 and kids[0].type == "image" and list_ctx is None:
            self.render_figure(kids[0], container)
            return
        # table caption
        if text.startswith("Table:") and list_ctx is None:
            self.pending_table_caption = text[len("Table:"):].strip()
            return
        p = self.new_par(container)
        fmt = {}
        if base_size:
            fmt["base_size"] = base_size
        if numpr is not None:
            self._apply_numpr(p, *numpr)
            p.paragraph_format.space_after = Pt(3)
        elif indent_level is not None:
            p.paragraph_format.left_indent = Emu(int((400 + 360 * indent_level) * 635))
            p.paragraph_format.space_after = Pt(3)
        prefix = getattr(self, "_callout_prefix", None)
        if prefix:
            label, color = prefix
            p._p.append(self._run_el(label + ": ", {"bold": True, "color": color, **fmt}))
            self._callout_prefix = None
        self.render_inline(p._p, inline, fmt, p)

    def _apply_numpr(self, p, num_id, level):
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(level))
        ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(num_id))
        numPr.append(il); numPr.append(ni)
        pPr.append(numPr)

    def render_list(self, node, container, list_ctx, base_size):
        ordered = node.type == "ordered_list"
        level = 0 if list_ctx is None else list_ctx["level"] + 1
        if list_ctx is not None and list_ctx["ordered"] == ordered:
            num_id = list_ctx["num"]
        else:
            num_id = self.new_num(ordered)
        ctx = {"level": level, "ordered": ordered, "num": num_id}
        for item in node.children:
            first = True
            for ch in item.children or []:
                if ch.type == "paragraph":
                    if first:
                        self.render_paragraph(ch, container, ctx, base_size, numpr=(num_id, level))
                        first = False
                    else:
                        self.render_paragraph(ch, container, ctx, base_size, indent_level=level)
                elif ch.type in ("bullet_list", "ordered_list"):
                    self.render_list(ch, container, ctx, base_size)
                else:
                    self.render_block(ch, container, ctx, base_size)

    def render_figure(self, img, container):
        src = img.attrs.get("src", "")
        title = img.attrs.get("title") or ""
        alt = self.inline_text(img) or img.attrs.get("alt", "") or ""
        path = (self.base / src) if not Path(src).is_absolute() else Path(src)
        frac = 1.0
        mm = re.search(r"width\s*=\s*([\d.]+)", title)
        if mm:
            frac = min(1.0, float(mm.group(1)))
        if not path.exists():
            self.warnings.append(f"missing figure: {src}")
            p = self.new_par(container)
            r = p.add_run(f"[Missing figure: {src}]"); r.font.color.rgb = rgb(self.C["red"])
            return
        p = self.new_par(container)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        width = int(self.text_width * frac)
        try:  # keep tall figures on one page
            from PIL import Image
            with Image.open(path) as im:
                iw, ih = im.size
            sec = self.doc.sections[-1]
            max_h = sec.page_height - sec.top_margin - sec.bottom_margin - Cm(3.2)
            if width * ih / iw > max_h:
                width = int(max_h * iw / ih)
        except Exception:
            pass
        run.add_picture(str(path), width=Emu(width))
        try:
            docpr = run._r.findall(".//" + qn("wp:docPr"))[0]
            docpr.set("descr", alt)
        except Exception:
            pass
        self.fig += 1
        self.stats["figures"] += 1
        cap = self.new_par(container, "Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {self.fig}. "); r.font.bold = True; r.font.italic = False; r.font.color.rgb = rgb(self.C["navy"])
        cap.add_run(alt)

    def render_toc_placeholder(self, container):
        p = self.new_par(container, "TOC Heading")
        p.add_run("Contents")
        self.toc_placeholder = self.new_par(container)

    def finalize_toc(self):
        if self.toc_placeholder is None or not self.toc_entries:
            return
        anchor = self.toc_placeholder._p
        paras = []
        for lvl, text, bm in self.toc_entries:
            p = OxmlElement("w:p")
            pPr = OxmlElement("w:pPr")
            ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), self.doc.styles[f"TOC {lvl}"].style_id)
            pPr.append(ps)
            p.append(pPr)
            hl = OxmlElement("w:hyperlink"); hl.set(qn("w:anchor"), bm); hl.set(qn("w:history"), "1")
            hl.append(self._run_el(text, {}))
            p.append(hl)
            paras.append(p)

        def fld(kind):
            r = OxmlElement("w:r"); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r.append(f); return r
        first, last = paras[0], paras[-1]
        instr_r = OxmlElement("w:r"); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
        it.text = ' TOC \\o "1-2" \\h \\z \\u '; instr_r.append(it)
        pPr = first.find(qn("w:pPr"))
        # order inside the first entry: pPr, begin, instr, separate, hyperlink...
        pPr.addnext(fld("begin"))
        begin = pPr.getnext()
        begin.addnext(instr_r)
        instr_r.addnext(fld("separate"))
        last.append(fld("end"))
        for p in paras:
            anchor.addprevious(p)
        # page break after the TOC
        br_p = OxmlElement("w:p"); r = OxmlElement("w:r"); b = OxmlElement("w:br"); b.set(qn("w:type"), "page"); r.append(b); br_p.append(r)
        anchor.addprevious(br_p)
        anchor.getparent().remove(anchor)

    def _table_weights(self, matrix):
        ncols = max(len(r) for r in matrix)
        weights = []
        for j in range(ncols):
            lens = [len(r[j]) for r in matrix if j < len(r)]
            mx = max(lens) if lens else 5
            avg = sum(lens) / len(lens) if lens else 5
            longest_word = max((len(w) for r in matrix if j < len(r) for w in r[j].split()), default=4)
            w = 0.55 * min(mx, 70) + 0.45 * min(avg, 70)
            w = max(w, min(longest_word, 38) * 0.9, 5)
            weights.append(w)
        return weights

    def render_table(self, node, container):
        rows = []  # list of (cells[node], is_header)
        for sect in node.children:
            for tr in sect.children:
                rows.append(([c for c in tr.children], sect.type == "thead"))
        if not rows:
            return
        ncols = max(len(r[0]) for r in rows)
        if self.pending_table_caption:
            self.tab += 1
            cap = self.new_par(container, "Caption")
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_after = Pt(4)
            r = cap.add_run(f"Table {self.tab}. "); r.font.bold = True; r.font.italic = False; r.font.color.rgb = rgb(self.C["navy"])
            cap.add_run(self.pending_table_caption)
            self.pending_table_caption = None
        self.stats["tables"] += 1
        table = container.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, {"top": (6, self.C["navy"]), "bottom": (6, self.C["navy"]), "insideH": (4, self.C["border"])})
        matrix = [[self.inline_text(c.children[0]) if c.children else "" for c in r[0]] for r in rows]
        weights = self._table_weights(matrix)
        tw = self.text_width if container is self.doc else int(self.text_width * 0.94)
        longest = [max((len(w) for r in matrix if j < len(r) for w in r[j].split()), default=4) for j in range(ncols)]
        mins = [max(Cm(1.6), Cm(min(longest[j], 28) * 0.2 + 0.45)) for j in range(ncols)]
        if sum(mins) > tw:
            mins = [m * tw / sum(mins) * 0.95 for m in mins]
        widths = [tw * w / sum(weights) for w in weights]
        # enforce minimum widths
        for _ in range(4):
            short = [i for i, w in enumerate(widths) if w < mins[i] - 1]
            if not short:
                break
            deficit = sum(mins[i] - widths[i] for i in short)
            for i in short:
                widths[i] = mins[i]
            others = [i for i in range(len(widths)) if i not in short]
            tot = sum(widths[i] for i in others)
            for i in others:
                widths[i] -= deficit * widths[i] / tot
        band = False
        for i, (cells, is_header) in enumerate(rows):
            row = table.rows[i]
            mark_row_cant_split(row)
            if is_header:
                mark_row_header(row)
            else:
                band = not band
            for j in range(ncols):
                cell = row.cells[j]
                set_cell_margins(cell, top=60, bottom=60, left=90, right=90)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                if j >= len(cells):
                    continue
                cn = cells[j]
                style = cn.attrs.get("style", "") if cn.attrs else ""
                if "right" in style:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif "center" in style:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if is_header:
                    set_cell_fill(cell, self.C["table_header"])
                    fmt = {"bold": True, "color": "FFFFFF", "base_size": 8.5}
                else:
                    if not band:
                        set_cell_fill(cell, self.C["table_band"])
                    fmt = {"base_size": 9}
                if cn.children:
                    self.render_inline(p._p, cn.children[0], fmt, p)
        set_table_fixed_widths(table, widths)
        after = self.new_par(container)
        after.paragraph_format.space_after = Pt(4)
        after.paragraph_format.line_spacing = 0.8

    def render_code(self, node, container):
        self.stats["code_blocks"] += 1
        table = container.add_table(rows=1, cols=1)
        cell = table.rows[0].cells[0]
        set_table_borders(table, {})
        set_cell_fill(cell, self.C["code_bg"])
        set_cell_borders(cell, left=(18, self.C["blue"]))
        set_cell_margins(cell, top=100, bottom=100, left=180, right=140)
        lines = node.content.rstrip("\n").split("\n")
        info = (node.info or "").strip()
        first = True
        if info:
            p = cell.paragraphs[0]; first = False
            r = p.add_run(info.upper()); r.font.size = Pt(7); r.font.bold = True; r.font.color.rgb = rgb(self.C["muted"])
            p.paragraph_format.space_after = Pt(2)
        for ln in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(ln if ln else " ")
            r.font.name = self.mono
            r._r.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:hAnsi"), self.mono)
            r.font.size = Pt(8.5)
            r.font.color.rgb = rgb(self.C["navy"])
        tw = self.text_width if container is self.doc else int(self.text_width * 0.94)
        set_table_fixed_widths(table, [tw])
        after = self.new_par(container)
        after.paragraph_format.space_after = Pt(2)
        after.paragraph_format.line_spacing = 0.8

    def render_callout(self, node, container):
        kind, title = "NOTE", None
        children = list(node.children or [])
        if children and children[0].type == "paragraph" and children[0].children:
            inline = children[0].children[0]
            kids = inline.children or []
            if kids and kids[0].type == "text":
                m = re.match(r"^\s*\[!(\w+)\]\s*(.*)$", kids[0].content)
                if m and m.group(1).upper() in self.brand["callouts"]:
                    kind = m.group(1).upper()
                    rest = m.group(2).strip()
                    # title = rest of first line if a line break follows
                    if len(kids) > 1 and kids[1].type in ("softbreak", "hardbreak"):
                        title = rest or None
                        inline.children = kids[2:]
                    elif len(kids) == 1:
                        if rest:
                            kids[0].token.content = rest
                        else:
                            inline.children = []
                    else:
                        kids[0].token.content = rest
                    if not inline.children:
                        children = children[1:]
        spec = self.brand["callouts"][kind]
        self.stats["callouts"] += 1
        table = container.add_table(rows=1, cols=1)
        set_table_borders(table, {})
        cell = table.rows[0].cells[0]
        set_cell_fill(cell, spec["fill"])
        set_cell_borders(cell, left=(30, spec["color"]))
        set_cell_margins(cell, top=110, bottom=110, left=200, right=160)
        cc = self.cell_container(cell)
        if title:
            p = self.new_par(cc)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p._p.append(self._run_el(f"{spec['label']}: {title}", {"bold": True, "color": spec["color"], "base_size": 9.5}))
        else:
            self._callout_prefix = (spec["label"], spec["color"])
        self.render_blocks(children, cc, None, base_size=9.5)
        self._callout_prefix = None
        # tighten last paragraph
        if cell.paragraphs:
            cell.paragraphs[-1].paragraph_format.space_after = Pt(0)
        tw = self.text_width if container is self.doc else int(self.text_width * 0.94)
        set_table_fixed_widths(table, [tw])
        after = self.new_par(container)
        after.paragraph_format.space_after = Pt(2)
        after.paragraph_format.line_spacing = 0.8

    def _trim_trailing_empty_paragraphs(self):
        """Spacer paragraphs at the very end can push an empty last page (seen after landscape sections)."""
        body = self.doc.element.body
        while len(body) >= 2:
            el = body[-2]
            if el.tag != qn("w:p"):
                break
            if el.findall(".//" + qn("w:t")) or el.findall(".//" + qn("w:drawing")) or el.findall(".//" + qn("w:sectPr")) \
                    or el.findall(".//" + qn("w:br")) or el.findall(".//" + qn("w:fldChar")):
                break
            body.remove(el)

    # ------------------------------------------------------------------ main
    def build(self, out: Path):
        self.setup_document()
        self.build_cover()
        self.build_header_footer()
        md = MarkdownIt("commonmark", {"html": True}).enable("table").enable("strikethrough")
        tokens = md.parse(self.body)
        tree = SyntaxTreeNode(tokens)
        if "[[TOC]]" not in self.body:
            self.render_toc_placeholder(self.doc)
        self.render_blocks(tree.children, self.doc)
        self.finalize_toc()
        self._trim_trailing_empty_paragraphs()
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        words = len(re.findall(r"\w+", self.body))
        self.stats["words_md"] = words
        return self.stats


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--brand", default=str(DEFAULT_BRAND))
    ap.add_argument("--pdf", action="store_true", help="also convert to PDF with LibreOffice (for visual QA)")
    a = ap.parse_args()
    brand = json.loads(Path(a.brand).read_text())
    b = HandoverBuilder(Path(a.markdown), brand)
    stats = b.build(Path(a.out))
    print(json.dumps({"output": a.out, **stats}, indent=2))
    for w in b.warnings:
        print("WARNING:", w, file=sys.stderr)
    if a.pdf:
        outdir = str(Path(a.out).parent)
        try:
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, a.out],
                           check=True, capture_output=True, timeout=300)
            print("PDF:", str(Path(a.out).with_suffix(".pdf")))
        except Exception as e:  # noqa
            print("PDF conversion failed:", e, file=sys.stderr)


if __name__ == "__main__":
    main()
