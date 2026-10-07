"""In-memory Excel exchange and finite natural-language volume matrix proposals."""

from io import BytesIO
from typing import Literal
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import Field
from production_wheel.rule_models import MatrixPair
from .models import StrictModel

MAX_WORKBOOK_BYTES = 2_000_000


def volume_ids(values):
    """Normalize positive decimal volume identities, rejecting duplicates and oversized axes."""
    result = [MatrixPair(volume_a=str(v), volume_b=str(v), status='Y').volume_a for v in values]
    if not 1 <= len(result) <= 100 or len(set(result)) != len(result):
        raise ValueError('Use 1–100 unique positive volumes in litres')
    return result


def validate_matrix_rows(rows):
    """Return a complete symmetric matrix with normalized axes and compatible diagonals."""
    values = [MatrixPair.model_validate(row) for row in rows]
    pairs = {(r.volume_a, r.volume_b): r.status for r in values}
    volumes = volume_ids(sorted({v for pair in pairs for v in pair}, key=float))
    if len(pairs) != len(values) or any((a, b) not in pairs or pairs[a, b] != pairs.get((b, a))
                                      for a in volumes for b in volumes):
        raise ValueError('Volume matrix must be complete, unique and symmetric; fill both halves')
    return [r.model_dump() for r in values]


class MatrixDescription(StrictModel):
    """Explicit replacement matrix: listed volumes, compatible families, default and exceptions."""

    volumes: list[str] = Field(min_length=1, max_length=100)
    compatible_groups: list[list[str]] = Field(default_factory=list, max_length=100)
    default_status: Literal['Y', 'AVOID', 'N']
    overrides: list[MatrixPair] = Field(default_factory=list, max_length=10000)


def expand_matrix_description(description):
    """Compile reviewed families into all ordered pairs without executing model code."""
    spec = MatrixDescription.model_validate(description)
    volumes = volume_ids(spec.volumes)
    pairs = {(a, b): 'Y' if a == b else spec.default_status for a in volumes for b in volumes}
    for family in spec.compatible_groups:
        family = volume_ids(family)
        if not set(family) <= set(volumes):
            raise ValueError('Compatible family refers to an unlisted volume')
        for a in family:
            for b in family:
                pairs[a, b] = 'Y'
    overrides = {}
    for row in spec.overrides:
        if row.volume_a not in volumes or row.volume_b not in volumes:
            raise ValueError('Matrix exception refers to an unlisted volume')
        key = tuple(sorted((row.volume_a, row.volume_b)))
        if key in overrides and overrides[key] != row.status:
            raise ValueError('Conflicting symmetric matrix exceptions')
        overrides[key] = row.status
        pairs[row.volume_a, row.volume_b] = pairs[row.volume_b, row.volume_a] = row.status
    return validate_matrix_rows([{'volume_a': a, 'volume_b': b, 'status': status}
                                 for (a, b), status in pairs.items()])


def matrix_workbook(rows):
    """Generate the editable Excel template from accepted profile rows using the app's Excel library."""
    rows = validate_matrix_rows(rows)
    volumes = sorted({r['volume_a'] for r in rows}, key=float)
    pairs = {(r['volume_a'], r['volume_b']): r['status'] for r in rows}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Volume compatibility'
    sheet.append(['Volume (litres)', *[float(v) for v in volumes]])
    for a in volumes:
        sheet.append([float(a), *[pairs[a, b] for b in volumes]])
    sheet.freeze_panes = 'B2'
    sheet.column_dimensions['A'].width = 20
    for column in range(2, len(volumes) + 2):
        sheet.column_dimensions[get_column_letter(column)].width = 10
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(horizontal='center')
            cell.font = Font(name='Aptos', size=11)
            if cell.row == 1 or cell.column == 1:
                cell.fill = PatternFill('solid', fgColor='16324F')
                cell.font = Font(name='Aptos', size=11, bold=True, color='FFFFFF')
    choices = DataValidation(type='list', formula1='"Y,AVOID,N"', allow_blank=False)
    choices.errorTitle = 'Choose a compatibility status'
    choices.error = 'Use Y, AVOID or N. Match the opposite cell; diagonal cells must be Y.'
    choices.showErrorMessage = True
    sheet.add_data_validation(choices)
    choices.add('B2:CW101')
    notes = workbook.create_sheet('Instructions')
    for text in (
        'Volume compatibility template',
        'Edit the Volume compatibility sheet. Volumes are positive numbers in litres.',
        'Y = compatible. AVOID = discouraged (blocked in HARD mode). N = forbidden in every mode.',
        'Fill every pair using Y, AVOID or N. Both halves must match; diagonal cells must be Y.',
        'To add or remove volumes, edit BOTH the header row and the matching first-column rows. Maximum 100 volumes.',
        'Keep the same volume order on both axes. Do not use formulas or merged cells in the matrix.',
        'Upload the edited .xlsx, review the preview and click Accept. Save profile then persists it.',
        'Accept replaces the whole matrix. Source products with volumes outside it will require a compatible profile.',
    ):
        notes.append([text])
    notes.column_dimensions['A'].width = 105
    for row in notes:
        row[0].font = Font(name='Aptos', size=11)
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
        notes.row_dimensions[row[0].row].height = 32
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_matrix_workbook(content):
    """Validate a bounded uploaded square matrix; formulas and partial matrices never pass."""
    if len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError('Compatibility workbook exceeds 2 MB')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 100 or sum(i.file_size for i in archive.infolist()) > 10_000_000:
                raise ValueError('Compatibility workbook expands beyond its size limit')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if 'Volume compatibility' not in workbook.sheetnames:
                raise ValueError('Use the template sheet named Volume compatibility')
            sheet = workbook['Volume compatibility']
            if sheet.max_row > 101 or sheet.max_column > 101:
                raise ValueError('Matrix supports at most 100 volumes; remove unused rows and columns')
            grid = list(sheet.values)
            while grid and all(value is None for value in grid[-1]):
                grid.pop()
            if not grid or len(grid) < 2:
                raise ValueError('Matrix is empty')
            width = max(i + 1 for row in grid for i, value in enumerate(row) if value is not None)
            grid = [row[:width] for row in grid]
            columns = volume_ids(grid[0][1:])
            rows = volume_ids([row[0] for row in grid[1:]])
            if rows != columns:
                raise ValueError('Row and column volumes must match in the same order')
            values = [{'volume_a': a, 'volume_b': b, 'status': str(grid[i][j] or '').strip().upper()}
                      for i, a in enumerate(rows, 1) for j, b in enumerate(columns, 1)]
            return validate_matrix_rows(values)
        finally:
            workbook.close()
    except (BadZipFile, KeyError, IndexError, TypeError, SyntaxError, EOFError) as exc:
        raise ValueError('Invalid compatibility .xlsx; download and use the template') from exc
