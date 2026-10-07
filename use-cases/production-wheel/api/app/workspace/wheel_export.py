"""Export one persisted Pareto point in its uploaded production-wheel column layout.

All inputs come from HANA-backed repository rows/artifacts. No local workbook
path is read at runtime and no original workbook is overwritten.
"""

from __future__ import annotations

import io
import warnings

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from production_wheel.extraction.common import as_text, find_table
from production_wheel.extraction.line_format import format_lines

# Used for migrated snapshots with canonical rows but no retained upload artifact.
SOURCE_FIELDS = {
    "Production Frequency": "production_frequency",
    "SEFI": "sefi",
    "SEFI description": "sefi_description",
    "Lot size considered": "lot_size_considered_litres",
    "Material": "material",
    "Material Description": "material_description",
    "Pallet in litres": "pallet_litres_resolved",
    "MRPCn": "mrp_controller",
    "Can code": "can_code",
    "Can description": "can_description",
    "Total FINI forecast (litres)": "forecast_litres_12m",
    "Avg daily demand": "avg_daily_demand",
    "Recipe Combination": "recipe_combination_original",
    "RESET recipe combination": "reset_recipe_flag",
    "Volume": "package_volume",
    "PCK code": "pck_code",
    "Decoration type": "decoration_type",
    "Can Shape": "can_shape",
    "Filling line": "filling_line_raw",
    "FINI/P-site": "fini_plant_key",
    "SEFI/Plant": "sefi_plant_key",
    "Subgroup": "proposed_subgroup_code",
    "Selected PV": "selected_pv",
    "RoC selected PV": "fini_adjusted_coverage_days",
}
PROPOSAL_FIELDS = {
    "Subgroup": "proposed_subgroup_code",
    "New recipe subgroup": "proposed_subgroup_code",
    "Selected PV": "selected_pv",
    "RoC selected PV": "fini_adjusted_coverage_days",
    "SEFI litres lot size of selected PV": "nominal_lot_litres",
    "Sefi qty as per recipe subgroup": "selected_pallet_allocation_litres",
    "Pallets as per subgroup": "selected_pallet_count",
    "Batches per week": "frequency_per_week",
}
EXTRA_FIELDS = {
    "Plant": "plant",
    "Point": "point_index",
    "Proposal group ID": "group_id",
    "Proposal status": "model_status",
    "Exclusion reason": "exclusion_reason",
    "Coverage basis": "active_coverage_basis",
    "Group coverage days": "coverage_days",
    "Effective batch litres": "effective_batch_litres",
    "Nominal FINI allocation litres": "nominal_allocation_litres",
    "Changeover contribution": "j_ch_contribution",
    "Selected filling line": "selected_line",
}


def source_layout(service, dataset_id: str) -> tuple[list[str], dict]:
    """Read uploaded main-table headers and source rows by sheet/row identity.

    A migrated snapshot without uploaded source metadata uses canonical headers.
    A missing promised artifact raises an error instead of silently losing columns.
    """
    metadata = service.repo.get("datasets", dataset_id)["metadata"]
    files = metadata.get("source_files", [])
    if not files:
        return list(SOURCE_FIELDS), {}
    for filename in files:
        content = service.repo.artifact(dataset_id, filename)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            book = load_workbook(io.BytesIO(content), data_only=True)
        try:
            if not any("FINI_grouping_for_wheel" in sheet.tables for sheet in book):
                continue
            table = find_table(book, "FINI_grouping_for_wheel")
            header_cells = book[table.sheet][
                table.reference.split(":")[0] : table.reference.split(":")[1]
            ][0]
            headers = [str(cell.value) for cell in header_cells]
            return headers, {
                (table.sheet, row["_source_row"]): row for row in table.records
            }
        finally:
            book.close()
    raise ValueError("Retained sources contain no production-wheel main table")


def wheel_rows(
    headers: list[str], originals: dict, members: list[dict], config: dict
) -> list[dict]:
    """Join source lineage to selected-point rows; overwrite proposal-dependent columns.

    All source FINIs remain present, including excluded/unassigned rows whose
    proposal fields stay blank. Original Filling line stays verbatim; corrected
    eligibility and group-common alternatives are separate, explicitly named columns.
    """
    output = []
    for member in sorted(
        members,
        key=lambda row: (
            str(row.get("source_sheet", "")),
            float(row.get("source_row") or 0),
        ),
    ):
        original = originals.get(
            (member.get("source_sheet"), int(float(member.get("source_row") or 0)))
        )
        if originals and original is None:
            raise ValueError("Source lineage is missing for an exported member")
        if original and as_text(original.get("Material")) != as_text(
            member.get("material")
        ):
            raise ValueError("Source lineage does not match the exported material")
        row = {
            header: (original or {}).get(
                header, member.get(SOURCE_FIELDS.get(header, ""))
            )
            for header in headers
        }
        selected = bool(member.get("proposed_subgroup"))
        for header, field in PROPOSAL_FIELDS.items():
            if header in headers:
                row[header] = member.get(field) if selected else None
        group_demand = member.get("group_demand_litres")
        if "% forecast within subgroup" in headers:
            row["% forecast within subgroup"] = (
                float(member.get("forecast_litres_12m") or 0) / float(group_demand)
                if selected and group_demand
                else None
            )
        if "Number of batches" in headers:
            row["Number of batches"] = (
                float(member.get("frequency_per_week") or 0)
                * config.get("productive_weeks", 50)
                if selected
                else None
            )
        if "SEFI/Plant/Recipe/Subgroup" in headers:
            row["SEFI/Plant/Recipe/Subgroup"] = (
                "/".join(
                    str(member.get(k) or "")
                    for k in (
                        "sefi",
                        "plant",
                        "effective_recipe_combination",
                        "proposed_subgroup_code",
                    )
                )
                if selected
                else None
            )
        row.update({label: member.get(field) for label, field in EXTRA_FIELDS.items()})
        if not selected:
            row["Proposal status"] = "No proposed assignment"
        raw = row.get("Filling line") or member.get("filling_line_raw")
        effective = member.get("filling_line_correction_raw") or raw
        row["Effective eligible lines"] = effective
        row["Common eligible lines"] = (
            format_lines(
                str(member.get("common_lines") or "").split("|")
                if member.get("common_lines")
                else [],
                effective,
            )
            if selected
            else None
        )
        output.append(row)
    return output


def _sheet(book, name: str, headers: list[str], rows: list[dict]) -> None:
    """Write typed values with readable headers, filters and frozen identity columns."""
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(key) for key in headers])
    for cells in sheet:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = (
                    "s"  # Workbook source/model strings must never become formulas.
                )
            elif isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="007D8A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 48
    for index, label in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            42, max(18, len(label) + 2)
        )
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions


def export_wheel(service, run_id: str, point_index: int) -> bytes:
    """Return an XLSX for a completed run's explicit point, preserving source layout."""
    result = service.results(run_id)
    if not any(int(p["point_index"]) == point_index for p in result["points"]):
        raise ValueError("Unknown frontier point")
    members = service.repo.rows(run_id, "members", {"point_index": point_index})
    if not members:
        raise ValueError("No persisted members for this point")
    headers, originals = source_layout(service, result["run"]["dataset_id"])
    columns = list(
        dict.fromkeys(
            [
                *headers,
                *EXTRA_FIELDS,
                "Effective eligible lines",
                "Common eligible lines",
            ]
        )
    )
    rows = wheel_rows(
        headers, originals, members, result["run"].get("request", {}).get("config", {})
    )
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "Production wheel", columns, rows)
    notes = [
        {"Field": "Run / point", "Meaning": f"{run_id} / {point_index}"},
        {
            "Field": "Source columns",
            "Meaning": "Original headers and source row order. Unchanged source columns retain uploaded cached values; original formulas, macros and pivot tables are not copied.",
        },
        {
            "Field": "Proposal columns",
            "Meaning": "Subgroup, selected PV, subgroup forecast share, selected-PV coverage, batch frequency/count, subgroup allocation and pallets are replaced by selected-point values. Unassigned FINIs have blank proposal fields.",
        },
        {
            "Field": "RoC selected PV",
            "Meaning": "Pallet-adjusted FINI coverage in days. Group coverage days separately uses the selected coverage basis.",
        },
        {
            "Field": "Filling line",
            "Meaning": "Original source text. Effective eligible lines includes any FixedLine correction. Common eligible lines uses that notation for alternatives shared by group members; it is not a scheduled line assignment.",
        },
        {
            "Field": "Changeover",
            "Meaning": "Recurring within-group FINI changes per productive week. Initial/between-group setups and cleaning time are excluded. Group contributions repeat per member and must not be summed across this sheet.",
        },
        {
            "Field": "Evidence",
            "Meaning": "Structural validity does not establish business acceptance or a complete globally optimal frontier. Review the run metadata and audit views.",
        },
    ]
    _sheet(book, "Export notes", ["Field", "Meaning"], notes)
    book["Export notes"].column_dimensions["B"].width = 110
    for cells in book["Export notes"].iter_rows(min_row=2):
        cells[1].alignment = Alignment(wrap_text=True, vertical="top")
        book["Export notes"].row_dimensions[cells[1].row].height = 64
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()
