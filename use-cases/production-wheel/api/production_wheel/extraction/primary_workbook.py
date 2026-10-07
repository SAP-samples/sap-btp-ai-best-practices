"""Extract canonical optimizer inputs from the primary production workbook."""

from __future__ import annotations

import collections
import itertools
import re
import warnings
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from production_wheel.extraction.common import (
    TableData,
    as_number,
    as_text,
    find_table,
    parse_lines,
)

CANONICAL_FACTOR = 0.90
HORIZON_MONTHS = 12
PRODUCTIVE_WEEKS = 50
DEMAND_DAYS = 250
MODELED_FREQUENCY_CODES = {"01W", "02W"}


def normalize_production_frequency(value: Any) -> str | None:
    """Return a zero-padded weekly frequency code from workbook text.

    Args:
        value: Source frequency such as ``01W. Weekly`` or ``2W``.

    Returns:
        A normalized code such as ``01W``, or ``None`` when unrecognized.
    """

    text = as_text(value)
    match = re.match(r"^(\d{1,2})\s*W\b", text or "", flags=re.IGNORECASE)
    return f"{int(match.group(1)):02d}W" if match else None


def _positive(value: Any) -> bool:
    """Return whether a source value is a finite positive number."""

    number = as_number(value)
    return number is not None and number > 0


def classify_fini_admission(
    source: Mapping[str, Any],
    production_versions: Iterable[Mapping[str, Any]] = (),
    *,
    require_resolved_pallet: bool = True,
    modeled_frequencies: Iterable[str] = MODELED_FREQUENCY_CODES,
) -> dict[str, Any]:
    """Classify one FINI by frequency first and data quality second.

    Args:
        source: Canonical FINI row.
        production_versions: Finite PV catalog used to validate fixed and
            optimized-PV eligibility.
        require_resolved_pallet: Whether unresolved pallet data is final.
        modeled_frequencies: Weekly codes admitted by this dataset; defaults retain
            the historical CLI scope.

    Returns:
        A copied row with fixed-PV and optimized-PV disposition fields.
    """

    row = dict(source)
    code = normalize_production_frequency(row.get("production_frequency"))
    row["production_frequency_code"] = code
    if code not in modeled_frequencies:
        reason = (
            f"production_frequency_{code.lower()}_out_of_scope"
            if code
            else "production_frequency_unrecognized_out_of_scope"
        )
        row.update(
            model_status="out_of_scope",
            exclusion_reason=reason,
            optimized_pv_model_status="out_of_scope",
            optimized_pv_exclusion_reason=reason,
        )
        return row

    if not row.get("material"):
        quality_reason = "missing_material"
    elif not row.get("sefi"):
        quality_reason = "missing_sefi"
    elif as_number(row.get("forecast_litres_12m")) is None:
        quality_reason = "missing_demand"
    elif not _positive(row.get("forecast_litres_12m")):
        quality_reason = "nonpositive_demand"
    elif row.get("line_data_status") != "known":
        quality_reason = "missing_line_evidence"
    elif not row.get("eligible_lines"):
        quality_reason = "no_eligible_line"
    elif as_number(row.get("package_volume")) is None:
        quality_reason = "missing_package_volume"
    elif not _positive(row.get("package_volume")):
        quality_reason = "nonpositive_package_volume"
    elif require_resolved_pallet and as_number(row.get("pallet_litres_resolved")) is None:
        quality_reason = "missing_pallet_data"
    elif require_resolved_pallet and not _positive(row.get("pallet_litres_resolved")):
        quality_reason = "nonpositive_pallet_data"
    else:
        quality_reason = None

    versions = tuple(production_versions)
    block_versions = [
        version
        for version in versions
        if str(version.get("plant")) == str(row.get("plant"))
        and str(version.get("sefi")) == str(row.get("sefi"))
        and _positive(version.get("lot_size_litres"))
    ]
    unique_lots: dict[str, set[float]] = collections.defaultdict(set)
    for version in block_versions:
        pv_id = as_text(version.get("production_version"))
        if pv_id:
            unique_lots[pv_id].add(float(as_number(version["lot_size_litres"])))
    optimized_available = any(len(lots) == 1 for lots in unique_lots.values())

    if quality_reason:
        fixed_reason = optimized_reason = quality_reason
    else:
        fixed_pv = as_text(row.get("fixed_pv"))
        if fixed_pv is None:
            fixed_reason = "missing_fixed_pv"
        elif versions and len(unique_lots.get(fixed_pv, set())) != 1:
            fixed_reason = "fixed_pv_not_unique_positive_lot"
        else:
            fixed_reason = None
        optimized_reason = (
            "no_unique_positive_optimized_pv"
            if versions and not optimized_available
            else None
        )
    row.update(
        model_status="excluded" if fixed_reason else "modeled",
        exclusion_reason=fixed_reason,
        optimized_pv_model_status="excluded" if optimized_reason else "modeled",
        optimized_pv_exclusion_reason=optimized_reason,
    )
    return row


def classify_fini_rows(
    rows: Iterable[Mapping[str, Any]],
    production_versions: Iterable[Mapping[str, Any]] = (),
    *,
    require_resolved_pallet: bool = True,
) -> list[dict[str, Any]]:
    """Classify canonical FINIs without using historical grouping fields."""

    versions = tuple(production_versions)
    return [
        classify_fini_admission(
            row,
            versions,
            require_resolved_pallet=require_resolved_pallet,
        )
        for row in rows
    ]


def _issue(
    rule_id: str,
    severity: str,
    entity: str,
    key: str | None,
    observed: Any,
    expected: Any,
    action: str,
) -> dict[str, Any]:
    """Build one structured extraction-validation issue."""

    return {
        "rule_id": rule_id,
        "severity": severity,
        "entity": entity,
        "entity_key": key,
        "observed": observed,
        "expected": expected,
        "action": action,
    }


def _material(value: Any) -> str | None:
    """Normalize a workbook material identifier."""

    return as_text(value)


def _simple_records(table: TableData, field_map: Mapping[str, str]) -> list[dict[str, Any]]:
    """Rename selected source-table fields while retaining source rows."""

    return [
        {target: record.get(source) for source, target in field_map.items()}
        | {"source_sheet": table.sheet, "source_row": record["_source_row"]}
        for record in table.records
    ]


def _canonical_fini(main: TableData, site: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Create FINI master and long-form legacy coverage records."""

    fini_rows: list[dict[str, Any]] = []
    legacy_coverage: list[dict[str, Any]] = []
    for record in main.records:
        material = _material(record.get("Material"))
        sefi = as_text(record.get("SEFI"))
        lines = parse_lines(record.get("Filling line"))
        subgroup = as_text(record.get("Subgroup"))
        fixed_pv = as_text(record.get("Selected PV"))
        legacy_baseline_group = as_text(record.get("SEFI/Plant/Recipe/Subgroup"))
        effective_recipe = as_text(record.get("RESET recipe combination")) or as_text(
            record.get("Recipe Combination")
        )
        baseline_group = (
            f"{sefi}/{site}/{effective_recipe}/{subgroup}"
            if sefi and effective_recipe and subgroup and fixed_pv
            else None
        )
        baseline_assigned = bool(subgroup and fixed_pv and baseline_group)
        line_known = lines is not None
        demand = as_number(record.get("Total FINI forecast (litres)"))
        package_volume = as_number(record.get("Volume"))
        pallet = as_number(record.get("Pallet in litres"))
        fini_rows.append(
            {
                "plant": site,
                "material": material,
                "fini_plant_key": as_text(record.get("FINI/P-site")) or f"{material}/{site}",
                "sefi": sefi,
                "sefi_plant_key": as_text(record.get("SEFI/Plant")) or f"{sefi}/{site}",
                "production_frequency": as_text(record.get("Production Frequency")),
                "production_frequency_code": normalize_production_frequency(
                    record.get("Production Frequency")
                ),
                "sefi_description": as_text(record.get("SEFI description")),
                "material_description": as_text(record.get("Material Description")),
                "mrp_controller": as_text(record.get("MRPCn")),
                "forecast_litres_12m": demand,
                "avg_daily_demand": as_number(record.get("Avg daily demand")),
                "sefi_recipe_forecast_litres": as_number(record.get("SEFI-recipe forecast")),
                "sefi_forecast_litres": as_number(record.get("SEFI Forecast (litres)")),
                "lot_size_considered_litres": as_number(record.get("Lot size considered")),
                "package_volume": package_volume,
                "pck_code": as_text(record.get("PCK code")),
                "can_code": as_text(record.get("Can code")),
                "can_description": as_text(record.get("Can description")),
                "decoration_type": as_text(record.get("Decoration type")),
                "can_shape": as_text(record.get("Can Shape")),
                "filling_line_raw": as_text(record.get("Filling line")),
                "eligible_lines": "|".join(map(str, lines)) if lines is not None else None,
                "line_data_status": "known" if line_known else "missing",
                "pallet_litres_primary": pallet,
                "pallet_litres_resolved": pallet,
                "pallet_resolution_source": "PRIMARY" if pallet and pallet > 0 else None,
                "pallet_candidate_count": 0,
                "pallet_conflict": 0,
                "recipe_combination_original": as_text(record.get("Recipe Combination")),
                "reset_recipe_flag": as_text(record.get("RESET recipe combination")),
                "effective_recipe_combination": effective_recipe,
                "baseline_subgroup": subgroup,
                "fixed_pv": fixed_pv,
                "baseline_group_key": baseline_group,
                "legacy_baseline_group_key": legacy_baseline_group,
                "baseline_assigned": int(baseline_assigned),
                "model_status": "pending_data_quality",
                "exclusion_reason": None,
                "optimized_pv_model_status": "pending_data_quality",
                "optimized_pv_exclusion_reason": None,
                "legacy_coverage_selected_pv_days": as_number(record.get("RoC selected PV")),
                "legacy_selected_pv_lot_litres": as_number(
                    record.get("SEFI litres lot size of selected PV")
                ),
                "legacy_pallets_per_subgroup": as_number(record.get("Pallets as per subgroup")),
                "source_sheet": main.sheet,
                "source_row": record["_source_row"],
            }
        )
        for number in range(1, 17):
            displayed = as_number(record.get(f"PV - {number}"))
            if displayed is None:
                continue
            legacy_coverage.append(
                {
                    "plant": site,
                    "material": material,
                    "production_version": f"PV - {number}",
                    "legacy_coverage_days_displayed": displayed,
                    "legacy_factor": 0.85 if number == 1 else CANONICAL_FACTOR,
                    "canonical_factor": CANONICAL_FACTOR,
                    "source_sheet": main.sheet,
                    "source_row": record["_source_row"],
                }
            )
    return fini_rows, legacy_coverage


def _line_tables(
    fini_rows: Iterable[Mapping[str, Any]], pck_table: TableData, site: str
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Create tri-state FINI and PCK filling-line bridges."""

    fini_bridge: list[dict[str, Any]] = []
    for row in fini_rows:
        known = row["line_data_status"] == "known"
        eligible = set(int(value) for value in (row.get("eligible_lines") or "").split("|") if value)
        for line in range(1, 7):
            fini_bridge.append(
                {
                    "plant": row["plant"],
                    "material": row["material"],
                    "pck_code": row["pck_code"],
                    "filling_line": line,
                    "eligible": int(line in eligible) if known else None,
                    "line_data_status": row["line_data_status"],
                    "source_pattern": row["filling_line_raw"],
                    "source_row": row["source_row"],
                }
            )
    pck_bridge: list[dict[str, Any]] = []
    for record in pck_table.records:
        pattern = record.get("Value")
        parsed = parse_lines(pattern)
        eligible = set(parsed or ())
        for line in range(1, 7):
            pck_bridge.append(
                {
                    "plant": site,
                    "pck_code": as_text(record.get("Name")),
                    "filling_line": line,
                    "eligible": int(line in eligible) if parsed is not None else None,
                    "line_data_status": "known" if parsed is not None else "missing",
                    "source_pattern": as_text(pattern),
                    "source_sheet": pck_table.sheet,
                    "source_row": record["_source_row"],
                }
            )
    return fini_bridge, pck_bridge


def _production_versions(table: TableData, site: str) -> list[dict[str, Any]]:
    """Create canonical finite production-version records."""

    result: list[dict[str, Any]] = []
    for record in table.records:
        production_version = as_text(record.get("Clasificador"))
        result.append(
            {
                "plant": site,
                "sefi": as_text(record.get("SEFI")),
                "sefi_plant_key": as_text(record.get("SEFI/P-site")),
                "production_version": production_version,
                "lot_size_litres": as_number(record.get("Lot size (Litres)")),
                "pv_lot_key": as_text(record.get("PV lot")),
                "canonical_factor": CANONICAL_FACTOR,
                "legacy_factor": 0.85 if production_version == "PV - 1" else CANONICAL_FACTOR,
                "source_sheet": table.sheet,
                "source_row": record["_source_row"],
            }
        )
    return result


def _package_volume_catalog(fini_rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Aggregate FINI count and demand for every observed package volume."""

    catalog: dict[tuple[float, str], dict[str, Any]] = {}
    for row in fini_rows:
        volume = as_number(row.get("package_volume"))
        if volume is None:
            continue
        for population, include in (
            ("all", True),
            ("modeled", row.get("model_status") == "modeled"),
        ):
            if not include:
                continue
            key = (volume, population)
            target = catalog.setdefault(
                key,
                {
                    "plant": row["plant"],
                    "package_volume": volume,
                    "population": population,
                    "fini_count": 0,
                    "forecast_litres_12m": 0.0,
                    "pck_codes": set(),
                },
            )
            target["fini_count"] += 1
            target["forecast_litres_12m"] += as_number(row.get("forecast_litres_12m")) or 0.0
            if row.get("pck_code"):
                target["pck_codes"].add(row["pck_code"])
    rows: list[dict[str, Any]] = []
    for key in sorted(catalog, key=lambda item: (item[1], item[0])):
        row = catalog[key]
        rows.append(
            row
            | {
                "forecast_litres_12m": round(row["forecast_litres_12m"], 6),
                "pck_count": len(row["pck_codes"]),
                "pck_codes": "|".join(sorted(row["pck_codes"])),
            }
        )
        rows[-1].pop("pck_codes", None)
        rows[-1]["pck_codes"] = "|".join(sorted(row["pck_codes"]))
    return rows


def _baseline_assignments(fini_rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    """Return one record per historically assigned FINI."""

    fields = (
        "plant",
        "material",
        "sefi",
        "baseline_subgroup",
        "fixed_pv",
        "baseline_group_key",
        "forecast_litres_12m",
        "model_status",
        "exclusion_reason",
        "source_row",
    )
    return [{field: row.get(field) for field in fields} for row in fini_rows if row["baseline_assigned"]]


def _group_metrics(
    fini_rows: Iterable[Mapping[str, Any]],
    production_versions: Iterable[Mapping[str, Any]],
    character_map: Mapping[tuple[str, str], str],
) -> list[dict[str, Any]]:
    """Recompute historical and modeled baseline group metrics from primitives."""

    pv_index = {
        (str(row["sefi"]), str(row["production_version"])): as_number(row["lot_size_litres"])
        for row in production_versions
    }
    assigned = [row for row in fini_rows if row["baseline_assigned"]]
    result: list[dict[str, Any]] = []
    for population in ("historical", "modeled"):
        included = assigned if population == "historical" else [row for row in assigned if row["model_status"] == "modeled"]
        groups: dict[str, list[Mapping[str, Any]]] = collections.defaultdict(list)
        for row in included:
            groups[str(row["baseline_group_key"])].append(row)
        for group_key, members in sorted(groups.items()):
            first = members[0]
            pv_values = {str(member["fixed_pv"]) for member in members if member.get("fixed_pv")}
            fixed_pv = next(iter(pv_values)) if len(pv_values) == 1 else None
            lot_size = pv_index.get((str(first["sefi"]), str(fixed_pv))) if fixed_pv else None
            demand = sum(as_number(member.get("forecast_litres_12m")) or 0.0 for member in members)
            line_sets = [
                set(int(value) for value in str(member["eligible_lines"]).split("|"))
                for member in members
                if member.get("eligible_lines")
            ]
            missing_line = len(line_sets) != len(members)
            common_lines = set.intersection(*line_sets) if line_sets and not missing_line else set()
            effective_batch = (lot_size or 0.0) * CANONICAL_FACTOR
            coverage = DEMAND_DAYS * effective_batch / demand if effective_batch > 0 and demand > 0 else None
            frequency = demand / (PRODUCTIVE_WEEKS * effective_batch) if effective_batch > 0 else None
            pcks = {str(member["pck_code"]) for member in members if member.get("pck_code")}
            result.append(
                {
                    "population": population,
                    "plant": first["plant"],
                    "sefi": first["sefi"],
                    "baseline_group_key": group_key,
                    "baseline_subgroup": first["baseline_subgroup"],
                    "sap_character": character_map.get(
                        (str(first["sefi"]), str(first["baseline_subgroup"]))
                    ),
                    "member_count": len(members),
                    "members": "|".join(sorted(str(member["material"]) for member in members)),
                    "fixed_pv": fixed_pv,
                    "fixed_pv_consistent": int(len(pv_values) == 1),
                    "nominal_lot_litres": lot_size,
                    "effective_batch_litres": effective_batch if lot_size else None,
                    "forecast_litres_12m": demand,
                    "coverage_days": coverage,
                    "runs_per_week": frequency,
                    "j_ch_contribution": frequency * (len(members) - 1) if frequency is not None else None,
                    "common_lines": "|".join(map(str, sorted(common_lines))) if common_lines else None,
                    "line_status": "missing" if missing_line else ("feasible" if common_lines else "infeasible"),
                    "distinct_pck_count": len(pcks),
                    "pck_codes": "|".join(sorted(pcks)),
                }
            )
    return result


def _character_tables(
    character_table: TableData,
    compatibility_sheet: Any,
    baseline_groups: Iterable[Mapping[str, Any]],
    site: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], dict[tuple[str, str], str]]:
    """Extract SAP-character catalogs, compatibility, and observed capacity."""

    character_map: dict[tuple[str, str], str] = {}
    character_catalog: list[dict[str, Any]] = []
    for record in character_table.records:
        sefi = as_text(record.get("SEFI"))
        subgroup = as_text(record.get("Subgroup"))
        character = as_text(record.get("SAP character"))
        if sefi and subgroup and character:
            character_map[(sefi, subgroup)] = character
        character_catalog.append(
            {
                "plant": site,
                "sefi": sefi,
                "subgroup": subgroup,
                "sap_character": character,
                "source_sheet": character_table.sheet,
                "source_row": record["_source_row"],
            }
        )
    column_labels: list[str] = []
    for column in range(2, compatibility_sheet.max_column + 1):
        value = as_text(compatibility_sheet.cell(1, column).value)
        if value is None:
            break
        column_labels.append(value)
    row_labels: list[str] = []
    for row in range(2, compatibility_sheet.max_row + 1):
        value = as_text(compatibility_sheet.cell(row, 1).value)
        if value is None:
            break
        row_labels.append(value)
    if row_labels != column_labels or not row_labels:
        raise ValueError("SAP-character compatibility matrix must have identical row/column labels")
    labels = row_labels
    source_range = f"A1:{get_column_letter(len(labels) + 1)}{len(labels) + 1}"
    compatibility = [
        {
            "char_a": labels[left],
            "char_b": labels[right],
            "compatible": compatibility_sheet.cell(left + 2, right + 2).value,
            "source_sheet": compatibility_sheet.title,
            "source_range": source_range,
        }
        for left in range(len(labels))
        for right in range(len(labels))
    ]
    capacity = _character_capacity(baseline_groups, site)
    return character_catalog, compatibility, capacity, character_map


def _character_capacity(
    baseline_groups: Iterable[Mapping[str, Any]], site: str
) -> list[dict[str, Any]]:
    """Rebuild the observed character/line capacity diagnostic."""

    groups = tuple(baseline_groups)
    # Capacity is an observed integration diagnostic, not grouping feasibility.
    line_sets_by_character: dict[str, set[int]] = collections.defaultdict(set)
    for group in groups:
        if group.get("population") != "modeled" or not group.get("sap_character") or not group.get("common_lines"):
            continue
        line_sets_by_character[str(group["sap_character"])].update(
            int(value) for value in str(group["common_lines"]).split("|")
        )
    capacity: list[dict[str, Any]] = []
    used: dict[tuple[str, str], int] = collections.Counter(
        (str(group["sefi"]), str(group["common_lines"]))
        for group in groups
        if group.get("population") == "modeled" and group.get("common_lines")
    )
    for (sefi, line_set), count in sorted(used.items()):
        target = tuple(int(value) for value in line_set.split("|"))
        available = sum(
            1 for lines in line_sets_by_character.values() if tuple(sorted(lines)) == target
        )
        capacity.append(
            {
                "plant": site,
                "sefi": sefi,
                "line_set": line_set,
                "groups_today": count,
                "characters_available_observed": available,
                "utilisation_pct": 100.0 * count / available if available else None,
                "status": "integration_diagnostic_only",
            }
        )
    return capacity


def _legacy_kpis(group_metrics: Iterable[Mapping[str, Any]], changeover_table: TableData) -> list[dict[str, Any]]:
    """Create recomputed and cached legacy KPI records."""

    modeled = [row for row in group_metrics if row["population"] == "modeled"]
    recalculated = {
        "basis": "modeled_recalculated_canonical_factor_0.90",
        "runs_per_week": sum(as_number(row.get("runs_per_week")) or 0.0 for row in modeled),
        "fini_changes_per_week": sum(as_number(row.get("j_ch_contribution")) or 0.0 for row in modeled),
        "pck_changes_per_week": None,
        "same_pck_changes_per_week": None,
        "status": "canonical_development",
    }
    source = changeover_table.records[0]
    cached = {
        "basis": "workbook_cached_changeover_summary",
        "runs_per_week": as_number(source.get("Batches per week")),
        "fini_changes_per_week": as_number(source.get("Weekly changes FINI-FINI")),
        "pck_changes_per_week": as_number(source.get("Weekly changes FINI-FINI with PCK change")),
        "same_pck_changes_per_week": as_number(source.get("Weekly FINI-FINI with same PCK")),
        "status": "legacy_audit_only",
    }
    return [recalculated, cached]


def refresh_scope_dependent_tables(
    tables: dict[str, list[dict[str, Any]]],
) -> None:
    """Refresh derived tables after final pallet-aware admission.

    Args:
        tables: Mutable canonical extraction table mapping.

    Returns:
        ``None``; model-population-dependent tables are replaced in place.
    """

    fini_rows = tables["fini_master"]
    versions = tables["production_versions"]
    character_map = {
        (str(row.get("sefi") or ""), str(row.get("subgroup") or "")): str(
            row.get("sap_character") or ""
        )
        for row in tables["sap_character_catalog"]
    }
    groups = _group_metrics(fini_rows, versions, character_map)
    tables["baseline_assignments"] = _baseline_assignments(fini_rows)
    tables["package_volume_catalog"] = _package_volume_catalog(fini_rows)
    tables["baseline_group_metrics"] = groups
    site = str(tables["site_parameters"][0]["plant"])
    tables["sap_character_capacity"] = _character_capacity(groups, site)
    if tables.get("legacy_kpi"):
        modeled = [row for row in groups if row["population"] == "modeled"]
        tables["legacy_kpi"][0] = {
            "basis": "modeled_recalculated_canonical_factor_0.90",
            "runs_per_week": sum(
                as_number(row.get("runs_per_week")) or 0.0 for row in modeled
            ),
            "fini_changes_per_week": sum(
                as_number(row.get("j_ch_contribution")) or 0.0 for row in modeled
            ),
            "pck_changes_per_week": None,
            "same_pck_changes_per_week": None,
            "status": "canonical_development",
        }


def extract_primary_workbook(
    path: Path, snapshot_expectations: Mapping[str, Any] | None = None
) -> tuple[dict[str, list[dict[str, Any]]], dict[str, Any], list[dict[str, Any]]]:
    """Extract all optimizer-relevant primary-workbook tables and validation evidence.

    Args:
        path: Primary production workbook path.
        snapshot_expectations: Optional reviewed figures of a frozen snapshot used as
            a regression guard. Supported keys: ``assigned``, ``excluded``,
            ``modeled``, ``modeled_groups``, ``modeled_demand``,
            ``excluded_demand``, ``modeled_blocks``, ``historical_groups`` and
            ``excluded_materials`` (list of material numbers). Each present key
            that differs from the workbook raises a ``SNAPSHOT_*`` error issue;
            absent keys are not checked.

    Returns:
        Tuple of emitted tables, workbook metadata, and validation issues.
    """

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", category=UserWarning)
        workbook = load_workbook(path, data_only=True, keep_links=False)
    required = {
        name: find_table(workbook, name)
        for name in (
            "Site",
            "Table11",
            "FINI_grouping_for_wheel",
            "Table15",
            "PCK_per_filling_line___{plant}_specific",
            "Change_over_estimation",
            "FINI_new_recipe_combination",
            "FINI_Cycle_stock_at_DC",
            "FINI_Cycle_stock_at_DC_without_pallet_conversion",
            "FINI_Planning_calendar_at_P_site",
            "FINI_Planning_calendar_at_P_site___no_calendar",
            "APO_SEFI_lot_size_per_recipe_combination",
            "Newrecipe",
            "Production_versions___List",
        )
    }
    site = as_text(required["Site"].records[0].get("Site")) or "UNKNOWN"
    fini_rows, legacy_coverage = _canonical_fini(required["FINI_grouping_for_wheel"], site)
    fini_line, pck_line = _line_tables(
        fini_rows, required["PCK_per_filling_line___{plant}_specific"], site
    )
    production_versions = _production_versions(required["Production_versions___List"], site)
    fini_rows = classify_fini_rows(
        fini_rows,
        production_versions,
        require_resolved_pallet=False,
    )

    # Build the character map first, then recompute groups with the characters attached.
    preliminary_character_map = {
        (as_text(record.get("SEFI")) or "", as_text(record.get("Subgroup")) or ""): as_text(
            record.get("SAP character")
        )
        for record in required["Newrecipe"].records
    }
    baseline_groups = _group_metrics(fini_rows, production_versions, preliminary_character_map)
    character_catalog, character_compatibility, character_capacity, character_map = _character_tables(
        required["Newrecipe"], workbook["Compatibility matrix"], baseline_groups, site
    )
    if character_map != preliminary_character_map:
        baseline_groups = _group_metrics(fini_rows, production_versions, character_map)

    tables: dict[str, list[dict[str, Any]]] = {
        "site_parameters": [
            {
                "plant": site,
                "forecast_horizon_months": HORIZON_MONTHS,
                "productive_weeks": PRODUCTIVE_WEEKS,
                "demand_days": DEMAND_DAYS,
                "canonical_factor": CANONICAL_FACTOR,
                "forecast_source": "PRIMARY",
                "status": "development_baseline",
            }
        ],
        "fini_master": fini_rows,
        "fini_line_eligibility": fini_line,
        "pck_line_eligibility": pck_line,
        "production_versions": production_versions,
        "baseline_assignments": _baseline_assignments(fini_rows),
        "package_volume_catalog": _package_volume_catalog(fini_rows),
        "baseline_group_metrics": baseline_groups,
        "legacy_fini_roc_by_pv": legacy_coverage,
        "sap_character_catalog": character_catalog,
        "sap_character_compatibility": character_compatibility,
        "sap_character_capacity": character_capacity,
        "sales_forecast_summary": _simple_records(
            required["Table11"],
            {"Measure": "measure", "SEFIs": "sefi_count", "FINIs": "fini_count", "Sales forecast": "forecast_litres"},
        ),
        "planning_calendar_inform": _simple_records(
            required["FINI_Planning_calendar_at_P_site"],
            {"Material": "material", "Plant": "plant", "Lot size": "lot_size_procedure", "Planning calendar": "planning_calendar", "Lot size Profile": "lot_size_profile"},
        ),
        "planning_calendar_no_calendar": _simple_records(
            required["FINI_Planning_calendar_at_P_site___no_calendar"],
            {"FINI": "material", "P-site": "plant", "Lot size": "lot_size_procedure", "Lot size Profile": "lot_size_profile", "Planning calendar": "planning_calendar"},
        ),
        "fini_cycle_stock_at_dc": _simple_records(
            required["FINI_Cycle_stock_at_DC"],
            {"Plnt": "plant", "Material": "material", "Lot size": "lot_size", "Min lot size at DC": "min_lot_size_at_dc", "Rounding value": "rounding_value", "Defined in wheel": "defined_in_wheel"},
        ),
        "fini_cycle_stock_without_pallet_conversion": _simple_records(
            required["FINI_Cycle_stock_at_DC_without_pallet_conversion"],
            {"Plnt": "plant", "Material": "material", "Lot size": "lot_size", "Defined in wheel": "defined_in_wheel"},
        ),
        "apo_sefi_lot_size": _simple_records(
            required["APO_SEFI_lot_size_per_recipe_combination"],
            {"SEFI": "sefi", "Recipe combination": "recipe_combination", "Lot size KG": "lot_size_kg"},
        ),
        "fini_new_recipe_combination": _simple_records(
            required["FINI_new_recipe_combination"],
            {"Plant": "plant", "Material": "material", "Recipe combination": "recipe_combination"},
        ),
        "baseline_groups_5_june": _simple_records(
            required["Table15"],
            {"Material": "material", "Subgroup": "subgroup", "RoC selected PV": "coverage_selected_pv", "Match?": "match"},
        ),
    }
    tables["legacy_kpi"] = _legacy_kpis(baseline_groups, required["Change_over_estimation"])

    modeled = [row for row in fini_rows if row["model_status"] == "modeled"]
    excluded = [row for row in fini_rows if row["model_status"] == "excluded"]
    assigned = [row for row in fini_rows if row["baseline_assigned"]]
    modeled_groups = [row for row in baseline_groups if row["population"] == "modeled"]
    issues: list[dict[str, Any]] = []
    expected_counts = {
        "fini_master": 1021,
        "baseline_assignments": 803,
        "planning_calendar_inform": 177,
        "planning_calendar_no_calendar": 844,
        "fini_cycle_stock_without_pallet_conversion": 175,
    }
    for table_name, expected in expected_counts.items():
        observed = len(tables[table_name])
        if observed != expected:
            issues.append(
                _issue(
                    f"COUNT_{table_name.upper()}",
                    "error",
                    table_name,
                    None,
                    observed,
                    expected,
                    "block acceptance until source range is reconciled",
                )
            )
    observed_summary = {
        "assigned": len(assigned),
        "excluded": len(excluded),
        "modeled": len(modeled),
        "modeled_groups": len(modeled_groups),
        "modeled_demand": round(sum(as_number(row["forecast_litres_12m"]) or 0.0 for row in modeled), 6),
        "excluded_demand": round(sum(as_number(row["forecast_litres_12m"]) or 0.0 for row in excluded), 6),
        "modeled_blocks": len({(row["plant"], row["sefi"]) for row in modeled}),
    }
    expectations = snapshot_expectations or {}
    historical_groups = sum(row["population"] == "historical" for row in baseline_groups)
    expected_historical = expectations.get("historical_groups")
    if expected_historical is not None and historical_groups != expected_historical:
        issues.append(
            _issue(
                "SNAPSHOT_HISTORICAL_GROUPS",
                "error",
                "snapshot",
                "historical_groups",
                historical_groups,
                expected_historical,
                "block acceptance until historical assignment keys are reconciled",
            )
        )
    for field, observed in observed_summary.items():
        expected = expectations.get(field)
        if expected is not None and observed != expected:
            issues.append(
                _issue(
                    f"SNAPSHOT_{field.upper()}",
                    "error",
                    "snapshot",
                    field,
                    observed,
                    expected,
                    "block acceptance until the frozen snapshot is reconciled",
                )
            )
    invalid_groups = [
        row["baseline_group_key"]
        for row in modeled_groups
        if row["line_status"] != "feasible"
    ]
    if invalid_groups:
        issues.append(
            _issue(
                "MODELED_GROUP_COMMON_LINE",
                "error",
                "baseline_group",
                "|".join(invalid_groups[:10]),
                len(invalid_groups),
                0,
                "exclude invalid members or correct line evidence before optimization",
            )
        )
    pv_keys = collections.Counter(
        (str(row["sefi"]), str(row["production_version"])) for row in production_versions
    )
    unresolved_pv = [
        row["material"]
        for row in modeled
        if pv_keys[(str(row["sefi"]), str(row["fixed_pv"]))] != 1
    ]
    if unresolved_pv:
        issues.append(
            _issue(
                "MODELED_FIXED_PV_UNIQUE",
                "error",
                "fini",
                "|".join(unresolved_pv[:10]),
                len(unresolved_pv),
                0,
                "block affected configurations; never substitute PV1",
            )
        )
    observed_exclusions = {str(row["material"]) for row in excluded}
    expected_exclusions = (
        {str(material) for material in expectations["excluded_materials"]}
        if "excluded_materials" in expectations
        else None
    )
    if expected_exclusions is not None and observed_exclusions != expected_exclusions:
        issues.append(
            _issue(
                "SNAPSHOT_EXCLUSION_IDENTITIES",
                "error",
                "snapshot",
                "excluded_materials",
                "|".join(sorted(observed_exclusions)),
                "|".join(sorted(expected_exclusions)),
                "block acceptance until exclusion identities are reviewed",
            )
        )
    duplicate_materials = [
        key for key, count in collections.Counter((row["plant"], row["material"]) for row in fini_rows).items() if count != 1
    ]
    if duplicate_materials:
        issues.append(
            _issue(
                "FINI_KEY_UNIQUE",
                "error",
                "fini",
                str(duplicate_materials[:10]),
                len(duplicate_materials),
                0,
                "block duplicate FINI/plant identities",
            )
        )
    malformed_modeled = [
        row["material"]
        for row in modeled
        if row["production_frequency_code"] not in MODELED_FREQUENCY_CODES
        or not (as_number(row["forecast_litres_12m"]) or 0) > 0
        or not (as_number(row["package_volume"]) or 0) > 0
        or not row["pck_code"]
    ]
    if malformed_modeled:
        issues.append(
            _issue(
                "MODELED_REQUIRED_FIELDS",
                "error",
                "fini",
                "|".join(map(str, malformed_modeled[:10])),
                len(malformed_modeled),
                0,
                "block affected FINIs until demand/frequency/package evidence is valid",
            )
        )
    invalid_pv_lots = [
        f"{row['sefi']}/{row['production_version']}"
        for row in production_versions
        if not (as_number(row["lot_size_litres"]) or 0) > 0
    ]
    if invalid_pv_lots:
        issues.append(
            _issue(
                "PV_LOT_POSITIVE",
                "error",
                "production_version",
                "|".join(invalid_pv_lots[:10]),
                len(invalid_pv_lots),
                0,
                "block invalid production versions",
            )
        )
    inconsistent_groups = [
        row["baseline_group_key"]
        for row in modeled_groups
        if not row["fixed_pv_consistent"]
    ]
    if inconsistent_groups:
        issues.append(
            _issue(
                "BASELINE_GROUP_FIXED_PV_CONSISTENT",
                "error",
                "baseline_group",
                "|".join(map(str, inconsistent_groups[:10])),
                len(inconsistent_groups),
                0,
                "block inconsistent baseline groups",
            )
        )
    missing_pallet = sum(1 for row in modeled if not as_number(row["pallet_litres_primary"]))
    if missing_pallet:
        issues.append(
            _issue(
                "MODELED_PALLET_ENRICHMENT_REQUIRED",
                "warning",
                "fini",
                None,
                missing_pallet,
                0,
                "resolve from one unique positive enrichment candidate",
            )
        )
    metadata = {
        "workbook": path.name,
        "sheets": workbook.sheetnames,
        "tables": {
            name: {
                "sheet": table.sheet,
                "range": table.reference,
                "record_count": len(table.records),
            }
            for name, table in required.items()
        },
        "snapshot": observed_summary,
        "canonical_factor": CANONICAL_FACTOR,
        "legacy_pv1_factor": 0.85,
    }
    return tables, metadata, issues
