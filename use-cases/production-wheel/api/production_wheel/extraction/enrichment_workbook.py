"""Read enrichment-workbook data without loading its unsupported pivot cache."""

from __future__ import annotations

import collections
import warnings
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from production_wheel.extraction.common import as_number, as_text, xlsx_table_catalog


def extract_pallet_candidates(
    path: Path, fini_keys: set[str]
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Extract pallet candidates matching the primary-workbook FINI/plant population.

    Args:
        path: Enrichment input workbook.
        fini_keys: FINI/plant keys to retain.

    Returns:
        Matching candidate rows and compact workbook metadata.
    """

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", category=UserWarning)
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    if "FINI_pallets" not in workbook.sheetnames:
        raise KeyError("Enrichment workbook does not contain sheet 'FINI_pallets'")
    worksheet = workbook["FINI_pallets"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [as_text(value) for value in next(rows)]
    required = {
        "Material",
        "Layer qty",
        "Pallet qty",
        "Volume",
        "Pallet in litres",
        "Producing plant",
        "FINI/P-site",
    }
    missing = required - {header for header in headers if header}
    if missing:
        raise KeyError(f"Enrichment FINI_pallets is missing headers: {sorted(missing)}")
    candidates: list[dict[str, Any]] = []
    for source_row, values in enumerate(rows, start=2):
        record = dict(zip(headers, values, strict=True))
        fini_key = as_text(record.get("FINI/P-site"))
        if fini_key not in fini_keys:
            continue
        source_pallet_litres = as_number(record.get("Pallet in litres"))
        pallet_qty = as_number(record.get("Pallet qty"))
        package_volume = as_number(record.get("Volume"))
        calculated_pallet_litres = (
            pallet_qty * package_volume
            if pallet_qty is not None and package_volume is not None
            else None
        )
        candidates.append(
            {
                "fini_plant_key": fini_key,
                "plant": as_text(record.get("Producing plant")),
                "material": as_text(record.get("Material")),
                "layer_qty": as_number(record.get("Layer qty")),
                "pallet_qty": pallet_qty,
                "package_volume": package_volume,
                "pallet_litres_source": source_pallet_litres,
                "pallet_litres_calculated": calculated_pallet_litres,
                "pallet_litres": calculated_pallet_litres,
                "source_calculation_matches": int(
                    source_pallet_litres is not None
                    and calculated_pallet_litres is not None
                    and abs(source_pallet_litres - calculated_pallet_litres) <= 1e-6
                ),
                "source_sheet": worksheet.title,
                "source_row": source_row,
            }
        )
    horizon = None
    if "Horizon selector" in workbook.sheetnames:
        horizon = as_text(workbook["Horizon selector"]["B2"].value)
    metadata = {
        "workbook": path.name,
        "tables": xlsx_table_catalog(path),
        "sheets": [
            {
                "name": sheet.title,
                "dimension": sheet.calculate_dimension(),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
            }
            for sheet in workbook.worksheets
        ],
        "forecast_horizon_selected": horizon,
        "pallet_candidate_rows_matched": len(candidates),
    }
    workbook.close()
    return candidates, metadata


def resolve_pallets(
    fini_rows: Iterable[Mapping[str, Any]],
    candidates: Iterable[Mapping[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    """Resolve pallet litres with primary-first, unique-positive-enrichment semantics.

    Args:
        fini_rows: Canonical FINI master records.
        candidates: Matching enrichment pallet candidate records.

    Returns:
        Updated FINI records, one resolution record per FINI, and issues.
    """

    by_key: dict[str, list[Mapping[str, Any]]] = collections.defaultdict(list)
    for candidate in candidates:
        by_key[str(candidate["fini_plant_key"])].append(candidate)
    updated: list[dict[str, Any]] = []
    resolutions: list[dict[str, Any]] = []
    issues: list[dict[str, Any]] = []
    nonblank_discrepancies: list[str] = []
    for source in fini_rows:
        row = dict(source)
        key = str(row["fini_plant_key"])
        primary_value = as_number(row.get("pallet_litres_primary"))
        raw_candidates = by_key.get(key, [])
        positive_values = sorted(
            {
                round(value, 9)
                for candidate in raw_candidates
                if (value := as_number(candidate.get("pallet_litres"))) is not None and value > 0
            }
        )
        conflict = len(positive_values) > 1
        if primary_value is not None:
            resolved = primary_value
            source_name = "PRIMARY"
            status = "resolved_primary" if primary_value > 0 else "invalid_primary_nonpositive"
            if primary_value > 0 and positive_values and not any(abs(value - primary_value) <= 1e-6 for value in positive_values):
                nonblank_discrepancies.append(key)
        elif len(positive_values) == 1:
            resolved = positive_values[0]
            source_name = "ENRICHMENT_FINI_pallets"
            status = "resolved_enrichment_unique"
        elif conflict:
            resolved = None
            source_name = None
            status = "conflicting_enrichment_candidates"
        else:
            resolved = None
            source_name = None
            status = "missing_candidate"
        row["pallet_litres_resolved"] = resolved
        row["pallet_resolution_source"] = source_name
        row["pallet_candidate_count"] = len(positive_values)
        row["pallet_conflict"] = int(conflict)
        updated.append(row)
        resolutions.append(
            {
                "plant": row["plant"],
                "material": row["material"],
                "fini_plant_key": key,
                "model_status": row["model_status"],
                "primary_pallet_litres": primary_value,
                "enrichment_positive_candidates": "|".join(f"{value:g}" for value in positive_values),
                "candidate_count": len(positive_values),
                "resolved_pallet_litres": resolved,
                "resolution_source": source_name,
                "resolution_status": status,
                "conflict": int(conflict),
            }
        )
    if nonblank_discrepancies:
        issues.append(
            {
                "rule_id": "PALLET_NONBLANK_SOURCE_DISCREPANCY",
                "severity": "warning",
                "entity": "fini",
                "entity_key": "|".join(nonblank_discrepancies[:10]),
                "observed": len(nonblank_discrepancies),
                "expected": 0,
                "action": "retain primary value and expose enrichment discrepancy for review",
            }
        )
    return updated, resolutions, issues
