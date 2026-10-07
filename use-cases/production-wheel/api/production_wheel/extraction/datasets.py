"""Deterministic in-memory multi-workbook dataset import for application storage.

Example::

    dataset = extract_workbooks(Path('wheel.xlsx'), settings={'demand_days': 250})
    inputs = canonical_inputs_from_tables(dataset['tables'])

No intermediate CSV, local output directory, network call or AI interpretation is
used. The historical primary-workbook command-line regression extractor remains separate.
"""

from __future__ import annotations

import collections
import warnings
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from production_wheel.extraction.common import as_text, find_table, sha256_file
from production_wheel.extraction.dataset_tables import (
    apply_line_evidence,
    detect_profile,
    is_supporting_table,
    parse_string_lines,
    supporting_header_map,
    supporting_rows,
)
from production_wheel.extraction.enrichment_workbook import extract_pallet_candidates, resolve_pallets
from production_wheel.extraction.primary_workbook import (
    _baseline_assignments,
    _canonical_fini,
    _package_volume_catalog,
    _production_versions,
    classify_fini_admission,
    normalize_production_frequency,
)

from production_wheel.extraction.profile_inputs import enrich_profile_inputs

PARSER_VERSION = "production-wheel-dataset-v2"
DEFAULT_SETTINGS = {
    "modeled_frequencies": ["01W", "02W"],
    "demand_days": 250,
    "productive_weeks": 50,
    "canonical_factor": 0.9,
}


def _settings(overrides: dict | None) -> tuple[dict, dict]:
    """Validate import parameter overrides and return values plus origin labels."""
    supplied = overrides or {}
    unknown = set(supplied) - set(DEFAULT_SETTINGS)
    if unknown:
        raise ValueError(f"Unknown dataset settings: {sorted(unknown)}")
    result = dict(DEFAULT_SETTINGS) | supplied
    for key in ("demand_days", "productive_weeks"):
        value = result[key]
        if (
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or value <= 0
            or int(value) != value
        ):
            raise ValueError(f"{key} must be a positive integer")
        result[key] = int(value)
    factor = result["canonical_factor"]
    if (
        isinstance(factor, bool)
        or not isinstance(factor, (int, float))
        or not 0 < factor <= 1
    ):
        raise ValueError("canonical_factor must be greater than zero and at most one")
    raw_codes = result["modeled_frequencies"]
    if not isinstance(raw_codes, (list, tuple)) or not raw_codes:
        raise ValueError("modeled_frequencies must be a nonempty list")
    codes = [normalize_production_frequency(value) for value in raw_codes]
    if any(code is None for code in codes):
        raise ValueError("modeled_frequencies contains an unrecognized weekly code")
    result["modeled_frequencies"] = sorted(set(codes))
    return result, {key: "user" if key in supplied else "default" for key in result}


def _plant(catalog: dict, main: Any) -> str:
    """Resolve a single plant from explicit Site or main-record key evidence."""
    evidence = set()
    if "Site" in catalog:
        evidence.update(as_text(row.get("Site")) for row in catalog["Site"].records)
    for row in main.records:
        evidence.add(as_text(row.get("Plant")) or as_text(row.get("P-site")))
        for field in ("FINI/P-site", "SEFI/Plant"):
            key = as_text(row.get(field))
            if key and "/" in key:
                evidence.add(key.rsplit("/", 1)[1].strip())
    evidence -= {None, ""}
    if len(evidence) != 1:
        raise ValueError(
            f"Workbook must identify exactly one plant; found {sorted(evidence)}"
        )
    return next(iter(evidence))


def _line_bridges(
    rows: list[dict], package_rows: list[dict], plant: str
) -> tuple[list, list]:
    """Build plant-scoped eligibility over only source-observed line identifiers."""
    universe = {
        line
        for row in rows
        for line in (row.get("eligible_lines") or "").split("|")
        if line
    }
    for source in package_rows:
        universe.update(parse_string_lines(source.get("value")) or ())
    fini_bridge = []
    for row in rows:
        eligible = set((row.get("eligible_lines") or "").split("|"))
        for line in sorted(universe):
            fini_bridge.append(
                {
                    "plant": plant,
                    "material": row["material"],
                    "pck_code": row["pck_code"],
                    "filling_line": line,
                    "eligible": line in eligible
                    if row["line_data_status"] == "known"
                    else None,
                    "line_data_status": row["line_data_status"],
                    "source_pattern": row["filling_line_raw"],
                    "source_sheet": row["source_sheet"],
                    "source_row": row["source_row"],
                }
            )
    package_bridge = []
    for row in package_rows:
        eligible = parse_string_lines(row.get("value"))
        for line in sorted(universe):
            package_bridge.append(
                {
                    "plant": plant,
                    "pck_code": as_text(row.get("name")),
                    "filling_line": line,
                    "eligible": line in eligible if eligible is not None else None,
                    "source_pattern": as_text(row.get("value")),
                    "source_sheet": row["source_sheet"],
                    "source_row": row["source_row"],
                }
            )
    return fini_bridge, package_bridge


def extract_workbooks(
    primary_path: Path,
    enrichment_path: Path | None = None,
    settings: dict | None = None,
) -> dict:
    """Extract canonical inputs from a primary workbook and optional pallet source.

    Args:
        primary_path: Workbook with common FINI and production-version tables.
        enrichment_path: Optional pallet enrichment workbook.
        settings: Frequency, demand-day, productive-week and PV-factor overrides.

    Returns:
        Dictionary with metadata, canonical table row lists and validation issues.
        Source row and sheet lineage are preserved. Historical sheets are optional.
    """
    parameters, origins = _settings(settings)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = load_workbook(primary_path, data_only=True, keep_links=False)
    try:
        catalog = {
            name: find_table(workbook, name)
            for sheet in workbook
            for name in sheet.tables
        }
        for required in ("FINI_grouping_for_wheel", "Production_versions___List"):
            if required not in catalog:
                raise ValueError(f"Required workbook table is missing: {required}")
        main = catalog["FINI_grouping_for_wheel"]
        plant = _plant(catalog, main)
        tables = supporting_rows(catalog)
        rows, legacy = _canonical_fini(main, plant)
        issues = apply_line_evidence(rows, tables["line_corrections"])
        versions = _production_versions(catalog["Production_versions___List"], plant)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            formula_book = load_workbook(primary_path, data_only=False, keep_links=False)
        try:
            formula_main = find_table(formula_book, "FINI_grouping_for_wheel")
            issues.extend(enrich_profile_inputs(rows, main, formula_main, versions))
        finally:
            formula_book.close()
        horizons = {row['source_horizon_days'] for row in rows if row.get('source_horizon_days') is not None}
        if len(horizons) == 1 and 'demand_days' not in (settings or {}):
            # Totals describe this source period, not an assumed twelve months.
            parameters['demand_days'] = next(iter(horizons))
            origins['demand_days'] = 'source_total_divided_by_daily_demand'
            if 'productive_weeks' not in (settings or {}):
                parameters['productive_weeks'] = parameters['demand_days'] / 5
                origins['productive_weeks'] = 'source_horizon_at_five_demand_days_per_week'
        for version in versions:
            version["canonical_factor"] = parameters["canonical_factor"]
        for record in legacy:
            record["canonical_factor"] = parameters["canonical_factor"]
        candidates = []
        if enrichment_path is not None:
            candidates, _ = extract_pallet_candidates(
                enrichment_path, {row["fini_plant_key"] for row in rows}
            )
        rows, resolutions, pallet_issues = resolve_pallets(rows, candidates)
        issues.extend(pallet_issues)
        for row in rows:
            demand = row["forecast_litres_12m"]
            row["avg_daily_demand_source"] = row["avg_daily_demand"]
            row["avg_daily_demand"] = (
                demand / parameters["demand_days"] if demand is not None else None
            )
            if row["pallet_resolution_source"] == "PRIMARY":
                row["pallet_resolution_source"] = "primary_workbook"
        rows = [
            classify_fini_admission(
                row, versions, modeled_frequencies=parameters["modeled_frequencies"]
            )
            for row in rows
        ]
        by_key = {row["fini_plant_key"]: row for row in rows}
        for resolution in resolutions:
            final_row = by_key[resolution["fini_plant_key"]]
            resolution["model_status"] = final_row["model_status"]
            resolution["resolution_source"] = final_row["pallet_resolution_source"]
        fini_bridge, package_bridge = _line_bridges(
            rows, tables.pop("pck_line_source"), plant
        )
        tables.update(
            fini_master=rows,
            production_versions=versions,
            legacy_fini_roc_by_pv=legacy,
            baseline_assignments=_baseline_assignments(rows),
            package_volume_catalog=_package_volume_catalog(rows),
            fini_line_eligibility=fini_bridge,
            pck_line_eligibility=package_bridge,
            enrichment_pallet_candidates=candidates,
            pallet_resolution=resolutions,
            site_parameters=[{"plant": plant, **parameters}],
        )
        duplicates = [
            key
            for key, count in collections.Counter(
                row["fini_plant_key"] for row in rows
            ).items()
            if count > 1
        ]
        if duplicates:
            issues.append(
                {
                    "rule_id": "FINI_KEY_UNIQUE",
                    "severity": "error",
                    "entity_key": "|".join(duplicates),
                    "action": "Resolve duplicate FINI/plant identities",
                }
            )
        for row in rows:
            if row["model_status"] == "excluded":
                issues.append(
                    {
                        "rule_id": "FINI_DATA_QUALITY",
                        "severity": "warning",
                        "entity_key": row["fini_plant_key"],
                        "reason": row["exclusion_reason"],
                        "action": "Review source evidence before admitting this FINI",
                    }
                )
        hashes = {"primary": sha256_file(primary_path)}
        if enrichment_path is not None:
            hashes["enrichment"] = sha256_file(enrichment_path)
        metadata = {
            "plant": plant,
            "format_id": detect_profile(catalog),
            "settings": parameters,
            "parameter_sources": origins,
            "counts": {name: len(records) for name, records in tables.items()},
            "input_hashes": hashes,
            "parser_version": PARSER_VERSION,
            "source_tables": {
                name: {
                    "sheet": table.sheet,
                    "range": table.reference,
                    **(
                        {"header_mapping": supporting_header_map(table)}
                        if is_supporting_table(name)
                        else {}
                    ),
                }
                for name, table in catalog.items()
            },
        }
        return {"metadata": metadata, "tables": tables, "issues": issues}
    finally:
        workbook.close()
