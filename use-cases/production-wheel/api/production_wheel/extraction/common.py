"""Shared, deterministic workbook-extraction utilities."""

from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import posixpath
import re
import struct
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook


@dataclass(frozen=True, slots=True)
class TableData:
    """Represent one Excel table and its extracted records.

    Attributes:
        name: Excel table name.
        sheet: Worksheet containing the table.
        reference: Excel range reference.
        records: Header-keyed data rows with empty rows removed.
    """

    name: str
    sheet: str
    reference: str
    records: tuple[dict[str, Any], ...]


def utc_now() -> str:
    """Return the current UTC timestamp in a stable manifest representation."""

    return datetime.now(UTC).isoformat(timespec="seconds").replace("+00:00", "Z")


def sha256_file(path: Path) -> str:
    """Return the SHA-256 digest of a local file.

    Args:
        path: File to hash.

    Returns:
        Lowercase hexadecimal SHA-256 digest.
    """

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def as_text(value: Any) -> str | None:
    """Normalize an Excel identifier or text value without adding decimal suffixes."""

    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_number(value: Any) -> float | None:
    """Return a finite numeric value or ``None`` for blanks/errors/non-numbers."""

    if isinstance(value, bool) or value in (None, ""):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number == number and abs(number) != float("inf") else None


def parse_lines(value: Any) -> tuple[int, ...] | None:
    """Parse a six-position filling-line pattern into sorted line identifiers.

    Args:
        value: Raw workbook cell such as ``"1  2  -  4"``.

    Returns:
        Sorted line tuple, or ``None`` when the source cell is blank.
    """

    text = as_text(value)
    if text is None:
        return None
    return tuple(sorted({int(token) for token in text.split() if token.isdigit()}))


# Some source workbooks embed a plant or site code in their Excel table names
# (for example ``PCK_per_filling_line___<plant>_specific``). Table contracts use
# the ``{plant}`` placeholder for that code so no plant is hard-coded; the token
# allows one optional underscore-separated suffix such as ``AB12_CD``.
PLANT_TOKEN = r"[A-Za-z0-9]+(?:_[A-Za-z0-9]+)?"


def table_name_matches(contract: str, name: str) -> bool:
    """Return whether an Excel table name satisfies a table-name contract.

    Args:
        contract: Exact table name, or a name containing ``{plant}`` placeholders.
        name: Actual Excel table name found in a workbook.

    Returns:
        True when the names are equal or the placeholders match a plant code.
    """

    if "{plant}" not in contract:
        return contract == name
    pattern = re.escape(contract).replace(re.escape("{plant}"), PLANT_TOKEN)
    return re.fullmatch(pattern, name) is not None


def find_table(workbook: Workbook, table_name: str) -> TableData:
    """Extract an Excel table by name using its dynamic range.

    Args:
        workbook: Loaded OpenPyXL workbook.
        table_name: Exact Excel table name, or a contract containing ``{plant}``
            placeholders (see ``table_name_matches``); the first match is used.

    Returns:
        Table metadata and non-empty records, named after the actual Excel table.

    Raises:
        KeyError: If the table is absent or contains no header row.
    """

    for worksheet in workbook.worksheets:
        matches = [name for name in worksheet.tables if table_name_matches(table_name, name)]
        if not matches:
            continue
        actual_name = matches[0]
        table = worksheet.tables[actual_name]
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        headers = [as_text(worksheet.cell(min_row, column).value) for column in range(min_col, max_col + 1)]
        if any(header is None for header in headers):
            raise KeyError(f"table {actual_name!r} contains a blank header")
        records: list[dict[str, Any]] = []
        for row in range(min_row + 1, max_row + 1):
            values = [worksheet.cell(row, column).value for column in range(min_col, max_col + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            record = dict(zip((str(header) for header in headers), values, strict=True))
            record["_source_row"] = row
            records.append(record)
        return TableData(actual_name, worksheet.title, table.ref, tuple(records))
    raise KeyError(f"Excel table not found: {table_name}")


def write_csv(path: Path, rows: Iterable[Mapping[str, Any]], fields: Sequence[str] | None = None) -> int:
    """Write dictionaries to a deterministic UTF-8 CSV and return its row count."""

    materialized = list(rows)
    if fields is None:
        ordered: list[str] = []
        seen: set[str] = set()
        for row in materialized:
            for key in row:
                if key not in seen:
                    ordered.append(key)
                    seen.add(key)
        fields = ordered
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fields), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(materialized)
    return len(materialized)


def write_json(path: Path, value: Any) -> None:
    """Write a JSON value with stable human-readable formatting."""

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")


def extract_power_query_m(path: Path) -> str | None:
    """Extract the embedded Power Query Section1 M source from an XLSX file."""

    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("customXml/item"):
                continue
            try:
                xml_text = archive.read(name).decode("utf-16", errors="ignore")
            except (KeyError, UnicodeError):
                continue
            match = re.search(r"DataMashup[^>]*>(.*?)</DataMashup>", xml_text, re.DOTALL)
            if not match:
                continue
            raw = base64.b64decode(re.sub(r"\s", "", match.group(1)))
            if len(raw) < 8:
                continue
            inner_length = struct.unpack("<I", raw[4:8])[0]
            with zipfile.ZipFile(io.BytesIO(raw[8 : 8 + inner_length])) as inner:
                for inner_name in inner.namelist():
                    if inner_name.lower().endswith(".m"):
                        return inner.read(inner_name).decode("utf-8", errors="ignore")
    return None


def xlsx_table_catalog(path: Path) -> list[dict[str, Any]]:
    """Read Excel table names, sheets, ranges, and row counts from OOXML parts.

    This avoids a normal-mode OpenPyXL load for workbooks whose pivot-cache
    extensions are unsafe or expensive to deserialize.

    Args:
        path: XLSX workbook path.

    Returns:
        One sorted record per Excel table.
    """

    relationship_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(path) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        workbook_rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        workbook_targets = {
            relation.attrib["Id"]: posixpath.normpath(posixpath.join("xl", relation.attrib["Target"]))
            for relation in workbook_rels.findall(f"{{{package_ns}}}Relationship")
        }
        sheet_paths = {
            sheet.attrib["name"]: workbook_targets[sheet.attrib[f"{{{relationship_ns}}}id"]]
            for sheet in workbook_root.findall(f".//{{{spreadsheet_ns}}}sheet")
        }
        result: list[dict[str, Any]] = []
        for sheet_name, sheet_path in sheet_paths.items():
            sheet_file = posixpath.basename(sheet_path)
            relations_path = posixpath.join(posixpath.dirname(sheet_path), "_rels", f"{sheet_file}.rels")
            if relations_path not in archive.namelist():
                continue
            relation_root = ET.fromstring(archive.read(relations_path))
            for relation in relation_root.findall(f"{{{package_ns}}}Relationship"):
                if not relation.attrib.get("Type", "").endswith("/table"):
                    continue
                table_path = posixpath.normpath(
                    posixpath.join(posixpath.dirname(sheet_path), relation.attrib["Target"])
                )
                table_root = ET.fromstring(archive.read(table_path))
                reference = table_root.attrib["ref"]
                _min_col, min_row, _max_col, max_row = range_boundaries(reference)
                result.append(
                    {
                        "name": table_root.attrib.get("displayName") or table_root.attrib.get("name"),
                        "sheet": sheet_name,
                        "range": reference,
                        "record_count": max(max_row - min_row, 0),
                        "ooxml_part": table_path,
                    }
                )
    return sorted(result, key=lambda row: (str(row["sheet"]), str(row["name"])))


_QUERY_RE = re.compile(r'^shared\s+(#"[^"]+"|[A-Za-z_][\w]*)\s*=', re.MULTILINE)
_LITERAL_SOURCE_RE = re.compile(r'(File\.Contents|Web\.Contents)\("([^"]+)"\)')
_PARAMETER_SOURCE_RE = re.compile(r"Files\[([^\]]+)\]")


def query_catalog(m_source: str, workbook_label: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Return query and external-dependency catalogs from Power Query M source.

    Args:
        m_source: Complete Section1 M source.
        workbook_label: Stable source label such as ``PRIMARY`` or ``ENRICHMENT``.

    Returns:
        Pair containing query rows and dependency rows.
    """

    matches = list(_QUERY_RE.finditer(m_source))
    queries: list[dict[str, Any]] = []
    dependencies: list[dict[str, Any]] = []
    for index, match in enumerate(matches):
        raw_name = match.group(1)
        name = raw_name[2:-1] if raw_name.startswith('#"') else raw_name
        end = matches[index + 1].start() if index + 1 < len(matches) else len(m_source)
        body = m_source[match.start() : end]
        activated = workbook_label == "ENRICHMENT" and name == "FINI pallets"
        diagnostic_names = {
            "FINI litres per pallets",
            "FINI pallets in pieces for products with forecast",
            "FINI pallets layer in pieces for products with forecast (2)",
            "SKU volume",
            "Producing Plant per DC",
            "Network___Primary_DC",
            "Files",
            "FirstMonth",
            "ForecastHorizon",
            "Weeks to consider",
            "Production versions - List",
            "Production_versions___List",
        }
        normalized = name.casefold()
        if activated:
            status = "activated"
            rationale = "Supplies governed pallet fallback candidates."
        elif name in diagnostic_names or "forecast" in normalized or "horizon" in normalized:
            status = "diagnostic_only"
            rationale = "Retained for source comparison; the primary workbook remains the canonical pilot horizon/PV source."
        else:
            status = "ignored"
            rationale = "Extracted for rule lineage and future review; not activated by the prototype."
        queries.append(
            {
                "workbook": workbook_label,
                "query": name,
                "implementation_status": status,
                "rationale": rationale,
                "source_sha256": hashlib.sha256(body.encode("utf-8")).hexdigest(),
            }
        )
        for kind, source in _LITERAL_SOURCE_RE.findall(body):
            dependencies.append(
                {
                    "workbook": workbook_label,
                    "query": name,
                    "kind": "url" if kind == "Web.Contents" else "file",
                    "source": source,
                    "source_form": "literal",
                }
            )
        for parameter in _PARAMETER_SOURCE_RE.findall(body):
            dependencies.append(
                {
                    "workbook": workbook_label,
                    "query": name,
                    "kind": "parameterized_file",
                    "source": parameter,
                    "source_form": "Files parameter",
                }
            )
    unique_dependencies = list(
        {
            (row["workbook"], row["query"], row["kind"], row["source"]): row
            for row in dependencies
        }.values()
    )
    return queries, unique_dependencies


def infer_field_dictionary(table_name: str, rows: Sequence[Mapping[str, Any]], source: str) -> list[dict[str, Any]]:
    """Infer a compact field dictionary for one emitted table.

    Args:
        table_name: Output table identifier.
        rows: Emitted records.
        source: Human-readable workbook/table lineage.

    Returns:
        One dictionary row per field.
    """

    fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    result: list[dict[str, Any]] = []
    for field in fields:
        values = [row.get(field) for row in rows if row.get(field) not in (None, "")]
        if values and all(isinstance(value, bool) for value in values):
            value_type = "boolean"
        elif values and all(isinstance(value, (int, float)) and not isinstance(value, bool) for value in values):
            value_type = "number"
        else:
            value_type = "text"
        lowered = field.casefold()
        if "litre" in lowered:
            unit = "litres"
        elif "coverage" in lowered or lowered.endswith("_days"):
            unit = "days"
        elif "week" in lowered or "runs_per_week" in lowered:
            unit = "per_week"
        elif "pct" in lowered or "share" in lowered:
            unit = "ratio_or_percent"
        else:
            unit = None
        if field.startswith("legacy_") or table_name.startswith("legacy_"):
            status = "legacy_audit"
        elif table_name.startswith("sap_character") or table_name in {
            "query_catalog",
            "external_dependencies",
            "enrichment_pallet_candidates",
            "validation_issues",
        }:
            status = "diagnostic"
        else:
            status = "canonical"
        result.append(
            {
                "table": table_name,
                "field": field,
                "type": value_type,
                "nullable": int(len(values) < len(rows)),
                "status": status,
                "unit": unit,
                "source": source,
                "derivation": (
                    "legacy workbook value retained without authority"
                    if status == "legacy_audit"
                    else "deterministically extracted or derived by the named table contract"
                ),
            }
        )
    return result
