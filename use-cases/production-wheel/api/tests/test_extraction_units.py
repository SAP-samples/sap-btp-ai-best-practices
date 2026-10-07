"""Unit tests for workbook extraction and pallet-resolution rules."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from production_wheel.extraction.common import (
    find_table,
    parse_lines,
    query_catalog,
    table_name_matches,
)
from production_wheel.extraction.enrichment_workbook import resolve_pallets
from production_wheel.extraction.primary_workbook import classify_fini_admission


def test_find_table_uses_dynamic_reference(tmp_path: Path) -> None:
    """Table extraction follows Excel metadata instead of fixed row limits."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Material", "Demand"])
    sheet.append(["A", 10])
    sheet.append(["B", 20])
    sheet.add_table(Table(displayName="DynamicInput", ref="A1:B3"))
    path = tmp_path / "dynamic.xlsx"
    workbook.save(path)

    extracted = find_table(workbook, "DynamicInput")

    assert extracted.reference == "A1:B3"
    assert [row["Material"] for row in extracted.records] == ["A", "B"]


def test_plant_placeholder_contract_matches_any_plant_code() -> None:
    """``{plant}`` contracts match plant-specific table names, exact names stay exact."""

    contract = "PCK_per_filling_line___{plant}_specific"
    assert table_name_matches(contract, "PCK_per_filling_line___PL01_specific")
    assert table_name_matches("{plant}_specific___ABC_in_{plant}", "PL01_SB_specific___ABC_in_PL02")
    assert not table_name_matches(contract, "PCK_per_filling_line___specific")
    assert not table_name_matches("Site", "Site2")

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Material", "Line"])
    sheet.append(["A", 1])
    sheet.add_table(Table(displayName="PCK_per_filling_line___PL01_specific", ref="A1:B2"))
    assert find_table(workbook, contract).name == "PCK_per_filling_line___PL01_specific"


def test_parse_lines_preserves_unknown_separately_from_no_lines() -> None:
    """A blank source is unknown while a populated pattern can contain no digits."""

    assert parse_lines(None) is None
    assert parse_lines("1  2  -  4") == (1, 2, 4)
    assert parse_lines("- - - - - -") == ()


def test_pallet_resolution_is_primary_first_and_requires_one_positive_candidate() -> None:
    """Enrichment fills blanks only and never overwrites a populated primary value."""

    fini = [
        {
            "plant": "P",
            "material": "A",
            "fini_plant_key": "A/P",
            "model_status": "modeled",
            "pallet_litres_primary": 50.0,
        },
        {
            "plant": "P",
            "material": "B",
            "fini_plant_key": "B/P",
            "model_status": "modeled",
            "pallet_litres_primary": None,
        },
        {
            "plant": "P",
            "material": "C",
            "fini_plant_key": "C/P",
            "model_status": "modeled",
            "pallet_litres_primary": None,
        },
    ]
    candidates = [
        {"fini_plant_key": "A/P", "pallet_litres": 60.0},
        {"fini_plant_key": "B/P", "pallet_litres": 80.0},
        {"fini_plant_key": "B/P", "pallet_litres": 80.0},
        {"fini_plant_key": "C/P", "pallet_litres": 90.0},
        {"fini_plant_key": "C/P", "pallet_litres": 100.0},
    ]

    updated, resolutions, issues = resolve_pallets(fini, candidates)

    assert updated[0]["pallet_litres_resolved"] == 50.0
    assert updated[0]["pallet_resolution_source"] == "PRIMARY"
    assert updated[1]["pallet_litres_resolved"] == 80.0
    assert updated[1]["pallet_resolution_source"] == "ENRICHMENT_FINI_pallets"
    assert updated[2]["pallet_litres_resolved"] is None
    assert updated[2]["pallet_conflict"] == 1
    assert resolutions[2]["resolution_status"] == "conflicting_enrichment_candidates"
    assert {issue["rule_id"] for issue in issues} == {
        "PALLET_NONBLANK_SOURCE_DISCREPANCY",
    }


def test_populated_nonpositive_primary_pallet_is_not_overwritten() -> None:
    """A populated invalid primary value is preserved and blocks modeled use."""

    fini = [
        {
            "plant": "P",
            "material": "A",
            "fini_plant_key": "A/P",
            "model_status": "modeled",
            "pallet_litres_primary": 0.0,
        }
    ]

    updated, resolutions, issues = resolve_pallets(
        fini, [{"fini_plant_key": "A/P", "pallet_litres": 80.0}]
    )

    assert updated[0]["pallet_litres_resolved"] == 0.0
    assert resolutions[0]["resolution_status"] == "invalid_primary_nonpositive"
    assert not issues


def _admission_row(**updates: object) -> dict[str, object]:
    """Return one valid weekly FINI row for admission tests."""

    return {
        "plant": "P1",
        "material": "F1",
        "sefi": "S1",
        "production_frequency": "01W. Weekly",
        "forecast_litres_12m": 1_000.0,
        "line_data_status": "known",
        "eligible_lines": "1|2",
        "package_volume": 1.0,
        "pallet_litres_resolved": 100.0,
        "fixed_pv": "PV1",
        "baseline_group_key": None,
    } | updates


def _pv_rows() -> tuple[dict[str, object], ...]:
    """Return one finite production version for admission tests."""

    return (
        {
            "plant": "P1",
            "sefi": "S1",
            "production_version": "PV1",
            "lot_size_litres": 500.0,
        },
    )


def test_frequency_alone_decides_business_scope_without_a_baseline() -> None:
    """Weekly/biweekly blanks are modeled while later frequencies are out of scope."""

    weekly = classify_fini_admission(_admission_row(), _pv_rows())
    biweekly = classify_fini_admission(
        _admission_row(production_frequency="2W. Every two weeks"), _pv_rows()
    )
    later = classify_fini_admission(
        _admission_row(production_frequency="03W. Every three weeks"), _pv_rows()
    )

    assert (weekly["model_status"], weekly["production_frequency_code"]) == (
        "modeled",
        "01W",
    )
    assert (biweekly["model_status"], biweekly["production_frequency_code"]) == (
        "modeled",
        "02W",
    )
    assert later["model_status"] == "out_of_scope"
    assert later["exclusion_reason"] == "production_frequency_03w_out_of_scope"


def test_in_scope_data_failures_are_exclusions_with_precise_reasons() -> None:
    """Required-data failures never masquerade as frequency scope decisions."""

    cases = {
        "missing_demand": {"forecast_litres_12m": None},
        "missing_line_evidence": {"line_data_status": "missing", "eligible_lines": None},
        "missing_package_volume": {"package_volume": None},
        "missing_pallet_data": {"pallet_litres_resolved": None},
    }
    for reason, updates in cases.items():
        result = classify_fini_admission(_admission_row(**updates), _pv_rows())
        assert (result["model_status"], result["exclusion_reason"]) == (
            "excluded",
            reason,
        )


def test_missing_fixed_pv_is_default_exclusion_but_optimized_sensitivity_scope() -> None:
    """An eligible block PV can support sensitivity without inventing a fixed PV."""

    result = classify_fini_admission(_admission_row(fixed_pv=None), _pv_rows())

    assert (result["model_status"], result["exclusion_reason"]) == (
        "excluded",
        "missing_fixed_pv",
    )
    assert result["optimized_pv_model_status"] == "modeled"
    assert result["optimized_pv_exclusion_reason"] is None


def test_query_catalog_assigns_one_governed_status_per_query() -> None:
    """Query classification keeps activated, diagnostic, and ignored rules distinct."""

    source = '''shared #"FINI pallets" = File.Contents("pallet.xlsx");
shared #"Forecast source" = Web.Contents("https://example.test");
shared Other = Files[CommonInput];
'''

    queries, dependencies = query_catalog(source, "ENRICHMENT")

    assert [row["implementation_status"] for row in queries] == [
        "activated",
        "diagnostic_only",
        "ignored",
    ]
    assert {(row["kind"], row["source"]) for row in dependencies} == {
        ("file", "pallet.xlsx"),
        ("url", "https://example.test"),
        ("parameterized_file", "CommonInput"),
    }
