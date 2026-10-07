"""Exercise workbook import without historical sheets or intermediate files."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from production_wheel.extraction.datasets import extract_workbooks
from production_wheel.scenarios import canonical_inputs_from_tables
from production_wheel.schemas import RunConfig


def _workbook(path: Path) -> None:
    """Write a minimal supported main/PV workbook to the supplied test path."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inputs"
    sheet.append(
        [
            "Material",
            "SEFI",
            "Plant",
            "Filling line",
            "Production Frequency",
            "Total FINI forecast (litres)",
            "Volume",
            "Pallet in litres",
            "Selected PV",
        ]
    )
    sheet.append(["100", "200", "TEST", "L1/L2/COROB", "01W", 10000, 5, 500, "PV - 1"])
    sheet.add_table(Table(displayName="FINI_grouping_for_wheel", ref="A1:I2"))
    versions = workbook.create_sheet("Versions")
    versions.append(["SEFI", "Clasificador", "Lot size (Litres)"])
    versions.append(["200", "PV - 1", 1000])
    versions.add_table(Table(displayName="Production_versions___List", ref="A1:C2"))
    workbook.save(path)


def test_in_memory_import_preserves_string_lines_and_settings(tmp_path):
    """Import plant string lines and configured demand calendar without output files."""
    path = tmp_path / "input.xlsx"
    _workbook(path)
    result = extract_workbooks(
        path,
        settings={"demand_days": 200, "productive_weeks": 40, "canonical_factor": 0.8},
    )
    row = result["tables"]["fini_master"][0]
    assert row["eligible_lines"] == "COROB|L1|L2"
    assert row["avg_daily_demand"] == 50
    assert row["source_sheet"] == "Inputs"
    assert result["metadata"]["plant"] == "TEST"
    assert result["metadata"]["parameter_sources"]["demand_days"] == "user"
    assert list(tmp_path.iterdir()) == [path]
    inputs = canonical_inputs_from_tables(result["tables"])
    assert len(inputs.members) == 1


def test_custom_calendar_is_valid():
    """Allow positive calendar and factor overrides in typed scenario settings."""
    config = RunConfig(demand_days=200, productive_weeks=40, canonical_factor=0.8)
    assert config.demand_days == 200


def test_fixed_line_override_and_optional_alias(tmp_path):
    """Use FixedLine source evidence and fixed-line calendar aliases without history."""
    from openpyxl import load_workbook

    path = tmp_path / "renamed.xlsx"
    _workbook(path)
    workbook = load_workbook(path)
    sheet = workbook.create_sheet("Corrections")
    sheet.append(["Material", "Material Description", "Correct Filling line"])
    sheet.append(["100", "Example", "COROB"])
    sheet.add_table(Table(displayName="FixedLine", ref="A1:C2"))
    workbook.save(path)
    result = extract_workbooks(path)
    row = result["tables"]["fini_master"][0]
    assert result["metadata"]["format_id"] == "fixed_line_profile"
    assert row["eligible_lines"] == "COROB"
    assert row["filling_line_raw"] == "L1/L2/COROB"
    assert row["line_correction_evidence"][0]["source_row"] == 2
    assert row["line_resolution_source"] == "FixedLine"


def test_dc_selling_profile_detected_from_tables(tmp_path):
    """Identify the DC-selling profile using workbook structure rather than its filename."""
    from openpyxl import load_workbook

    path = tmp_path / "renamed.xlsx"
    _workbook(path)
    workbook = load_workbook(path)
    sheet = workbook.create_sheet("Support")
    sheet.append(["Material", "Plant", "Lot size"])
    sheet.append(["100", "TEST", "EX"])
    sheet.add_table(Table(displayName="PL01_SB_specific___ABC_in_PL01", ref="A1:C2"))
    calendar = workbook.create_sheet("Calendar")
    calendar.append(["Material", "Plant", "Planning calendar"])
    calendar.append(["100", "TEST", "WEEK"])
    calendar.add_table(
        Table(displayName="SB_FINI_Planning_calendar_at_P_site", ref="A1:C2")
    )
    workbook.save(path)
    result = extract_workbooks(path)
    assert result["metadata"]["format_id"] == "dc_selling_profile"
    assert (
        result["tables"]["planning_calendar_inform"][0]["planning_calendar"] == "WEEK"
    )


def test_configurable_frequency_and_unknown_line_are_separate(tmp_path):
    """Include requested frequency but exclude records without line evidence."""
    from openpyxl import load_workbook

    path = tmp_path / "input.xlsx"
    _workbook(path)
    workbook = load_workbook(path)
    workbook["Inputs"]["E2"] = "04W"
    workbook.save(path)
    assert (
        extract_workbooks(path)["tables"]["fini_master"][0]["model_status"]
        == "out_of_scope"
    )
    assert (
        extract_workbooks(path, settings={"modeled_frequencies": ["04W"]})["tables"][
            "fini_master"
        ][0]["model_status"]
        == "modeled"
    )
    workbook["Inputs"]["D2"] = None
    workbook.save(path)
    row = extract_workbooks(path, settings={"modeled_frequencies": ["04W"]})["tables"][
        "fini_master"
    ][0]
    assert row["exclusion_reason"] == "missing_line_evidence"


def test_typed_and_csv_inputs_have_equal_boolean_semantics(tmp_path):
    """Normalize boolean database rows and CSV text to identical solver inputs."""
    path = tmp_path / "input.xlsx"
    _workbook(path)
    tables = extract_workbooks(path)["tables"]
    tables["fini_master"][0]["baseline_assigned"] = True
    tables["production_versions"][0]["active"] = False
    typed = canonical_inputs_from_tables(tables)
    strings = {
        name: [
            {key: str(value) if value is not None else "" for key, value in row.items()}
            for row in records
        ]
        for name, records in tables.items()
    }
    csv = canonical_inputs_from_tables(strings)
    assert typed.production_versions == csv.production_versions
    assert typed.baseline_evidence_members == csv.baseline_evidence_members
    assert typed.production_versions[0].active is False


def test_required_lines_normalizes_legacy_integers():
    """Keep numeric and named required-line witnesses in the same string contract."""
    from production_wheel.schemas import RequiredLinesConstraint

    rule = RequiredLinesConstraint(
        constraint_id="line-test",
        kind="required_lines",
        filling_lines=[1, "L1", "COROB"],
    )
    assert rule.filling_lines == ("1", "L1", "COROB")


def test_supporting_headers_preserve_distinct_symbols_and_normalization_collisions():
    """Keep every distinct source column when punctuation or names normalize alike."""
    import re
    from production_wheel.extraction.common import TableData
    from production_wheel.extraction.dataset_tables import supporting_header_map, supporting_rows

    long_header = "🧪" * 100
    table = TableData(
        "Compatibility",
        "Recipe naming logic",
        "A1:K2",
        ({
            "Group count (%)": 4,
            "Group count": 5,
            "Group-count": 6,
            ")": "right",
            "(": "left",
            "/": "slash",
            "😀": "emoji",
            "field_": "literal",
            "source_sheet": "source value",
            "owner_id": "owner value",
            long_header: "long value",
            "_source_row": 2,
        },),
    )
    mapping = supporting_header_map(table)
    rows = supporting_rows({table.name: table})["sap_character_compatibility"]
    assert set(mapping) == set(table.records[0]) - {"_source_row"}
    assert len(set(mapping.values())) == len(mapping)
    assert all(re.fullmatch(r"[a-z_][a-z_0-9]*", name) and len(name) <= 127 for name in mapping.values())
    assert {source: rows[0][target] for source, target in mapping.items()} == {
        key: value for key, value in table.records[0].items() if key != "_source_row"
    }
    assert rows[0]["source_sheet"] == "Recipe naming logic"
    assert rows[0]["source_row"] == 2


def test_symbol_headers_survive_full_workbook_extraction(tmp_path):
    """Expose source-to-stored symbol mapping and data through the public importer."""
    from openpyxl import load_workbook

    path = tmp_path / "input.xlsx"
    _workbook(path)
    workbook = load_workbook(path)
    sheet = workbook.create_sheet("Symbols")
    sheet.append(["Compatible", ")", "(", ":", ";", "+", "=", "/", "★"])
    sheet.append(["A", "yes", "no", "yes", "no", "yes", "no", "yes", "star"])
    sheet.add_table(Table(displayName="Compatibility", ref="A1:I2"))
    workbook.save(path)

    result = extract_workbooks(path)
    mapping = result["metadata"]["source_tables"]["Compatibility"]["header_mapping"]
    row = result["tables"]["sap_character_compatibility"][0]
    assert len(set(mapping.values())) == len(mapping)
    assert {header: row[column] for header, column in mapping.items()} == {
        "Compatible": "A",
        ")": "yes",
        "(": "no",
        ":": "yes",
        ";": "no",
        "+": "yes",
        "=": "no",
        "/": "yes",
        "★": "star",
    }


def test_hana_numeric_flags_and_array_lines_preserve_canonical_evidence(tmp_path):
    """Retain baseline/PV flags after HANA DOUBLE storage without changing IDs."""
    import json

    path = tmp_path / "input.xlsx"
    _workbook(path)
    tables = extract_workbooks(path)["tables"]
    row = tables["fini_master"][0]
    row["baseline_assigned"] = 1.0
    row["pallet_conflict"] = "0.0"
    row["material"] = "001.0"
    row["eligible_lines"] = ["L1", "COROB"]
    evidence = [{"material": "001.0", "correct_filling_line": "COROB", "source_row": 2}]
    row["line_correction_evidence"] = evidence
    tables["production_versions"][0]["active"] = 0.0
    inputs = canonical_inputs_from_tables(tables)
    assert len(inputs.baseline_evidence_members) == 1
    assert inputs.baseline_evidence_members[0].fini_id == "001.0"
    assert inputs.production_versions[0].active is False
    assert inputs.members[0].eligible_lines == frozenset({"L1", "COROB"})
    assert inputs.fini_rows[0]["baseline_assigned"] == "1"
    assert inputs.fini_rows[0]["pallet_conflict"] == "0"
    assert json.loads(inputs.fini_rows[0]["line_correction_evidence"]) == evidence
    row["baseline_assigned"] = "1.0"
    tables["production_versions"][0]["active"] = "1.0"
    assert len(canonical_inputs_from_tables(tables).baseline_evidence_members) == 1
    assert canonical_inputs_from_tables(tables).production_versions[0].active is True
    row["baseline_assigned"] = "0.0"
    tables["production_versions"][0]["active"] = "0.0"
    assert canonical_inputs_from_tables(tables).baseline_evidence_members == ()
    assert canonical_inputs_from_tables(tables).production_versions[0].active is False
