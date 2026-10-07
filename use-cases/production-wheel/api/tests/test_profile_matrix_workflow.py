"""Excel and natural-language compatibility use the same strict matrix contract."""

from io import BytesIO
import pytest
from openpyxl import load_workbook
from app.workspace.profile_matrix import matrix_workbook, read_matrix_workbook, expand_matrix_description
from app.workspace.profile_interpretation import Interpretation, InterpretationInput, finalize_interpretation


def test_excel_matrix_roundtrip_changed_volumes_and_invalid_cells():
    """Accept fewer/new volume axes and reject incomplete, asymmetric or formula cells."""
    rows = expand_matrix_description({'volumes': ['1', '2', '9'],
        'compatible_groups': [['1', '2']], 'default_status': 'N'})
    workbook = load_workbook(BytesIO(matrix_workbook(rows)))
    sheet = workbook['Volume compatibility']
    sheet.delete_cols(4); sheet.delete_rows(4)
    sheet.cell(1, 3, 3); sheet.cell(3, 1, 3)
    output = BytesIO(); workbook.save(output)
    parsed = read_matrix_workbook(output.getvalue())
    assert {r['volume_a'] for r in parsed} == {'1', '3'}
    assert len(parsed) == 4
    for invalid in (None, 'MAYBE', '=1+1', 'N'):
        sheet.cell(2, 3).value = invalid
        output = BytesIO(); workbook.save(output)
        with pytest.raises(ValueError):
            read_matrix_workbook(output.getvalue())


def test_clarification_never_exposes_partial_rules_or_matrix_for_acceptance():
    """An unresolved intent blocks the complete proposal, even if a model emits rules."""
    body = InterpretationInput(text='Keep similar products together', plant='P')
    parsed = Interpretation(interpretation='Please define similarity.', clarification_required=True,
        rules=[{'kind': 'group_rule', 'constraint_id': 'partial',
                'assertion': {'op': 'literal', 'value': True}}])
    result = finalize_interpretation(parsed, body)
    assert result['rules'] == [] and result['matrix_rows'] is None
    assert result['clarification_required']
    matrix = Interpretation(interpretation='1 and 2 compatible; all other pairs forbidden.',
        volume_compatibility={'volumes': ['1', '2', '3'],
                             'compatible_groups': [['1', '2']], 'default_status': 'N'})
    result = finalize_interpretation(matrix, body)
    assert not result['clarification_required']
    assert len(result['matrix_rows']) == 9
    assert next(r['status'] for r in result['matrix_rows'] if r['volume_a'] == '1' and r['volume_b'] == '3') == 'N'


def test_provider_string_null_means_no_matrix_change():
    """A provider's literal null string can only remove a proposal, never create one."""
    result = Interpretation(interpretation='Clarification needed.', clarification_required=True,
                            volume_compatibility='null')
    assert result.volume_compatibility is None


def test_matrix_http_preview_and_explicit_profile_save():
    """Downloads/uploads preview only; explicit save persists accepted matrix-only prose."""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.routers.workspace import router
    from app.workspace.dependencies import get_service
    from app.workspace.repository import MemoryRepository
    from app.workspace.service import WorkspaceService
    service = WorkspaceService(MemoryRepository())
    app = FastAPI(); app.include_router(router)
    app.dependency_overrides[get_service] = lambda: service
    client = TestClient(app)
    rows = expand_matrix_description({'volumes': ['1', '2', '99'],
        'compatible_groups': [['1', '2']], 'default_status': 'N'})
    download = client.post('/api/plant-profiles/matrix-template', json={'matrix_rows': rows})
    assert download.status_code == 200
    preview = client.post('/api/plant-profiles/matrix-preview', files={'file': ('matrix.xlsx', download.content)})
    assert preview.status_code == 200 and preview.json()['volume_count'] == 3
    assert service.list_plant_profiles() == []
    text = '1 and 2 compatible, all other different-volume pairs forbidden.'
    body = {'name': 'Preview test', 'plant': 'P', 'matrix_rows': preview.json()['matrix_rows'],
            'rules_text': text, 'matrix_rules_text': text}
    saved = client.post('/api/plant-profiles', json=body)
    assert saved.status_code == 200 and saved.json()['rules'] == []
    assert service.get_plant_profile(saved.json()['profile_id'])['matrix_rows'] == rows
    assert client.post('/api/plant-profiles/matrix-preview', files={'file': ('bad.xlsx', b'not an xlsx')}).status_code == 422
    help_text = client.get('/api/plant-profile-help').json()
    assert 'horizon_days' in help_text and 'rules-text' in help_text
