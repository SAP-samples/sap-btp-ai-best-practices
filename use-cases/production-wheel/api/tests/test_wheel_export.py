"""Source-style line and production-wheel export contracts."""

import io

import pytest
from app.workspace.repository import MemoryRepository
from app.workspace.service import WorkspaceService
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from production_wheel.extraction.dataset_tables import parse_string_lines


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("- 2  3   -   -   -", ("2", "3")),
        ("11 / 12 / 18 / 23", ("11", "12", "18", "23")),
        ("L1 / L2", ("L1", "L2")),
        (None, None),
        ("- - -", ()),
    ],
)
def test_lines_are_exact_string_identifiers(raw, expected):
    """Numeric alternatives, positional slots and named lines remain distinct."""
    assert parse_string_lines(raw) == expected


def test_subset_format_keeps_source_convention():
    """Group-common lines retain original spacing, positional slots and named tokens."""
    from production_wheel.extraction.line_format import format_lines

    assert format_lines(("2", "3"), "- 2  3   -   -   -") == "- 2  3   -   -   -"
    assert format_lines(("3",), "- 2  3   -   -   -") == "- -  3   -   -   -"
    assert format_lines(("12", "23"), "11 / 12 / 18 / 23") == "12 / 23"
    assert format_lines(("L2",), "L1 / L2") == "L2"


def test_wheel_export_uses_source_columns_and_selected_solution():
    """Original line text survives while proposed PV/group values replace historical ones."""
    from app.workspace.wheel_export import export_wheel

    book = Workbook()
    sheet = book.active
    sheet.title = "Source"
    headers = [
        "Material",
        "SEFI",
        "Filling line",
        "Subgroup",
        "Selected PV",
        "RoC selected PV",
    ]
    sheet.append(headers)
    sheet.append(["100", "s", "- 2  3   -   -   -", "old", "old-pv", 99])
    sheet.add_table(Table(displayName="FINI_grouping_for_wheel", ref="A1:F2"))
    source = io.BytesIO()
    book.save(source)
    repo = MemoryRepository()
    service = WorkspaceService(repo)
    repo.insert(
        "datasets",
        "d",
        {"dataset_id": "d", "metadata": {"source_files": ["source.xlsx"]}},
    )
    repo.put_artifact("d", "source.xlsx", source.getvalue())
    repo.insert(
        "runs",
        "r",
        {
            "run_id": "r",
            "dataset_id": "d",
            "results_ready": True,
            "metadata": {},
            "request": {"config": {"productive_weeks": 50}},
        },
    )
    repo.replace_tables(
        "r",
        {
            "solutions": [{"point_index": 1}],
            "members": [
                {
                    "point_index": 1,
                    "plant": "p",
                    "sefi": "s",
                    "material": "100",
                    "source_row": 2,
                    "source_sheet": "Source",
                    "proposed_subgroup": "g",
                    "proposed_subgroup_code": "G001",
                    "selected_pv": "new-pv",
                    "fini_adjusted_coverage_days": 12,
                    "common_lines": "3",
                    "selected_line": "3",
                    "model_status": "modeled",
                }
            ],
        },
    )
    output = load_workbook(io.BytesIO(export_wheel(service, "r", 1)), data_only=True)
    sheet = output["Production wheel"]
    values = list(sheet.values)
    assert list(values[0][:6]) == headers
    row = dict(zip(values[0], values[1]))
    assert row["Filling line"] == "- 2  3   -   -   -"
    assert row["Subgroup"] == "G001"
    assert row["Selected PV"] == "new-pv"
    assert row["RoC selected PV"] == 12
    assert row["Common eligible lines"] == "- -  3   -   -   -"
    assert row["Point"] == 1
    assert row["Selected filling line"] == "3"
    with pytest.raises(ValueError, match="point"):
        export_wheel(service, "r", 2)


def test_export_values_keep_unassigned_rows_blank_and_source_strings_literal():
    """Never carry old group assignments into excluded rows or execute source formulas."""
    from app.workspace.wheel_export import _sheet, wheel_rows

    rows = wheel_rows(
        ["Material", "Subgroup", "Selected PV", "Filling line"],
        {},
        [
            {
                "material": "=UNTRUSTED()",
                "model_status": "excluded",
                "fixed_pv": "old",
                "filling_line_raw": "L1 / L2",
            }
        ],
        {},
    )
    assert rows[0]["Subgroup"] is None
    assert rows[0]["Selected PV"] is None
    assert rows[0]["Proposal status"] == "No proposed assignment"
    book = Workbook()
    _sheet(book, "Safe", list(rows[0]), rows)
    assert book["Safe"]["A2"].data_type == "s"


def test_source_layout_finds_primary_after_enrichment_for_migrated_snapshots():
    """Legacy artifact order must not make an enrichment workbook the output template."""
    from app.workspace.wheel_export import source_layout

    repo = MemoryRepository()
    service = WorkspaceService(repo)
    repo.insert(
        "datasets",
        "d",
        {"metadata": {"source_files": ["enrichment.xlsx", "source.xlsx"]}},
    )
    book = Workbook()
    stream = io.BytesIO()
    book.save(stream)
    repo.put_artifact("d", "enrichment.xlsx", stream.getvalue())
    book.active.append(["Material", "Filling line"])
    book.active.append(["001", "L1 / L2"])
    book.active.add_table(Table(displayName="FINI_grouping_for_wheel", ref="A1:B2"))
    stream = io.BytesIO()
    book.save(stream)
    repo.put_artifact("d", "source.xlsx", stream.getvalue())
    headers, rows = source_layout(service, "d")
    assert headers == ["Material", "Filling line"]
    assert len(rows) == 1
