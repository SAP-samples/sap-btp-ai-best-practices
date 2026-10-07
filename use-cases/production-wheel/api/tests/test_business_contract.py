"""Tests for typed settings, volume compatibility, and KPI formulas."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError

from production_wheel.matrix import (
    BASELINE_EMPIRICAL_MATRIX_VERSION,
    CUSTOMER_OPERATIONAL_VOLUME_FAMILIES_V1,
    CUSTOMER_OPERATIONAL_VOLUME_FAMILY_MATRIX_V1,
    CUSTOMER_VOLUME_CATALOG,
    MODELED_VOLUMES,
    SYNTHETIC_VOLUME_MATRIX_V1,
    UNKNOWN_LINE_FEASIBLE,
    UNKNOWN_NO_CURRENT_LINE_OVERLAP,
    YES_OBSERVED_BASELINE,
    build_baseline_empirical_matrix,
    customer_volume_families,
    matrix_compatible,
    matrix_exceptions,
)
from production_wheel.candidates import CandidateMember
from production_wheel.metrics import (
    changeover_contribution,
    coverage_summary,
    effective_batch,
    group_coverage,
    group_frequency,
    pallet_allocations,
    proportional_allocations,
    target_band_key,
)
from production_wheel.schemas import (
    ApprovalStatus,
    CannotLinkConstraint,
    ConstraintScope,
    CoverageBasis,
    GroupSizeMode,
    GroupSizePolicy,
    MatrixMode,
    MaxGroupSizeConstraint,
    MustLinkConstraint,
    PalletFormula,
    RunConfig,
    SolveRequest,
    VersionIdentifiers,
    VolumeCompatibilityOverrideConstraint,
)
from production_wheel.validation import (
    ChangeClass,
    ValidationContext,
    ValidationStatus,
    classify_config_change,
    get_capabilities,
    validate_request,
)


def test_synthetic_matrix_is_complete_symmetric_and_diagonal() -> None:
    """The governed synthetic matrix contains every ordered volume pair."""

    assert len(MODELED_VOLUMES) == 23
    assert len(SYNTHETIC_VOLUME_MATRIX_V1) == 23 * 23
    lookup = {
        (row["volume_a"], row["volume_b"]): row["compatible"]
        for row in SYNTHETIC_VOLUME_MATRIX_V1
    }
    for volume_a in MODELED_VOLUMES:
        for volume_b in MODELED_VOLUMES:
            left, right = format(volume_a, "f"), format(volume_b, "f")
            assert lookup[(left, right)] == lookup[(right, left)]
            if left == right:
                assert lookup[(left, right)] == "YES"
    assert matrix_compatible("0.25", "1")
    assert not matrix_compatible("1", "2.5")


def test_matrix_exceptions_count_unordered_fini_pairs() -> None:
    """Compatibility diagnostics count each incompatible FINI pair once."""

    exceptions = matrix_exceptions([("A", "1"), ("B", "2.5"), ("C", "3")])

    assert [(row["material_a"], row["material_b"]) for row in exceptions] == [
        ("A", "B"),
        ("A", "C"),
    ]


def test_customer_operational_matrix_matches_approved_family_counts() -> None:
    """Inclusive overlapping families produce the exact governed 34-volume matrix."""

    rows = CUSTOMER_OPERATIONAL_VOLUME_FAMILY_MATRIX_V1
    assert len(CUSTOMER_VOLUME_CATALOG) == 34
    assert len(rows) == 34 * 34
    lookup = {
        (Decimal(str(row["volume_a"])), Decimal(str(row["volume_b"]))): row["status"]
        for row in rows
    }
    for left in CUSTOMER_VOLUME_CATALOG:
        assert lookup[(left, left)] == "Y"
        for right in CUSTOMER_VOLUME_CATALOG:
            assert lookup[(left, right)] == lookup[(right, left)]

    preferred = sum(
        lookup[(left, right)] == "Y"
        for index, left in enumerate(CUSTOMER_VOLUME_CATALOG)
        for right in CUSTOMER_VOLUME_CATALOG[index + 1 :]
    )
    assert (preferred, 561 - preferred) == (181, 380)
    modeled_preferred = sum(
        lookup[(left, right)] == "Y"
        for index, left in enumerate(MODELED_VOLUMES)
        for right in MODELED_VOLUMES[index + 1 :]
    )
    assert (modeled_preferred, 253 - modeled_preferred) == (78, 175)
    assert set(customer_volume_families("4")) == {"2_TO_5_L", "4_TO_10_L"}
    assert set(customer_volume_families("10")) == {"4_TO_10_L", "10_TO_15_L"}
    assert lookup[(Decimal("0.5"), Decimal("1"))] == "Y"
    assert lookup[(Decimal("2.5"), Decimal("5"))] == "Y"
    assert lookup[(Decimal("4"), Decimal("10"))] == "Y"
    assert lookup[(Decimal("10"), Decimal("15"))] == "Y"
    assert lookup[(Decimal("0.75"), Decimal("5"))] == "AVOID"
    assert lookup[(Decimal("1"), Decimal("10"))] == "AVOID"
    assert "N" not in lookup.values()

    exceptions = matrix_exceptions((("A", "1"), ("B", "10")), rows)
    assert exceptions[0]["status"] == "AVOID"
    assert exceptions[0]["matrix_version"] == CUSTOMER_OPERATIONAL_VOLUME_FAMILIES_V1
    assert exceptions[0]["family_a"] == "LE_1_L"
    assert exceptions[0]["family_b"] == "4_TO_10_L|10_TO_15_L"
    capabilities = get_capabilities()
    customer = capabilities["matrix_versions"][  # type: ignore[index]
        CUSTOMER_OPERATIONAL_VOLUME_FAMILIES_V1
    ]
    assert len(customer["package_volumes"]) == 34


def test_customer_matrix_workbook_matches_generated_rule_and_validation() -> None:
    """The delivered workbook mirrors code and exposes controlled edit values."""

    repository_root = Path(__file__).resolve().parents[2]
    workbook_path = (
        repository_root
        / "outputs"
        / "volume-matrix-request-20260828"
        / "customer-operational-volume-families-v1.xlsx"
    )
    if not workbook_path.is_file():
        pytest.skip("delivered matrix workbook is not part of this checkout")
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    assert workbook.sheetnames == [
        "Volume Matrix",
        "Instructions",
        "Conditional Rules",
    ]
    matrix = workbook["Volume Matrix"]
    expected_volumes = [format(volume, "f") for volume in CUSTOMER_VOLUME_CATALOG]
    assert [str(matrix.cell(1, column).value) for column in range(2, 36)] == (
        expected_volumes
    )
    assert [str(matrix.cell(row, 1).value) for row in range(2, 36)] == (
        expected_volumes
    )
    expected_statuses = {
        (str(row["volume_a"]), str(row["volume_b"])): str(row["status"])
        for row in CUSTOMER_OPERATIONAL_VOLUME_FAMILY_MATRIX_V1
    }
    for row_index, volume_a in enumerate(expected_volumes, start=2):
        for column_index, volume_b in enumerate(expected_volumes, start=2):
            assert matrix.cell(row_index, column_index).value == expected_statuses[
                (volume_a, volume_b)
            ]
    validations = tuple(matrix.data_validations.dataValidation)
    assert [(str(item.sqref), item.formula1) for item in validations] == [
        ("B2:AI35", '"Y,AVOID,N,REVIEW"')
    ]
    instructions = workbook["Instructions"]
    assert instructions["B7"].value == CUSTOMER_OPERATIONAL_VOLUME_FAMILIES_V1
    assert instructions["B9"].value == "Y"
    assert instructions["B10"].value == "N"


def test_empirical_matrix_is_complete_symmetric_and_never_infers_no() -> None:
    """Baseline co-membership is positive evidence and absence stays unknown."""

    def evidence_member(
        fini_id: str,
        volume: float,
        lines: tuple[str, ...],
        baseline: str | None,
    ) -> CandidateMember:
        """Create one member carrying only the evidence used by the matrix."""

        return CandidateMember(
            fini_id=fini_id,
            plant="P1",
            sefi="S1",
            eligible_lines=frozenset(lines),
            demand_litres=100.0,
            pallet_litres=25.0,
            package_volume=volume,
            baseline_group=baseline,
        )

    modeled = (
        evidence_member("A", 0.25, ("1",), "G1"),
        evidence_member("B", 4.0, ("1",), "G1"),
        evidence_member("C", 2.5, ("1",), None),
    )
    historical = (
        *modeled,
        evidence_member("D", 0.5, (), "G2"),
        evidence_member("E", 3.0, (), "G2"),
    )
    rows = build_baseline_empirical_matrix(modeled, historical)
    assert len(rows) == 23 * 23
    assert {row["matrix_version"] for row in rows} == {
        BASELINE_EMPIRICAL_MATRIX_VERSION
    }
    lookup = {
        (row["volume_a"], row["volume_b"]): row["status"] for row in rows
    }
    assert all(lookup[(right, left)] == status for (left, right), status in lookup.items())
    assert lookup[("0.25", "4")] == YES_OBSERVED_BASELINE
    assert lookup[("0.5", "3")] == YES_OBSERVED_BASELINE
    assert lookup[("0.25", "2.5")] == UNKNOWN_LINE_FEASIBLE
    assert lookup[("0.5", "2.5")] == UNKNOWN_NO_CURRENT_LINE_OVERLAP
    assert lookup[("0.25", "15")] == UNKNOWN_NO_CURRENT_LINE_OVERLAP
    assert "NO" not in lookup.values()


def test_pallet_and_coverage_fixture_matches_hand_calculation() -> None:
    """Both pallet formulas and all coverage bases match a hand-worked case."""

    demands = {"A": Decimal("100"), "B": Decimal("300")}
    pallets = {"A": Decimal("60"), "B": Decimal("100")}
    batch = effective_batch(400)
    nominal = proportional_allocations(demands, batch)
    minimum = pallet_allocations(nominal, pallets, PalletFormula.MINIMUM_ONLY)
    whole = pallet_allocations(nominal, pallets, PalletFormula.WHOLE_PALLET_ROUNDING)

    assert batch == Decimal("360.00")
    assert nominal == {"A": Decimal("90.00"), "B": Decimal("270.00")}
    assert minimum == nominal
    assert whole == {"A": Decimal("120"), "B": Decimal("300")}
    assert group_coverage(CoverageBasis.BASE_GROUP, demands, batch, whole) == Decimal("225.00")
    assert group_coverage(CoverageBasis.ADJUSTED_GROUP, demands, batch, whole) == Decimal("262.5")
    assert group_coverage(CoverageBasis.WORST_FINI, demands, batch, whole) == Decimal("300")
    frequency = group_frequency(400, batch)
    assert frequency == Decimal(1) / Decimal(45)
    assert changeover_contribution(frequency, 2) == Decimal(1) / Decimal(45)


def test_summary_and_target_band_conventions_are_fixed() -> None:
    """P90 uses nearest rank and target excess uses unnormalized days."""

    coverages = [1, 10, 20, 400]
    demands = [1, 1, 1, 7]
    summary = coverage_summary(coverages, demands)
    key = target_band_key(coverages, demands, RunConfig().target_band)

    assert summary.maximum == 400
    assert summary.p90 == 400
    assert summary.median == 15
    assert summary.demand_weighted_mean == 283.1
    assert key.as_tuple() == (2.0, 35.0, 39.0, 283.1)


def test_validation_reports_manual_review_and_rule_contradictions() -> None:
    """Unsupported caps and conflicting material rules never reach the solver silently."""

    request = SolveRequest(
        config=RunConfig(
            group_size=GroupSizePolicy(
                mode=GroupSizeMode.BOUNDED_RELAXATION,
                max_excess=3,
            )
        )
    )
    assert validate_request(request).status == ValidationStatus.MANUAL_REVIEW

    contradictory = SolveRequest(
        constraints=(
            MustLinkConstraint(
                kind="must_link",
                constraint_id="must",
                materials=("A", "B"),
                approval_status=ApprovalStatus.APPROVED,
            ),
            CannotLinkConstraint(
                kind="cannot_link",
                constraint_id="cannot",
                materials=("A", "B"),
                approval_status=ApprovalStatus.APPROVED,
            ),
        )
    )
    result = validate_request(
        contradictory,
        ValidationContext(materials=frozenset({"A", "B"})),
    )
    assert result.status == ValidationStatus.INVALID
    assert {message.code for message in result.messages} >= {
        "MUST_CANNOT_CONTRADICTION",
        "TRANSITIVE_LINK_CONTRADICTION",
    }


def test_customer_matrix_accepts_its_complete_volume_catalog() -> None:
    """Typed overrides resolve the full 34-volume customer catalog."""

    request = SolveRequest(
        config=RunConfig(
            versions=VersionIdentifiers(
                matrix_version=CUSTOMER_OPERATIONAL_VOLUME_FAMILIES_V1
            )
        ),
        constraints=(
            VolumeCompatibilityOverrideConstraint(
                kind="volume_compatibility_override",
                constraint_id="full-catalog-pair",
                volume_a="0.93",
                volume_b="2",
                compatible=False,
                approval_status=ApprovalStatus.APPROVED,
            ),
        ),
    )

    assert validate_request(request).status is ValidationStatus.VALID
    assert request.config.versions.ruleset_version == "PROTOTYPE_RULESET_V5"
    assert request.config.versions.schema_version == "PROTOTYPE_SCHEMA_V9"


def test_config_change_distinguishes_reuse_narrowing_and_widening() -> None:
    """Only feasible-space widening requires candidate regeneration."""

    baseline = RunConfig(
        matrix_mode=MatrixMode.HARD,
        versions={"matrix_version": "SYNTHETIC_VOLUME_MATRIX_V1"},
    )
    wider = RunConfig(
        group_size=GroupSizePolicy(
            mode=GroupSizeMode.BOUNDED_RELAXATION,
            max_excess=1,
        )
    )
    diagnostic = RunConfig(
        matrix_mode=MatrixMode.DIAGNOSTIC,
        versions={"matrix_version": "SYNTHETIC_VOLUME_MATRIX_V1"},
    )
    objective_only = baseline.model_copy(
        update={"coverage_basis": CoverageBasis.WORST_FINI}
    )

    assert classify_config_change(baseline, wider).classification == ChangeClass.WIDENING
    assert classify_config_change(baseline, diagnostic).classification == ChangeClass.WIDENING
    assert classify_config_change(diagnostic, baseline).classification == ChangeClass.NARROWING
    assert classify_config_change(baseline, objective_only).classification == ChangeClass.REUSE

    flexible = diagnostic.model_copy(update={"matrix_mode": MatrixMode.FLEXIBLE})
    assert classify_config_change(diagnostic, flexible).classification == ChangeClass.REUSE
    assert (
        diagnostic.structural_ruleset_fingerprint()
        == flexible.structural_ruleset_fingerprint()
    )


def test_request_schema_rejects_executable_or_unknown_fields_and_hashes_stably() -> None:
    """The future-agent boundary accepts only the finite typed constraint DSL."""

    with pytest.raises(ValidationError):
        SolveRequest.model_validate(
            {
                "constraints": [
                    {
                        "kind": "max_group_size",
                        "constraint_id": "cap",
                        "maximum": 7,
                        "python": "exec('unsafe')",
                    }
                ]
            }
        )
    first = SolveRequest()
    second = SolveRequest.model_validate(first.model_dump(mode="json"))
    assert first.request_id() == second.request_id()
    assert first.config.configuration_id() == second.config.configuration_id()


def test_operations_first_prioritizes_jch_before_weighted_mean() -> None:
    """Operations-first preserves violation severity but then minimizes J_CH."""

    from production_wheel.metrics import objective_key
    from production_wheel.schemas import CoverageMode

    key = objective_key(
        CoverageMode.OPERATIONS_FIRST,
        [10, 20],
        [1, 9],
        3,
        RunConfig().target_band,
    )
    assert key == (0.0, 0.0, 0.0, 3.0, 19.0)


def test_constraint_governance_and_malformed_values_fail_safely() -> None:
    """Draft/rejected rules pause or fail and malformed values never raise."""

    draft = SolveRequest(
        constraints=(
            MustLinkConstraint(
                kind="must_link",
                constraint_id="draft",
                materials=("A", "B"),
            ),
        )
    )
    assert validate_request(
        draft, ValidationContext(materials=frozenset({"A", "B"}))
    ).status == ValidationStatus.MANUAL_REVIEW

    cap = SolveRequest(
        constraints=(
            MaxGroupSizeConstraint(
                kind="max_group_size",
                constraint_id="cap10",
                maximum=10,
                approval_status=ApprovalStatus.APPROVED,
            ),
        )
    )
    assert validate_request(cap).status == ValidationStatus.MANUAL_REVIEW

    malformed = SolveRequest(
        constraints=(
            VolumeCompatibilityOverrideConstraint(
                kind="volume_compatibility_override",
                constraint_id="volume",
                volume_a="formula()",
                volume_b="1",
                compatible=True,
                approval_status=ApprovalStatus.APPROVED,
            ),
        )
    )
    assert validate_request(malformed).status == ValidationStatus.INVALID


def test_configuration_hash_excludes_scenario_label() -> None:
    """Named scenarios deduplicate when every executable setting is identical."""

    assert RunConfig(scenario_id="one").configuration_id() == RunConfig(
        scenario_id="two"
    ).configuration_id()
