"""Tests for Cloud Foundry-safe Office document extractors."""

from email.message import EmailMessage
from pathlib import Path

from docx import Document
from fpdf import FPDF
from openpyxl import Workbook

from app.services.document_extractors import (
    extract_docx,
    extract_eml,
    extract_pdf,
    extract_xlsx,
)


def _write_email(path: Path, message: EmailMessage) -> None:
    """Write an RFC 822 email fixture to disk.

    Inputs:
        path: Destination path for the `.eml` fixture.
        message: EmailMessage instance with headers, body, and optional parts.

    Outputs:
        None. The serialized message is written as bytes.
    """

    path.write_bytes(message.as_bytes())


def test_extract_docx_preserves_headings_and_tables(tmp_path: Path) -> None:
    """Verify DOCX extraction keeps heading context and table text.

    Args:
        tmp_path: Temporary directory provided by pytest for creating the DOCX.

    Returns:
        None. Assertions validate the extracted document metadata and blocks.
    """

    document_path = tmp_path / "strategy.docx"
    document = Document()
    document.add_heading("Strategic Planning", level=1)
    document.add_paragraph("Review the yearly strategic priorities.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Budget review"
    table.cell(0, 1).text = "CFO"
    document.save(document_path)

    extracted = extract_docx(document_path)

    assert extracted.file_name == "strategy.docx"
    assert any(
        block.section_label == "Strategic Planning" for block in extracted.blocks
    )
    assert any(
        block.block_type == "table" and "Budget review" in block.text
        for block in extracted.blocks
    )


def test_extract_xlsx_preserves_sheet_and_cells(tmp_path: Path) -> None:
    """Verify XLSX extraction keeps sheet names, cells, and row ranges.

    Args:
        tmp_path: Temporary directory provided by pytest for creating the XLSX.

    Returns:
        None. Assertions validate the extracted sheet block metadata and text.
    """

    workbook_path = tmp_path / "objectives.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Objectives"
    sheet["C5"] = "Budget approval"
    sheet["D5"] = "CFO"
    sheet["E6"] = "=1+1"
    workbook.save(workbook_path)
    workbook.close()

    extracted = extract_xlsx(workbook_path)
    block = extracted.blocks[0]

    assert extracted.file_name == "objectives.xlsx"
    assert block.sheet_name == "Objectives"
    assert block.table_name == "Objectives!C5:E6"
    assert block.row_start == 5
    assert block.row_end == 6
    assert "C5=Budget approval" in block.text
    assert "E6==1+1" in block.text


def test_extract_eml_preserves_headers_and_plain_body(tmp_path: Path) -> None:
    """Verify EML extraction keeps important headers and body text.

    Inputs:
        tmp_path: Temporary directory for creating the email fixture.

    Outputs:
        None. Assertions confirm email metadata and extracted blocks.
    """

    email_path = tmp_path / "governance.eml"
    message = EmailMessage()
    message["Subject"] = "Board governance evidence"
    message["From"] = "chair@example.com"
    message["To"] = "audit@example.com"
    message["Cc"] = "cfo@example.com"
    message["Date"] = "Mon, 1 Jun 2026 10:00:00 +0000"
    message.set_content("Board minutes confirm annual strategy oversight.")
    _write_email(email_path, message)

    extracted = extract_eml(email_path)

    assert extracted.file_name == "governance.eml"
    assert extracted.document_type == "eml"
    assert extracted.blocks[0].block_id == "eml-block-0001"
    assert extracted.blocks[0].block_type == "headers"
    assert "Subject: Board governance evidence" in extracted.blocks[0].text
    assert extracted.blocks[1].block_type == "body"
    assert "annual strategy oversight" in extracted.blocks[1].text
    assert extracted.metadata["has_attachments"] is False
    assert extracted.metadata["attachment_count"] == 0


def test_extract_eml_falls_back_to_sanitized_html_body(tmp_path: Path) -> None:
    """Verify HTML-only emails produce readable tag-free body text.

    Inputs:
        tmp_path: Temporary directory for creating the email fixture.

    Outputs:
        None. Assertions confirm HTML tags are not embedded.
    """

    email_path = tmp_path / "html-only.eml"
    message = EmailMessage()
    message["Subject"] = "HTML governance evidence"
    message.set_content(
        "<html><body><h1>Governance</h1><p>Board &amp; management review risks.</p></body></html>",
        subtype="html",
    )
    _write_email(email_path, message)

    extracted = extract_eml(email_path)
    body = next(block for block in extracted.blocks if block.block_type == "body")

    assert "Governance" in body.text
    assert "Board & management review risks." in body.text
    assert "<p>" not in body.text
    assert "</html>" not in body.text


def test_extract_eml_prefers_plain_text_in_multipart_email(tmp_path: Path) -> None:
    """Verify multipart emails prefer plain text over HTML alternatives.

    Inputs:
        tmp_path: Temporary directory for creating the email fixture.

    Outputs:
        None. Assertions confirm the selected body is the plain part.
    """

    email_path = tmp_path / "multipart.eml"
    message = EmailMessage()
    message["Subject"] = "Multipart governance evidence"
    message.set_content("Plain body confirms ownership reviews.")
    message.add_alternative(
        "<html><body>HTML body should not be selected.</body></html>",
        subtype="html",
    )
    _write_email(email_path, message)

    extracted = extract_eml(email_path)
    body = next(block for block in extracted.blocks if block.block_type == "body")

    assert "Plain body confirms ownership reviews." in body.text
    assert "HTML body should not be selected." not in body.text


def test_extract_eml_records_skipped_attachments_without_indexing_content(
    tmp_path: Path,
) -> None:
    """Verify email attachments are captured as metadata but not embedded.

    Inputs:
        tmp_path: Temporary directory for creating the email fixture.

    Outputs:
        None. Assertions confirm attachment content is skipped in v1.
    """

    email_path = tmp_path / "with-attachment.eml"
    message = EmailMessage()
    message["Subject"] = "Evidence with attachment"
    message.set_content("Email body is the only indexed email content.")
    message.add_attachment(
        b"attachment text must not enter retrieval",
        maintype="application",
        subtype="pdf",
        filename="policy.pdf",
    )
    _write_email(email_path, message)

    extracted = extract_eml(email_path)
    combined_text = "\n".join(block.text for block in extracted.blocks)

    assert extracted.metadata["has_attachments"] is True
    assert extracted.metadata["attachment_count"] == 1
    assert extracted.metadata["skipped_attachments"] == [
        {"file_name": "policy.pdf", "content_type": "application/pdf"}
    ]
    assert extracted.metadata["warnings"] == [
        "Email attachments are recorded as metadata but not extracted."
    ]
    assert "attachment text must not enter retrieval" not in combined_text


def _write_text_pdf(path: Path, lines: list[str]) -> None:
    """Write a simple text PDF for extraction tests.

    Inputs:
        path: Destination PDF path.
        lines: Text lines to place on the first PDF page.

    Outputs:
        None. The PDF is written to disk.
    """

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    for line in lines:
        pdf.cell(0, 10, text=line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(path)


def test_extract_pdf_returns_page_blocks_with_quality_metadata(
    tmp_path: Path,
) -> None:
    """Verify PDF extraction returns traceable page-level text blocks.

    Inputs:
        tmp_path: Temporary directory for creating a PDF fixture.

    Outputs:
        None. Assertions confirm page metadata, text, and quality signals.
    """

    pdf_path = tmp_path / "financial-report.pdf"
    _write_text_pdf(
        pdf_path,
        [
            "Strategic plan includes decarbonization objectives.",
            "Board reviews progress annually.",
        ],
    )

    extracted = extract_pdf(pdf_path)

    assert extracted.file_name == "financial-report.pdf"
    assert extracted.document_type == "pdf"
    assert extracted.blocks[0].block_id == "pdf-page-0001"
    assert extracted.blocks[0].page == 1
    assert "decarbonization objectives" in extracted.blocks[0].text
    assert extracted.metadata["page_count"] == 1
    assert extracted.metadata["extracted_page_count"] == 1
    assert extracted.metadata["empty_page_count"] == 0
    assert extracted.metadata["low_text_density"] is False


def test_extract_pdf_flags_low_text_density(tmp_path: Path) -> None:
    """Verify sparse PDFs are marked as low text density.

    Inputs:
        tmp_path: Temporary directory for creating a sparse PDF fixture.

    Outputs:
        None. Assertions confirm extraction quality metadata flags sparse text.
    """

    pdf_path = tmp_path / "scanned-like.pdf"
    _write_text_pdf(pdf_path, ["x"])

    extracted = extract_pdf(pdf_path)

    assert extracted.metadata["page_count"] == 1
    assert extracted.metadata["low_text_density"] is True
