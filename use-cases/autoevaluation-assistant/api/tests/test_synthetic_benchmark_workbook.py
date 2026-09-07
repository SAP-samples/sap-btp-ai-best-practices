"""Tests for the tracked deterministic synthetic benchmark workbook."""

from __future__ import annotations

import hashlib
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.models.assessment import AssessmentQuestion
from app.services import question_scoring as shared_question_scoring
from app.services.benchmark_import import parse_benchmark_workbook
from app.services.benchmark_import.constants import REQUIRED_COLUMNS
from app.services.customer_class_scope import max_allowed_level
from app.services.framework_importer import (
    load_framework_seed,
    load_italian_framework_translations,
)
from scripts import generate_synthetic_benchmark_workbook as benchmark_generator

generate_workbook = benchmark_generator.generate_workbook


def _italian_framework(repo_root: Path) -> list[AssessmentQuestion]:
    """Load canonical questions with Italian source-comparison text.

    Inputs:
        repo_root: Repository root containing framework seed files.

    Outputs:
        list[AssessmentQuestion]: Canonical IDs/dimensions with Italian labels.
    """

    seed = load_framework_seed(
        workbook_path=repo_root / "data" / "sanitized" / "assessment_framework.xlsx",
        explanations_path=repo_root / "data" / "sanitized" / "assessment_question_explanations.csv",
    )
    translations = load_italian_framework_translations(
        italian_dir=repo_root / "data" / "sanitized" / "IT",
        questions=seed.questions,
    )
    localized: list[AssessmentQuestion] = []
    for question in seed.questions:
        question_translation = translations.question_translations[question.question_id]
        localized.append(
            question.model_copy(
                update={
                    "section": question_translation.section,
                    "question": question_translation.question,
                    "answer_items": [
                        item.model_copy(
                            update={
                                "text": translations.answer_item_translations[
                                    item.answer_item_id
                                ]
                            }
                        )
                        for item in question.answer_items
                    ],
                }
            )
        )
    return localized


def _read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    """Read exact headers and data rows from a benchmark fixture.

    Inputs:
        path: XLSX workbook path containing the ``Estrazione`` sheet.

    Outputs:
        tuple: Ordered header list and row mappings for all non-header rows.
    """

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Estrazione"]
        headers = [
            str(value)
            for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        ]
        rows = [
            dict(zip(headers, values, strict=True))
            for values in sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in values)
        ]
        return headers, rows
    finally:
        workbook.close()


def _tracked_fixture(repo_root: Path) -> Path:
    """Return the required tracked synthetic workbook path.

    Inputs:
        repo_root: Repository root containing backend tests.

    Outputs:
        Path: Canonical fixture destination under ``api/tests/fixtures``.
    """

    return (
        repo_root
        / "api"
        / "tests"
        / "fixtures"
        / "benchmark"
        / "DB_extraction_clustering_NACE_synthetic.xlsx"
    )


def test_tracked_synthetic_fixture_has_complete_exact_source_shape(
    repo_root: Path,
) -> None:
    """Verify the tracked fixture has 24 complete class-3 questionnaires.

    Inputs:
        repo_root: Repository root used to locate the tracked fixture.

    Outputs:
        None. Assertions cover exact columns, counts, IDs, and selection bounds.
    """

    fixture = _tracked_fixture(repo_root)
    assert fixture.is_file(), "Required tracked synthetic benchmark fixture is missing"
    headers, rows = _read_rows(fixture)

    assert headers == list(REQUIRED_COLUMNS)
    assert len(rows) == 24 * 551
    assert {row["Classe"] for row in rows} == {"3"}
    company_ids = {row["ID Impresa"] for row in rows}
    questionnaire_ids = {row["ID Questionario"] for row in rows}
    question_ids = {row["ID Domanda"] for row in rows}
    assert len(company_ids) == 24
    assert len(questionnaire_ids) == 24
    assert len(question_ids) == 38
    assert all(str(value).startswith("SYN-COMPANY-") for value in company_ids)
    assert all(
        str(value).startswith("SYN-QUESTIONNAIRE-")
        for value in questionnaire_ids
    )
    assert all(str(row["ID Risposta"]).startswith("SYN-R-") for row in rows)
    assert set(Counter(row["ID Questionario"] for row in rows).values()) == {551}
    assert all(
        str(row["Settore Operativo (NACE) 1"]).startswith("Synthetic ")
        for row in rows
    )
    assert all(
        int(row["Livello"])
        <= max_allowed_level(str(row["ID Domanda"]), "class_3")
        for row in rows
        if row["Valore Risposta"] == "Si"
    )
    selected_counts = Counter(
        row["ID Questionario"]
        for row in rows
        if row["Valore Risposta"] == "Si"
    )
    assert len(set(selected_counts.values())) > 1


def test_tracked_fixture_reimports_without_score_reconciliation_errors(
    repo_root: Path,
) -> None:
    """Verify the synthetic source passes the real pure importer end to end.

    Inputs:
        repo_root: Repository root containing fixture and framework seed files.

    Outputs:
        None. Parsed counts match and supplied dimension scores reconcile.
    """

    fixture = _tracked_fixture(repo_root)
    assert fixture.is_file(), "Required tracked synthetic benchmark fixture is missing"

    dataset = parse_benchmark_workbook(
        fixture.read_bytes(),
        fixture.name,
        _italian_framework(repo_root),
    )

    assert dataset.summary.success is True
    assert dataset.summary.row_count == 13_224
    assert dataset.summary.accepted_count == 13_224
    assert dataset.summary.rejected_count == 0
    assert dataset.summary.company_count == 24
    assert dataset.summary.questionnaire_count == 24
    assert dataset.summary.question_count == 38
    assert len(dataset.scores) == 24 * (38 + 7 + 1)
    assert "score_reconciliation" not in {
        warning.code for warning in dataset.summary.warnings
    }


def test_generator_is_byte_deterministic_and_does_not_copy_profiles(
    repo_root: Path,
    tmp_path: Path,
) -> None:
    """Verify identical seed inputs produce identical fully synthetic workbooks.

    Inputs:
        repo_root: Repository root containing the real or tracked template.
        tmp_path: Temporary directory for two generated outputs.

    Outputs:
        None. Byte hashes match and real company/profile identities do not leak.
    """

    real_template = (
        repo_root / "data" / "report data" / "DB_extraction_clustering_NACE.xlsx"
    )
    template = real_template if real_template.exists() else _tracked_fixture(repo_root)
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"

    generate_workbook(template, first, companies=3, seed=42)
    generate_workbook(template, second, companies=3, seed=42)

    assert hashlib.sha256(first.read_bytes()).hexdigest() == hashlib.sha256(
        second.read_bytes()
    ).hexdigest()

    formatted_workbook = load_workbook(first, read_only=False, data_only=True)
    try:
        formatted_sheet = formatted_workbook["Estrazione"]
        assert formatted_sheet.freeze_panes == "A2"
        assert formatted_sheet.auto_filter.ref == "A1:AD1654"
        assert all(cell.font.bold for cell in formatted_sheet[1])
        assert all(cell.fill.fill_type == "solid" for cell in formatted_sheet[1])
        widths = {
            column: formatted_sheet.column_dimensions[column].width
            for column in formatted_sheet.column_dimensions
        }
        assert set(widths) == {
            chr(ord("A") + offset) if offset < 26 else f"A{chr(ord('A') + offset - 26)}"
            for offset in range(len(REQUIRED_COLUMNS))
        }
        assert all(width is not None and 10 <= width <= 60 for width in widths.values())
        assert widths["W"] >= 40
        assert widths["AB"] >= 40
        assert formatted_sheet["W2"].alignment.wrap_text is True
        assert formatted_sheet["AB2"].alignment.wrap_text is True
    finally:
        formatted_workbook.close()

    headers, generated_rows = _read_rows(first)
    assert headers == list(REQUIRED_COLUMNS)
    assert len(generated_rows) == 3 * 551
    assert all(
        str(row["ID Impresa"]).startswith("SYN-COMPANY-")
        and str(row["ID Questionario"]).startswith("SYN-QUESTIONNAIRE-")
        and str(row["ID Risposta"]).startswith("SYN-R-")
        for row in generated_rows
    )

    if real_template.exists():
        _real_headers, real_rows = _read_rows(real_template)
        real_company_ids = {str(row["ID Impresa"]) for row in real_rows}
        real_questionnaire_ids = {str(row["ID Questionario"]) for row in real_rows}
        real_profiles = {
            (
                row["Fatturato"],
                row["Nr Dipendenti"],
                row["Settore Operativo (NACE) 1"],
                row["Settore Operativo (NACE) 2"],
                row["Settore Operativo (NACE) 3"],
                row["Forma Giuridica"],
            )
            for row in real_rows
        }
        generated_profiles = {
            (
                row["Fatturato"],
                row["Nr Dipendenti"],
                row["Settore Operativo (NACE) 1"],
                row["Settore Operativo (NACE) 2"],
                row["Settore Operativo (NACE) 3"],
                row["Forma Giuridica"],
            )
            for row in generated_rows
        }
        assert real_company_ids.isdisjoint(
            {str(row["ID Impresa"]) for row in generated_rows}
        )
        assert real_questionnaire_ids.isdisjoint(
            {str(row["ID Questionario"]) for row in generated_rows}
        )
        assert real_profiles.isdisjoint(generated_profiles)


def test_generator_reuses_shared_maturity_primitive_and_reconciles(
    repo_root: Path,
    tmp_path: Path,
    monkeypatch: Any,
) -> None:
    """Verify generated audit scores call the shared primitive and match import.

    Inputs:
        repo_root: Repository root containing the tracked benchmark template.
        tmp_path: Temporary output directory for one generated workbook.
        monkeypatch: Pytest fixture used to observe the generator call boundary.

    Outputs:
        None. Every class-3 question is scored through the common lower-level
        primitive and the resulting workbook has no reconciliation warnings.
    """

    primitive = getattr(shared_question_scoring, "calculate_maturity_score", None)
    assert callable(primitive), "Shared lower-level maturity primitive is missing"
    assert getattr(benchmark_generator, "calculate_maturity_score", None) is primitive

    calls: list[tuple[tuple[Any, ...], dict[str, Any]]] = []

    def recording_primitive(*args: Any, **kwargs: Any) -> Any:
        """Record one generator score request and delegate to the real primitive."""

        calls.append((args, kwargs))
        return primitive(*args, **kwargs)

    monkeypatch.setattr(
        benchmark_generator,
        "calculate_maturity_score",
        recording_primitive,
    )
    output = tmp_path / "shared-scoring.xlsx"
    generate_workbook(
        _tracked_fixture(repo_root),
        output,
        companies=1,
        seed=42,
    )
    dataset = parse_benchmark_workbook(
        output.read_bytes(),
        output.name,
        _italian_framework(repo_root),
    )

    assert len(calls) == 38
    assert dataset.summary.success is True
    assert "score_reconciliation" not in {
        warning.code for warning in dataset.summary.warnings
    }


def test_synthetic_dimension_scores_are_constant_within_questionnaire_scope(
    repo_root: Path,
) -> None:
    """Verify each generated questionnaire/dimension repeats one audit score.

    Inputs:
        repo_root: Repository root used to locate the tracked fixture.

    Outputs:
        None. Every derived scope has exactly one supplied integer score.
    """

    _headers, rows = _read_rows(_tracked_fixture(repo_root))
    scores: dict[tuple[str, str], set[int]] = defaultdict(set)
    for row in rows:
        scores[(str(row["ID Questionario"]), str(row["Dimensione"]))].add(
            int(row["Dimensione Score"])
        )

    assert len(scores) == 24 * 7
    assert all(len(values) == 1 for values in scores.values())


def test_generator_script_runs_directly_from_api_directory(repo_root: Path) -> None:
    """Verify documented direct Python invocation resolves application imports.

    Inputs:
        repo_root: Repository root containing the backend script and package.

    Outputs:
        None. ``--help`` exits successfully without touching a workbook.
    """

    result = subprocess.run(
        [
            sys.executable,
            str(repo_root / "api" / "scripts" / "generate_synthetic_benchmark_workbook.py"),
            "--help",
        ],
        cwd=repo_root / "api",
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    assert "--companies" in result.stdout
