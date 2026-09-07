"""Tests for the per-question AI review LangGraph workflow."""

import base64
from pathlib import Path
from typing import Any

from docx import Document
from langchain_core.messages import AIMessage
from openpyxl import Workbook
import pytest

from app.models.ai_review import (
    AnswerItemDecision,
    LevelReviewResult,
    QuestionReviewResult,
    RetrievedEvidenceDecision,
    RetrievalRoundAssessment,
)
from app.models.assessment import AnswerItem, AssessmentQuestion
from app.services.ai_review_graph import (
    QuestionReviewInput,
    run_rag_question_review,
    run_question_review,
    run_question_review_with_routing,
    run_task_from_hana,
)
from app.services.batch_question_react_graph import run_manual_question_react_review
from app.services.document_extractors import ExtractedBlock, ExtractedDocument
from app.services.genai_responses import ReviewContextLimitError


class FakeReviewClient:
    """Deterministic review client used to avoid live Gen AI Hub calls.

    Inputs:
        None. The fake stores the latest input messages passed by the graph.

    Outputs:
        A test double exposing ``review_question`` with the same boundary as
        ``GenAiReviewClient``.
    """

    def __init__(self) -> None:
        """Initialize an empty input capture slot.

        Inputs:
            None.

        Outputs:
            None. ``last_input`` is ready to record graph model calls.
        """

        self.last_input: list[dict[str, Any]] | None = None
        self.last_language: str | None = None

    def review_question(
        self,
        input_messages: list[dict[str, Any]],
        language: str = "en",
    ) -> QuestionReviewResult:
        """Return a deterministic question review result for tests.

        Inputs:
            input_messages: Prompt messages assembled by the LangGraph workflow.
            language: Review output language requested by the worker task.

        Outputs:
            QuestionReviewResult: Stable result selecting the expected L2 item.
        """

        self.last_input = input_messages
        self.last_language = language
        return QuestionReviewResult(
            question_id="Q.STR.01.01",
            model="fake-gpt-5.4",
            overall_status="supported",
            highest_supported_level=2,
            current_selected_answer_item_ids=["Q.STR.01.01-L1-001"],
            verified_selected_answer_item_ids=["Q.STR.01.01-L2-001"],
        )


def test_run_question_review_extracts_docx_and_calls_client(tmp_path: Path) -> None:
    """Verify DOCX evidence is extracted into the prompt sent to the client.

    Inputs:
        tmp_path: Pytest temporary directory used to create a DOCX evidence file.

    Outputs:
        None. Assertions validate the fake client result and prompt assembly.
    """

    document_path = tmp_path / "strategy-evidence.docx"
    document = Document()
    document.add_heading("Strategy Evidence", level=1)
    document.add_paragraph("The company has a documented transition roadmap.")
    document.save(document_path)
    review_client = FakeReviewClient()

    review_input = QuestionReviewInput(
        task_id="task-1",
        question=AssessmentQuestion(
            question_id="Q.STR.01.01",
            dimension="Strategy",
            section="Governance",
            question="Is there a transition strategy?",
            explanation="Review whether a strategy is documented.",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.01.01-L2-001",
                    question_id="Q.STR.01.01",
                    level=2,
                    item_index=1,
                    text="Transition roadmap is documented.",
                )
            ],
        ),
        current_selected_answer_item_ids=["Q.STR.01.01-L1-001"],
        attachment_paths=[document_path],
    )

    result = run_question_review(review_input, review_client=review_client)

    assert result.verified_selected_answer_item_ids == ["Q.STR.01.01-L2-001"]
    assert review_client.last_input is not None
    prompt_text = "\n".join(
        block["text"]
        for message in review_client.last_input
        for block in message["content"]
        if block["type"] == "input_text"
    )
    assert "Strategy Evidence" in prompt_text
    assert "documented transition roadmap" in prompt_text
    assert "Q.STR.01.01-L2-001" in prompt_text


def test_run_question_review_keeps_pdf_as_raw_input_file_block(
    tmp_path: Path,
) -> None:
    """Verify direct graph review still sends PDFs as raw Responses files.

    Inputs:
        tmp_path: Pytest temporary directory used to create a PDF attachment.

    Outputs:
        None. Assertions confirm Task 9 does not change the legacy direct seam.
    """

    pdf_path = tmp_path / "strategy-evidence.pdf"
    pdf_bytes = b"%PDF-1.7 direct review evidence"
    pdf_path.write_bytes(pdf_bytes)
    review_client = FakeReviewClient()

    review_input = QuestionReviewInput(
        task_id="task-pdf-direct",
        question=AssessmentQuestion(
            question_id="Q.STR.01.01",
            dimension="Strategy",
            section="Governance",
            question="Is there a transition strategy?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.01.01-L1-001",
                    question_id="Q.STR.01.01",
                    level=1,
                    item_index=1,
                    text="A strategy exists in evidence.",
                )
            ],
        ),
        attachment_paths=[pdf_path],
    )

    run_question_review(review_input, review_client=review_client)

    assert review_client.last_input is not None
    file_blocks = [
        block
        for block in review_client.last_input[0]["content"]
        if block["type"] == "input_file"
    ]
    assert len(file_blocks) == 1
    assert file_blocks[0]["filename"] == "strategy-evidence.pdf"
    assert file_blocks[0]["file_data"] == (
        "data:application/pdf;base64,"
        f"{base64.b64encode(pdf_bytes).decode('ascii')}"
    )


def test_run_question_review_extracts_xlsm_evidence(tmp_path: Path) -> None:
    """Verify XLSM evidence is extracted like XLSX content.

    Inputs:
        tmp_path: Pytest temporary directory used to create a workbook file.

    Outputs:
        None. Assertions validate the fake client received extracted sheet text.
    """

    workbook_path = tmp_path / "objectives-evidence.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Objectives"
    sheet["B3"] = "Quarterly strategy review"
    workbook.save(workbook_path)
    workbook.close()
    review_client = FakeReviewClient()

    review_input = QuestionReviewInput(
        task_id="task-2",
        question=AssessmentQuestion(
            question_id="Q.STR.01.01",
            dimension="Strategy",
            section="Governance",
            question="Is there a transition strategy?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.01.01-L2-001",
                    question_id="Q.STR.01.01",
                    level=2,
                    item_index=1,
                    text="Transition roadmap is documented.",
                )
            ],
        ),
        attachment_paths=[workbook_path],
    )

    run_question_review(review_input, review_client=review_client)

    assert review_client.last_input is not None
    prompt_text = "\n".join(
        block["text"]
        for message in review_client.last_input
        for block in message["content"]
        if block["type"] == "input_text"
    )
    assert "Objectives" in prompt_text
    assert "Quarterly strategy review" in prompt_text


@pytest.mark.parametrize(
    ("file_name", "mime_type", "image_bytes"),
    [
        ("factory-signage.png", "image/png", b"\x89PNG\r\n\x1a\nimage evidence"),
        ("permit-photo.jpeg", "image/jpeg", b"\xff\xd8\xff\xe0image evidence"),
    ],
)
def test_run_question_review_sends_images_as_responses_image_blocks(
    tmp_path: Path,
    file_name: str,
    mime_type: str,
    image_bytes: bytes,
) -> None:
    """Verify image evidence is sent to GPT-5.4 as Responses image blocks.

    Inputs:
        tmp_path: Pytest temporary directory used to create an image evidence
            file.
        file_name: Image evidence filename under test.
        mime_type: MIME type expected in the Responses API data URL.
        image_bytes: Binary image payload used to verify base64 encoding.

    Outputs:
        None. Assertions validate image metadata text and ``input_image`` block
        assembly without calling the live Gen AI Hub deployment.
    """

    image_path = tmp_path / file_name
    image_path.write_bytes(image_bytes)
    review_client = FakeReviewClient()

    review_input = QuestionReviewInput(
        task_id="task-image",
        question=AssessmentQuestion(
            question_id="Q.STR.01.01",
            dimension="Strategy",
            section="Governance",
            question="Is there a transition strategy?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.01.01-L1-001",
                    question_id="Q.STR.01.01",
                    level=1,
                    item_index=1,
                    text="The strategy is visible in company evidence.",
                )
            ],
        ),
        attachment_paths=[image_path],
    )

    run_question_review(review_input, review_client=review_client)

    assert review_client.last_input is not None
    content_blocks = review_client.last_input[0]["content"]
    image_metadata = [
        block["text"]
        for block in content_blocks
        if block["type"] == "input_text" and file_name in block["text"]
    ]
    image_blocks = [
        block for block in content_blocks if block["type"] == "input_image"
    ]

    assert image_metadata
    assert "visual evidence review" in image_metadata[0]
    assert len(image_blocks) == 1
    expected_data_url = (
        f"data:{mime_type};base64,{base64.b64encode(image_bytes).decode('ascii')}"
    )
    assert image_blocks[0]["image_url"] == expected_data_url


def test_run_question_review_passes_language_to_review_client(
    tmp_path: Path,
) -> None:
    """Verify the graph preserves the task language for model generation.

    Inputs:
        tmp_path: Pytest temporary directory reserved for attachment setup.

    Outputs:
        None. Assertions confirm Italian task language reaches the review
        client, which builds the final model instructions.
    """

    review_client = FakeReviewClient()
    review_input = QuestionReviewInput(
        task_id="task-3",
        language="it",
        question=AssessmentQuestion(
            question_id="Q.STR.01.01",
            dimension="Strategy",
            section="Pianificazione Strategica",
            question="La tua organizzazione è dotata di una strategia?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.01.01-L1-001",
                    question_id="Q.STR.01.01",
                    level=1,
                    item_index=1,
                    text="La strategia è documentata.",
                )
            ],
        ),
        attachment_paths=[],
    )

    run_question_review(review_input, review_client=review_client)

    assert review_client.last_language == "it"


def test_run_question_review_normalizes_verified_selections_from_decisions() -> None:
    """Verify selectable AI decisions become verified selections in catalog order.

    Inputs:
        None. The test injects a fake review client returning an intentionally
        inconsistent verified selection list and level decisions.

    Outputs:
        None. Assertions confirm the graph normalizes verified selections from
        keep-selected, select, and low-confidence decisions while excluding
        unsupported or unclear decisions.
    """

    class DecisionReviewClient:
        """Fake client returning decisions that need graph-level normalization.

        Inputs:
            None. The fake ignores the prompt and emits deterministic decisions.

        Outputs:
            A review client boundary compatible with ``run_question_review``.
        """

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Return a review result whose verified IDs must be recomputed.

            Inputs:
                input_messages: Prompt messages assembled by the graph.
                language: Requested output language, unused by this fake.

            Outputs:
                QuestionReviewResult: Result with selectable decisions and an
                incorrect pre-normalization verified ID list.
            """
            return QuestionReviewResult(
                question_id="Q.STR.03.01",
                model="fake-gpt-5.4",
                overall_status="partially_supported",
                highest_supported_level=3,
                current_selected_answer_item_ids=[],
                verified_selected_answer_item_ids=["wrong-id"],
                level_results=[
                    LevelReviewResult(
                        level=3,
                        level_status="supported",
                        level_reasoning=(
                            "Reasoning and evidence: objectives.xlsx sheet "
                            "Objectives supports top executive monitoring."
                        ),
                        answer_item_decisions=[
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L3-001",
                                decision="select",
                                confidence=0.88,
                                rationale="Direct evidence supports the item.",
                            )
                        ],
                    ),
                    LevelReviewResult(
                        level=1,
                        level_status="low_confidence",
                        level_reasoning=(
                            "Reasoning and evidence: objectives.xlsx sheet "
                            "Objectives implies annual objective definition."
                        ),
                        answer_item_decisions=[
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L1-001",
                                decision="low_confidence",
                                confidence=0.57,
                                rationale="Higher-level evidence implies this prerequisite.",
                            ),
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L1-002",
                                decision="unsupported",
                                confidence=0.93,
                                rationale="No evidence supports this item.",
                            ),
                        ],
                    ),
                    LevelReviewResult(
                        level=2,
                        level_status="partially_supported",
                        level_reasoning=(
                            "Reasoning and evidence: objectives.xlsx sheet "
                            "Objectives keeps monitoring selected."
                        ),
                        answer_item_decisions=[
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L2-002",
                                decision="unclear",
                                confidence=0.86,
                                rationale="The current selection lacks enough evidence.",
                            ),
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L2-001",
                                decision="keep_selected",
                                confidence=0.91,
                                rationale="The current selection remains supported.",
                            ),
                        ],
                    ),
                ],
            )

    review_input = QuestionReviewInput(
        task_id="task-normalize",
        question=AssessmentQuestion(
            question_id="Q.STR.03.01",
            dimension="Strategy",
            section="Objectives",
            question="Are objectives defined and monitored?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L1-001",
                    question_id="Q.STR.03.01",
                    level=1,
                    item_index=1,
                    text="Objectives are defined at least annually.",
                ),
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L1-002",
                    question_id="Q.STR.03.01",
                    level=1,
                    item_index=2,
                    text="A separate unsupported annual objective item.",
                ),
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L2-001",
                    question_id="Q.STR.03.01",
                    level=2,
                    item_index=1,
                    text="Objectives are monitored at least annually.",
                ),
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L2-002",
                    question_id="Q.STR.03.01",
                    level=2,
                    item_index=2,
                    text="Management defines response actions.",
                ),
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L3-001",
                    question_id="Q.STR.03.01",
                    level=3,
                    item_index=1,
                    text="Objectives are monitored by the Top Executive.",
                ),
            ],
        ),
    )

    result = run_question_review(
        review_input,
        review_client=DecisionReviewClient(),
    )

    assert result.verified_selected_answer_item_ids == [
        "Q.STR.03.01-L1-001",
        "Q.STR.03.01-L2-001",
        "Q.STR.03.01-L3-001",
    ]


def test_run_question_review_clips_results_to_scoped_answer_levels() -> None:
    """Verify model output above the scoped answer catalog cannot persist."""

    class ScopedReviewClient:
        """Fake review client returning one unsupported-by-scope high level."""

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Return L3 output for a question catalog scoped to L1-L2."""
            return QuestionReviewResult(
                question_id="Q.STR.03.01",
                model="fake-gpt-5.4",
                overall_status="supported",
                highest_supported_level=3,
                current_selected_answer_item_ids=[],
                verified_selected_answer_item_ids=["Q.STR.03.01-L3-001"],
                level_results=[
                    LevelReviewResult(
                        level=3,
                        level_status="supported",
                        level_reasoning="Top executive monitoring is claimed.",
                        answer_item_decisions=[
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L3-001",
                                decision="select",
                                confidence=0.91,
                                rationale="The fake model selected L3.",
                            )
                        ],
                    ),
                    LevelReviewResult(
                        level=2,
                        level_status="supported",
                        level_reasoning="Annual monitoring is supported.",
                        answer_item_decisions=[
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L2-001",
                                decision="select",
                                confidence=0.9,
                                rationale="The fake model selected L2.",
                            )
                        ],
                    ),
                ],
            )

    review_input = QuestionReviewInput(
        task_id="task-scoped",
        question=AssessmentQuestion(
            question_id="Q.STR.03.01",
            dimension="Strategy",
            section="Objectives",
            question="Are objectives defined and monitored?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L1-001",
                    question_id="Q.STR.03.01",
                    level=1,
                    item_index=1,
                    text="Objectives are defined.",
                ),
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L2-001",
                    question_id="Q.STR.03.01",
                    level=2,
                    item_index=1,
                    text="Objectives are monitored.",
                ),
            ],
        ),
    )

    result = run_question_review(review_input, review_client=ScopedReviewClient())

    assert result.highest_supported_level == 2
    assert [level.level for level in result.level_results] == [2]
    assert "Q.STR.03.01-L3-001" not in result.verified_selected_answer_item_ids


def test_run_question_review_filters_verified_ids_without_decision_rows() -> None:
    """Verify fallback verified IDs still stay inside the scoped catalog."""

    class MinimalScopedReviewClient:
        """Fake review client returning no answer decision rows."""

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Return an L3 verified ID for a question scoped to L1-L2."""
            return QuestionReviewResult(
                question_id="Q.STR.03.01",
                model="fake-gpt-5.4",
                overall_status="supported",
                highest_supported_level=3,
                current_selected_answer_item_ids=[],
                verified_selected_answer_item_ids=["Q.STR.03.01-L3-001"],
                level_results=[
                    LevelReviewResult(
                        level=3,
                        level_status="supported",
                        level_reasoning="Top executive monitoring is claimed.",
                        answer_item_decisions=[],
                    ),
                    LevelReviewResult(
                        level=2,
                        level_status="supported",
                        level_reasoning="Annual monitoring is supported.",
                        answer_item_decisions=[],
                    ),
                ],
            )

    review_input = QuestionReviewInput(
        task_id="task-scoped-no-decisions",
        question=AssessmentQuestion(
            question_id="Q.STR.03.01",
            dimension="Strategy",
            section="Objectives",
            question="Are objectives defined and monitored?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L1-001",
                    question_id="Q.STR.03.01",
                    level=1,
                    item_index=1,
                    text="Objectives are defined.",
                ),
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L2-001",
                    question_id="Q.STR.03.01",
                    level=2,
                    item_index=1,
                    text="Objectives are monitored.",
                ),
            ],
        ),
    )

    result = run_question_review(
        review_input,
        review_client=MinimalScopedReviewClient(),
    )

    assert result.highest_supported_level == 2
    assert [level.level for level in result.level_results] == [2]
    assert result.verified_selected_answer_item_ids == []


def test_run_question_review_clears_verified_selections_when_all_decisions_are_out_of_scope() -> None:
    """Verify clipped-only decision rows clear stale verified answer IDs."""

    class OutOfScopeOnlyReviewClient:
        """Fake review client returning only decisions above the scoped catalog."""

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Return one L3 decision for a question catalog scoped to L1-L2."""
            return QuestionReviewResult(
                question_id="Q.STR.03.01",
                model="fake-gpt-5.4",
                overall_status="supported",
                highest_supported_level=3,
                current_selected_answer_item_ids=[],
                verified_selected_answer_item_ids=["Q.STR.03.01-L3-001"],
                level_results=[
                    LevelReviewResult(
                        level=3,
                        level_status="supported",
                        level_reasoning="Top executive monitoring is claimed.",
                        answer_item_decisions=[
                            AnswerItemDecision(
                                answer_item_id="Q.STR.03.01-L3-001",
                                decision="select",
                                confidence=0.91,
                                rationale="The fake model selected L3.",
                            )
                        ],
                    )
                ],
            )

    review_input = QuestionReviewInput(
        task_id="task-scoped-empty",
        question=AssessmentQuestion(
            question_id="Q.STR.03.01",
            dimension="Strategy",
            section="Objectives",
            question="Are objectives defined and monitored?",
            answer_items=[
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L1-001",
                    question_id="Q.STR.03.01",
                    level=1,
                    item_index=1,
                    text="Objectives are defined.",
                ),
                AnswerItem(
                    answer_item_id="Q.STR.03.01-L2-001",
                    question_id="Q.STR.03.01",
                    level=2,
                    item_index=1,
                    text="Objectives are monitored.",
                ),
            ],
        ),
    )

    result = run_question_review(
        review_input,
        review_client=OutOfScopeOnlyReviewClient(),
    )

    assert result.highest_supported_level == 2
    assert result.level_results == []
    assert result.verified_selected_answer_item_ids == []


def test_run_task_from_hana_preserves_original_attachment_file_names(
    monkeypatch: Any,
) -> None:
    """Verify worker temp files keep uploaded filenames visible to the model.

    Inputs:
        monkeypatch: Pytest helper used to replace the graph runner and capture
        the assembled ``QuestionReviewInput``.

    Outputs:
        None. Assertions confirm duplicate upload names are isolated by
        directory while each temporary file basename remains unchanged.
    """

    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L2-001",
                question_id="Q.STR.01.01",
                level=2,
                item_index=1,
                text="Transition roadmap is documented.",
            )
        ],
    )

    class RepositoryStub:
        """Minimal repository exposing framework questions for the worker path.

        Inputs:
            None. The stub closes over the test question.

        Outputs:
            Repository-like object with ``list_questions`` for
            ``run_task_from_hana``.
        """

        def list_questions(self, dimension: str) -> list[AssessmentQuestion]:
            """Return the single test question for the requested dimension.

            Inputs:
                dimension: Dimension requested by the worker graph adapter.

            Outputs:
                list[AssessmentQuestion]: Single matching Strategy question.
            """

            assert dimension == "Strategy"
            return [question]

        def save_attachment_extraction(self, **kwargs: Any) -> str:
            """Accept extraction persistence calls from the worker adapter.

            Inputs:
                **kwargs: Attachment extraction metadata.

            Outputs:
                str: Fake extraction identifier.
            """

            captured_calls.append({"extraction": kwargs})
            return "extraction-1"

    captured_calls: list[dict[str, Any]] = []

    def fake_extract_evidence_files(paths: list[Path]) -> list[ExtractedDocument]:
        """Capture worker temp paths and skip real PDF parsing in this test.

        Inputs:
            paths: Temporary attachment paths written by the worker adapter.

        Outputs:
            list[ExtractedDocument]: Two extracted documents with duplicate
            basenames so attachment ID order can be verified.
        """

        captured_calls.append({"attachment_paths": list(paths)})
        return [
            ExtractedDocument(
                file_name=path.name,
                document_type="pdf",
                blocks=[
                    ExtractedBlock(
                        block_id=f"pdf-page-{index:04d}",
                        block_type="page",
                        text=f"Evidence text from duplicate attachment {index}.",
                        page=1,
                    )
                ],
            )
            for index, path in enumerate(paths, start=1)
        ]

    def fake_run_manual_question_react_review(
        review_input: QuestionReviewInput,
        extracted_documents: list[ExtractedDocument],
        attachment_id_by_file: dict[str, str],
        repository: Any,
        worker_id: str = "",
        review_client: Any | None = None,
        embedding_client: Any | None = None,
        llm: Any | None = None,
        attachment_ids_by_document: list[str] | None = None,
    ) -> QuestionReviewResult:
        """Capture manual ReAct RAG worker input and return a result."""

        captured_calls.append(
            {
                "review_input": review_input,
                "extracted_documents": extracted_documents,
                "attachment_id_by_file": attachment_id_by_file,
                "repository": repository,
                "worker_id": worker_id,
                "attachment_ids_by_document": attachment_ids_by_document,
                "review_client": review_client,
                "embedding_client": embedding_client,
                "llm": llm,
            }
        )
        return QuestionReviewResult(
            question_id=review_input.question.question_id,
            model="fake-gpt-5.4",
            overall_status="supported",
            highest_supported_level=1,
            current_selected_answer_item_ids=[],
            verified_selected_answer_item_ids=[],
        )

    monkeypatch.setattr(
        "app.services.ai_review_graph.extract_evidence_files",
        fake_extract_evidence_files,
    )
    monkeypatch.setattr(
        "app.services.ai_review_graph.run_question_review_with_routing",
        lambda *args, **kwargs: (_ for _ in ()).throw(
            AssertionError("manual worker should use ReAct RAG, not direct routing")
        ),
    )
    monkeypatch.setattr(
        "app.services.batch_question_react_graph.run_manual_question_react_review",
        fake_run_manual_question_react_review,
        raising=False,
    )

    result = run_task_from_hana(
        repository=RepositoryStub(),
        task={
            "task_id": "task-1",
            "dimension": "Strategy",
            "question_id": "Q.STR.01.01",
            "current_selected_answer_item_ids": [],
        },
        attachments=[
            {
                "file_name": "strategy-evidence.pdf",
                "content": b"%PDF-1.7",
            },
            {
                "file_name": "nested/path/strategy-evidence.pdf",
                "content": b"%PDF-1.7 duplicate name",
            },
        ],
    )

    assert result.question_id == "Q.STR.01.01"
    react_call = captured_calls[-1]
    attachment_paths = react_call["review_input"].attachment_paths
    assert [path.name for path in attachment_paths] == [
        "strategy-evidence.pdf",
        "strategy-evidence.pdf",
    ]
    assert attachment_paths[0].parent != attachment_paths[1].parent
    assert react_call["attachment_id_by_file"] == {
        "strategy-evidence.pdf": "attachment-001"
    }
    assert react_call["attachment_ids_by_document"] == [
        "attachment-001",
        "attachment-002",
    ]
    extraction_calls = [
        call["extraction"] for call in captured_calls if "extraction" in call
    ]
    assert [call["attachment_id"] for call in extraction_calls] == [
        "attachment-001",
        "attachment-002",
    ]
    assert [call["extracted"].file_name for call in extraction_calls] == [
        "strategy-evidence.pdf",
        "strategy-evidence.pdf",
    ]


def test_run_task_from_hana_publishes_extraction_progress(
    monkeypatch: Any,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Verify manual workers expose document extraction progress for polling.

    Inputs:
        monkeypatch: Pytest helper used to replace extraction and RAG calls.

    Outputs:
        None. Assertions confirm the manual task status is updated before the
        potentially slow extraction step and before handoff to ReAct RAG.
    """

    question = AssessmentQuestion(
        question_id="Q.STR.03.01",
        dimension="Strategy",
        section="Objectives",
        question="Are objectives defined and monitored?",
        answer_items=[],
    )

    class RepositoryStub:
        """Repository stub capturing manual progress and extraction audit rows."""

        def __init__(self) -> None:
            """Initialize empty progress and extraction stores."""

            self.progress_updates: list[dict[str, Any]] = []
            self.extractions: list[dict[str, Any]] = []
            self.commits = 0

        def list_questions(self, dimension: str) -> list[AssessmentQuestion]:
            """Return the fixture question for the requested dimension."""

            assert dimension == "Strategy"
            return [question]

        def update_question_progress(self, **kwargs: Any) -> None:
            """Record one manual progress update."""

            self.progress_updates.append(kwargs)

        def save_attachment_extraction(self, **kwargs: Any) -> None:
            """Record one extraction audit payload."""

            self.extractions.append(kwargs)

        def commit(self) -> None:
            """Record that progress was committed for polling visibility."""

            self.commits += 1

    def fake_extract_evidence_files(paths: list[Path]) -> list[ExtractedDocument]:
        """Return one extracted document for the uploaded PDF path."""

        assert [path.name for path in paths] == ["objectives.pdf"]
        return [
            ExtractedDocument(
                file_name="objectives.pdf",
                document_type="pdf",
                blocks=[
                    ExtractedBlock(
                        block_id="pdf-page-0001",
                        block_type="page",
                        text="Objectives are monitored quarterly.",
                        page=1,
                    )
                ],
            )
        ]

    def fake_run_manual_question_react_review(**kwargs: Any) -> QuestionReviewResult:
        """Return a deterministic result after progress handoff."""

        assert kwargs["worker_id"] == "worker-manual"
        return QuestionReviewResult(
            question_id="Q.STR.03.01",
            model="fake-gpt-5.4",
            overall_status="supported",
            highest_supported_level=3,
            current_selected_answer_item_ids=[],
            verified_selected_answer_item_ids=[],
        )

    monkeypatch.setattr(
        "app.services.ai_review_graph.extract_evidence_files",
        fake_extract_evidence_files,
    )
    monkeypatch.setattr(
        "app.services.batch_question_react_graph.run_manual_question_react_review",
        fake_run_manual_question_react_review,
        raising=False,
    )

    caplog.set_level("INFO", logger="app.services.ai_review_graph")
    repository = RepositoryStub()
    run_task_from_hana(
        repository=repository,
        task={
            "task_id": "task-progress",
            "job_id": "job-progress",
            "dimension": "Strategy",
            "question_id": "Q.STR.03.01",
            "current_selected_answer_item_ids": [],
            "lease_owner": "worker-manual",
        },
        attachments=[
            {
                "attachment_id": "attachment-001",
                "file_name": "objectives.pdf",
                "content": b"%PDF-1.7",
            }
        ],
    )

    assert repository.progress_updates[0]["status"] == "extracting_documents"
    assert "Extracting text from 1 uploaded document" in repository.progress_updates[0][
        "progress_message"
    ]
    assert repository.progress_updates[-1]["status"] == "embedding_documents"
    assert "Creating evidence chunks" in repository.progress_updates[-1][
        "progress_message"
    ]
    assert repository.commits >= 2
    assert any(
        "manual_ai_review_stage_timing stage=extract_documents" in record.message
        and "duration_ms=" in record.message
        for record in caplog.records
    )


def test_run_rag_question_review_preserves_duplicate_attachment_ids() -> None:
    """Verify RAG chunk persistence keeps duplicate-basename attachment IDs.

    Inputs:
        None. The test uses fake repository, embedding, and review boundaries.

    Outputs:
        None. Assertions validate chunk rows point to the ordered source
        attachments rather than a collapsed filename map.
    """

    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L2-001",
                question_id="Q.STR.01.01",
                level=2,
                item_index=1,
                text="Transition roadmap is documented.",
            )
        ],
    )
    extracted_documents = [
        ExtractedDocument(
            file_name="strategy-evidence.pdf",
            document_type="pdf",
            blocks=[
                ExtractedBlock(
                    block_id="pdf-page-0001",
                    block_type="page",
                    text="First attachment describes the transition strategy.",
                    page=1,
                )
            ],
        ),
        ExtractedDocument(
            file_name="strategy-evidence.pdf",
            document_type="pdf",
            blocks=[
                ExtractedBlock(
                    block_id="pdf-page-0001",
                    block_type="page",
                    text="Second attachment describes governance approvals.",
                    page=1,
                )
            ],
        ),
    ]

    class RepositoryStub:
        """Repository stub capturing chunk and retrieval persistence."""

        def __init__(self) -> None:
            """Initialize persistence capture lists."""

            self.saved_chunks: list[dict[str, Any]] = []
            self.retrieval_rounds: list[dict[str, Any]] = []

        def save_evidence_chunks(
            self,
            task_id: str,
            chunks: list[dict[str, Any]],
        ) -> None:
            """Record chunk rows submitted for HANA persistence.

            Inputs:
                task_id: Review task identifier.
                chunks: Chunk payloads with source attachment IDs.

            Outputs:
                None.
            """

            assert task_id == "task-duplicate-rag"
            self.saved_chunks.extend(chunks)

        def search_evidence_chunks(
            self,
            task_id: str,
            query_embedding: list[float],
            top_k: int,
        ) -> list[dict[str, Any]]:
            """Return saved chunks as vector-search candidates.

            Inputs:
                task_id: Review task identifier.
                query_embedding: Query embedding vector, unused.
                top_k: Maximum number of rows requested.

            Outputs:
                list[dict[str, Any]]: Retrieved chunk rows.
            """

            assert task_id == "task-duplicate-rag"
            return [
                {
                    "chunk_id": chunk["chunk_id"],
                    "attachment_id": chunk["attachment_id"],
                    "file_name": chunk["file_name"],
                    "document_type": chunk["document_type"],
                    "chunk_text": chunk["chunk_text"],
                    "location_json": chunk["location_json"],
                    "similarity_score": 0.9,
                }
                for chunk in self.saved_chunks[:top_k]
            ]

        def save_retrieval_round(self, **kwargs: Any) -> str:
            """Record one retrieval-round payload.

            Inputs:
                **kwargs: Retrieval round metadata.

            Outputs:
                str: Fake retrieval-round ID.
            """

            self.retrieval_rounds.append(kwargs)
            return f"round-{len(self.retrieval_rounds)}"

    class EmbeddingClientStub:
        """Embedding client stub returning deterministic vectors."""

        model_name = "fake-embedding"

        def embed_batch(self, texts: list[str]) -> list[list[float]]:
            """Return one stable vector per chunk text.

            Inputs:
                texts: Chunk texts to embed.

            Outputs:
                list[list[float]]: Fake embedding vectors.
            """

            return [[1.0, float(index)] for index, _text in enumerate(texts)]

        def embed_query(self, text: str) -> list[float]:
            """Return a stable query vector.

            Inputs:
                text: Retrieval query text.

            Outputs:
                list[float]: Fake query embedding.
            """

            return [0.5, 0.25]

    class RagReviewClientStub:
        """Review client stub for retrieval assessment and final review."""

        model = "fake-gpt-5.4"

        def __init__(self) -> None:
            """Initialize final prompt capture storage."""

            self.final_prompt_text = ""

        def assess_retrieved_evidence(
            self,
            input_messages: list[dict[str, object]],
            language: str = "en",
        ) -> RetrievalRoundAssessment:
            """Stop the RAG loop after one sufficient retrieval round.

            Inputs:
                input_messages: Retrieval assessment prompt.
                language: Requested review language.

            Outputs:
                RetrievalRoundAssessment: Sufficient assessment that accepts
                every retrieved duplicate-basename chunk.
            """

            return RetrievalRoundAssessment(
                accepted_evidence=[
                    RetrievedEvidenceDecision(
                        source_chunk_id=str(chunk["chunk_id"]),
                        relevance="support",
                        related_answer_item_ids=["Q.STR.01.01-L2-001"],
                        rationale="The chunk supports the strategy evidence.",
                        confidence=0.9,
                    )
                    for chunk in repository.saved_chunks
                ],
                rejected_evidence=[],
                evidence_gaps=[],
                refined_queries=[],
                sufficient_for_final_review=True,
                stop_reason="sufficient_evidence",
            )

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Return a deterministic final question review.

            Inputs:
                input_messages: Final compact RAG prompt.
                language: Requested review language.

            Outputs:
                QuestionReviewResult: Stable result for the test.
            """

            self.final_prompt_text = "\n".join(
                block["text"]
                for message in input_messages
                for block in message["content"]
                if block["type"] == "input_text"
            )
            return QuestionReviewResult(
                question_id=question.question_id,
                model=self.model,
                overall_status="supported",
                highest_supported_level=2,
                current_selected_answer_item_ids=[],
                verified_selected_answer_item_ids=[],
            )

    repository = RepositoryStub()
    review_client = RagReviewClientStub()

    result = run_rag_question_review(
        review_input=QuestionReviewInput(
            task_id="task-duplicate-rag",
            question=question,
            attachment_paths=[
                Path("/tmp/one/strategy-evidence.pdf"),
                Path("/tmp/two/strategy-evidence.pdf"),
            ],
        ),
        extracted_documents=extracted_documents,
        attachment_id_by_file={"strategy-evidence.pdf": "collapsed-id"},
        attachment_ids_by_document=["attachment-001", "attachment-002"],
        repository=repository,
        review_client=review_client,
        embedding_client=EmbeddingClientStub(),
    )

    assert result.overall_status == "supported"
    assert [chunk["attachment_id"] for chunk in repository.saved_chunks] == [
        "attachment-001",
        "attachment-002",
    ]
    assert len({chunk["chunk_id"] for chunk in repository.saved_chunks}) == 2
    assert "Attachment ID: attachment-001" in review_client.final_prompt_text
    assert "Attachment ID: attachment-002" in review_client.final_prompt_text


def test_manual_react_rag_persists_chunks_tool_audit_and_progress(
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Verify manual uploads use task-scoped ReAct RAG and audit retrieval.

    Inputs:
        None. The test uses fake model, embedding, and repository boundaries.

    Outputs:
        None. Assertions confirm the manual wrapper stores a RAG attempt,
        evidence chunks, retrieval audit rows, progress updates, and timing
        logs in the legacy task tables.
    """
    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L2-001",
                question_id="Q.STR.01.01",
                level=2,
                item_index=1,
                text="Transition roadmap is documented.",
            )
        ],
    )

    class RepositoryStub:
        """Repository stub capturing task-scoped RAG persistence."""

        def __init__(self) -> None:
            """Initialize empty persistence capture stores."""

            self.attempts: list[dict[str, Any]] = []
            self.saved_chunks: list[dict[str, Any]] = []
            self.retrieval_rounds: list[dict[str, Any]] = []
            self.progress_updates: list[dict[str, Any]] = []
            self.commits = 0

        def save_question_attempt(self, **kwargs: Any) -> str:
            """Record one RAG attempt row."""

            self.attempts.append(kwargs)
            return f"attempt-{len(self.attempts)}"

        def save_evidence_chunks(
            self,
            task_id: str,
            chunks: list[dict[str, Any]],
        ) -> None:
            """Record the task-scoped chunks submitted for persistence."""

            assert task_id == "task-react-rag"
            self.saved_chunks = list(chunks)

        def search_evidence_chunks(
            self,
            task_id: str,
            query_embedding: list[float],
            top_k: int,
        ) -> list[dict[str, Any]]:
            """Return persisted chunks as vector-search candidates."""

            assert task_id == "task-react-rag"
            _ = query_embedding
            return [
                {
                    "chunk_id": chunk["chunk_id"],
                    "attachment_id": chunk["attachment_id"],
                    "file_name": chunk["file_name"],
                    "document_type": chunk["document_type"],
                    "chunk_text": chunk["chunk_text"],
                    "location_json": chunk["location_json"],
                    "similarity_score": 0.94,
                }
                for chunk in self.saved_chunks[:top_k]
            ]

        def save_retrieval_round(self, **kwargs: Any) -> str:
            """Record one manual RAG tool-call audit row."""

            self.retrieval_rounds.append(kwargs)
            return f"retrieval-round-{len(self.retrieval_rounds)}"

        def update_question_progress(self, **kwargs: Any) -> None:
            """Record one manual question progress update."""

            self.progress_updates.append(kwargs)

        def commit(self) -> None:
            """Record an explicit progress commit boundary."""

            self.commits += 1

    class EmbeddingClientStub:
        """Embedding client returning one deterministic vector per text."""

        model_name = "fake-embedding"

        def embed_batch(self, texts: list[str]) -> list[list[float]]:
            """Return stable vectors for chunks and RAG queries."""

            return [[float(index), 0.25] for index, _ in enumerate(texts, start=1)]

    class ToolCallingLlmStub:
        """Fake chat model that calls rag_search once before finalizing."""

        def __init__(self) -> None:
            """Initialize fake tool-call state."""

            self.tool_calls = 0

        def bind_tools(self, tools: list[object]) -> "ToolCallingLlmStub":
            """Return self for the LangChain tool-binding interface."""

            self.tools = tools
            return self

        def invoke(self, messages: list[object]) -> AIMessage:
            """Request one RAG tool call, then stop the graph."""

            tool_messages = [
                message
                for message in messages
                if getattr(message, "type", "") == "tool"
            ]
            if not tool_messages:
                self.tool_calls += 1
                return AIMessage(
                    content="",
                    tool_calls=[
                        {
                            "name": "rag_search",
                            "id": "manual-rag-call-1",
                            "args": {
                                "queries": [
                                    {
                                        "query": "transition strategy roadmap",
                                        "answer_item_id": "Q.STR.01.01-L2-001",
                                        "level": 2,
                                    }
                                ],
                                "top_k_per_query": 8,
                            },
                        }
                    ],
                )
            return AIMessage(content="retrieval complete")

    class ReviewClientStub:
        """Review client returning a deterministic structured result."""

        model = "fake-gpt-5.4"

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Return a stable result after the RAG tool call."""

            prompt_text = input_messages[0]["content"][0]["text"]
            assert "Retrieved evidence chunks:" in prompt_text
            return QuestionReviewResult(
                question_id=question.question_id,
                model=self.model,
                overall_status="supported",
                highest_supported_level=2,
                current_selected_answer_item_ids=[],
                verified_selected_answer_item_ids=["Q.STR.01.01-L2-001"],
            )

    caplog.set_level("INFO", logger="app.services.batch_question_react_graph")
    repository = RepositoryStub()
    result = run_manual_question_react_review(
        review_input=QuestionReviewInput(
            task_id="task-react-rag",
            question=question,
            current_selected_answer_item_ids=[],
        ),
        extracted_documents=[
            ExtractedDocument(
                file_name="strategy.pdf",
                document_type="pdf",
                blocks=[
                    ExtractedBlock(
                        block_id="pdf-page-0001",
                        block_type="page",
                        text="The transition strategy roadmap is documented.",
                        page=1,
                    )
                ],
            )
        ],
        attachment_id_by_file={"strategy.pdf": "attachment-001"},
        repository=repository,
        worker_id="worker-react",
        review_client=ReviewClientStub(),
        embedding_client=EmbeddingClientStub(),
        llm=ToolCallingLlmStub(),
    )

    assert result.overall_status == "supported"
    assert repository.attempts[0]["mode"] == "rag"
    assert len(repository.saved_chunks) == 1
    assert len(repository.retrieval_rounds) == 1
    assert repository.retrieval_rounds[0]["queries"] == [
        "transition strategy roadmap"
    ]
    assert repository.retrieval_rounds[0]["accepted_evidence_json"]["query_count"] == 1
    statuses = [update["status"] for update in repository.progress_updates]
    assert statuses == [
        "embedding_documents",
        "in_progress",
        "retrieving_evidence",
        "finalizing_answer",
        "finalizing_answer",
    ]
    assert repository.progress_updates[0]["progress_message"] == (
        "Creating embeddings for 1 evidence chunk."
    )
    assert repository.progress_updates[2]["rag_call_count"] == 1
    assert repository.progress_updates[2]["query_count"] == 1
    assert repository.progress_updates[3]["retrieved_chunk_count"] == 1
    assert repository.commits >= len(repository.progress_updates)
    logged_stages = {
        getattr(record, "stage", None)
        for record in caplog.records
        if record.message.startswith("manual_ai_review_stage_timing")
    }
    assert {
        "chunk_documents",
        "embed_chunks",
        "save_chunks",
        "react_rag",
        "structured_review",
    }.issubset(logged_stages)
    assert any(
        "manual_ai_review_stage_timing stage=embed_chunks" in record.message
        and "duration_ms=" in record.message
        for record in caplog.records
    )


def test_manual_react_rag_returns_limitation_for_empty_extracted_text() -> None:
    """Verify scan-only evidence completes with an explicit limitation result.

    Inputs:
        None. The test uses an empty extracted PDF payload.

    Outputs:
        None. Assertions confirm no verified answers are selected and retry
        state is cleared through an empty chunk save.
    """
    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L1-001",
                question_id="Q.STR.01.01",
                level=1,
                item_index=1,
                text="Strategy is documented.",
            )
        ],
    )

    class RepositoryStub:
        """Repository stub that records attempt and chunk persistence."""

        def __init__(self) -> None:
            """Initialize empty capture stores."""

            self.attempts: list[dict[str, Any]] = []
            self.saved_chunks: list[dict[str, Any]] | None = None

        def save_question_attempt(self, **kwargs: Any) -> str:
            """Record one RAG attempt."""

            self.attempts.append(kwargs)
            return "attempt-1"

        def save_evidence_chunks(
            self,
            task_id: str,
            chunks: list[dict[str, Any]],
        ) -> None:
            """Record the empty chunk replacement."""

            assert task_id == "task-empty-text"
            self.saved_chunks = list(chunks)

    class ReviewClientStub:
        """Review client stub that should not receive a final review call."""

        model = "fake-gpt-5.4"

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Fail if final review is attempted without evidence chunks."""

            raise AssertionError("final review should not run without chunks")

    repository = RepositoryStub()
    result = run_manual_question_react_review(
        review_input=QuestionReviewInput(
            task_id="task-empty-text",
            question=question,
            current_selected_answer_item_ids=["Q.STR.01.01-L1-001"],
        ),
        extracted_documents=[
            ExtractedDocument(
                file_name="scan-only.pdf",
                document_type="pdf",
                blocks=[],
                metadata={"low_text_density": True, "page_count": 2},
            )
        ],
        attachment_id_by_file={"scan-only.pdf": "attachment-001"},
        repository=repository,
        review_client=ReviewClientStub(),
    )

    assert result.overall_status == "insufficient_extractable_text"
    assert result.current_selected_answer_item_ids == ["Q.STR.01.01-L1-001"]
    assert result.verified_selected_answer_item_ids == []
    assert "did not produce extractable text" in result.warnings[0]
    assert repository.attempts[0]["mode"] == "rag"
    assert repository.saved_chunks == []


def test_routed_direct_review_keeps_sparse_pdf_raw_file_block(
    tmp_path: Path,
) -> None:
    """Verify sparse PDF extraction preserves raw PDF evidence on direct route.

    Inputs:
        tmp_path: Pytest temporary directory used to create a PDF-like evidence
            file without relying on PDF parsing.

    Outputs:
        None. Assertions confirm the direct worker route sends extracted
        metadata and the original PDF file block to the review client.
    """

    pdf_path = tmp_path / "scanned-strategy.pdf"
    pdf_bytes = b"%PDF-1.7 scanned visual evidence"
    pdf_path.write_bytes(pdf_bytes)
    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L2-001",
                question_id="Q.STR.01.01",
                level=2,
                item_index=1,
                text="Transition roadmap is documented.",
            )
        ],
    )

    class RepositoryStub:
        """Repository stub capturing routed attempt persistence."""

        def __init__(self) -> None:
            """Initialize empty attempt capture storage."""

            self.attempts: list[dict[str, Any]] = []

        def save_question_attempt(self, **kwargs: Any) -> str:
            """Record one question-attempt payload.

            Inputs:
                **kwargs: Attempt metadata emitted by routing.

            Outputs:
                str: Fake attempt identifier.
            """

            self.attempts.append(kwargs)
            return f"attempt-{len(self.attempts)}"

    class CapturingReviewClient:
        """Review client stub capturing direct prompt content blocks."""

        model = "fake-gpt-5.4"

        def __init__(self) -> None:
            """Initialize prompt capture storage."""

            self.last_input: list[dict[str, Any]] | None = None

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Capture the prompt and return a deterministic review result.

            Inputs:
                input_messages: Direct review prompt messages.
                language: Requested review language.

            Outputs:
                QuestionReviewResult: Stable result for assertions.
            """

            self.last_input = input_messages
            return QuestionReviewResult(
                question_id=question.question_id,
                model=self.model,
                overall_status="supported",
                highest_supported_level=2,
                current_selected_answer_item_ids=[],
                verified_selected_answer_item_ids=[],
            )

    repository = RepositoryStub()
    review_client = CapturingReviewClient()

    result = run_question_review_with_routing(
        review_input=QuestionReviewInput(
            task_id="task-sparse-pdf",
            question=question,
            attachment_paths=[pdf_path],
        ),
        extracted_documents=[
            ExtractedDocument(
                file_name="scanned-strategy.pdf",
                document_type="pdf",
                blocks=[],
                metadata={
                    "page_count": 2,
                    "extracted_page_count": 0,
                    "empty_page_count": 2,
                    "total_characters": 0,
                    "low_text_density": True,
                    "warnings": [],
                },
            )
        ],
        attachment_id_by_file={"scanned-strategy.pdf": "attachment-001"},
        repository=repository,
        review_client=review_client,
    )

    assert result.overall_status == "supported"
    assert repository.attempts[0]["mode"] == "direct"
    assert review_client.last_input is not None
    content_blocks = review_client.last_input[0]["content"]
    assert any(
        block["type"] == "input_text"
        and "Text extraction for this PDF was sparse" in block["text"]
        for block in content_blocks
    )
    assert [
        block
        for block in content_blocks
        if block["type"] == "input_file"
        and block["filename"] == "scanned-strategy.pdf"
    ] == [
        {
            "type": "input_file",
            "filename": "scanned-strategy.pdf",
            "file_data": (
                "data:application/pdf;base64,"
                f"{base64.b64encode(pdf_bytes).decode('ascii')}"
            ),
        }
    ]


def test_rag_review_does_not_finalize_when_evidence_is_insufficient() -> None:
    """Verify insufficient RAG assessment skips final answer selection.

    Inputs:
        None. The test uses fake repository, embedding, and review boundaries.

    Outputs:
        None. Assertions confirm a bounded RAG stop without sufficient evidence
        returns an explicit warning result and does not call final review.
    """

    question = AssessmentQuestion(
        question_id="Q.STR.02.01",
        dimension="Strategy",
        section="Objectives",
        question="Does the organization set strategic objectives?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.02.01-L1-001",
                question_id="Q.STR.02.01",
                level=1,
                item_index=1,
                text="Objectives are documented.",
            )
        ],
    )

    class RepositoryStub:
        """Repository stub for an insufficient RAG path."""

        def __init__(self) -> None:
            """Initialize saved chunk and retrieval round stores."""

            self.saved_chunks: list[dict[str, Any]] = []
            self.rounds: list[dict[str, Any]] = []

        def save_evidence_chunks(
            self,
            task_id: str,
            chunks: list[dict[str, Any]],
        ) -> None:
            """Record chunks submitted for persistence.

            Inputs:
                task_id: Review task identifier.
                chunks: Chunk dictionaries.

            Outputs:
                None.
            """

            self.saved_chunks = list(chunks)

        def search_evidence_chunks(
            self,
            task_id: str,
            query_embedding: list[float],
            top_k: int,
        ) -> list[dict[str, Any]]:
            """Return persisted chunks as retrieval candidates."""

            return [
                {
                    "chunk_id": chunk["chunk_id"],
                    "attachment_id": chunk["attachment_id"],
                    "file_name": chunk["file_name"],
                    "document_type": chunk["document_type"],
                    "chunk_text": chunk["chunk_text"],
                    "location_json": chunk["location_json"],
                    "similarity_score": 0.88,
                }
                for chunk in self.saved_chunks[:top_k]
            ]

        def save_retrieval_round(self, **kwargs: Any) -> str:
            """Record a retrieval round and return a fake ID."""

            self.rounds.append(kwargs)
            return f"round-{len(self.rounds)}"

    class EmbeddingClientStub:
        """Embedding client returning deterministic vectors."""

        model_name = "fake-embedding"

        def embed_batch(self, texts: list[str]) -> list[list[float]]:
            """Return one fake vector per chunk text."""

            return [[float(index)] for index, _text in enumerate(texts, start=1)]

        def embed_query(self, text: str) -> list[float]:
            """Return a fake query vector."""

            return [0.5]

    class InsufficientReviewClient:
        """Review client that never permits final review."""

        model = "fake-gpt-5.4"

        def assess_retrieved_evidence(
            self,
            input_messages: list[dict[str, object]],
            language: str = "en",
        ) -> RetrievalRoundAssessment:
            """Return an insufficient assessment with no refined queries."""

            return RetrievalRoundAssessment(
                accepted_evidence=[],
                rejected_evidence=[],
                evidence_gaps=["Need board-approved objective evidence."],
                refined_queries=[],
                sufficient_for_final_review=False,
                stop_reason="needs_more_evidence",
            )

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Fail if final answer selection is attempted."""

            raise AssertionError("final review should not be called")

    result = run_rag_question_review(
        review_input=QuestionReviewInput(
            task_id="task-insufficient-rag",
            question=question,
            current_selected_answer_item_ids=["Q.STR.02.01-L1-001"],
        ),
        extracted_documents=[
            ExtractedDocument(
                file_name="objectives.pdf",
                document_type="pdf",
                blocks=[
                    ExtractedBlock(
                        block_id="pdf-page-0001",
                        block_type="page",
                        text="Background description without objective approval.",
                        page=1,
                    )
                ],
            )
        ],
        attachment_id_by_file={"objectives.pdf": "attachment-001"},
        repository=RepositoryStub(),
        review_client=InsufficientReviewClient(),
        embedding_client=EmbeddingClientStub(),
    )

    assert result.overall_status == "insufficient_evidence"
    assert result.current_selected_answer_item_ids == ["Q.STR.02.01-L1-001"]
    assert result.verified_selected_answer_item_ids == []
    assert "Need board-approved objective evidence." in result.warnings[0]


def test_run_task_from_hana_always_uses_manual_react_rag(
    monkeypatch: Any,
) -> None:
    """Verify worker review delegates to manual ReAct RAG after extraction.

    Inputs:
        monkeypatch: Pytest helper used to replace extraction and ReAct helpers.

    Outputs:
        None. Assertions validate extraction persistence and ReAct invocation
        without direct-vs-RAG route selection.
    """

    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L2-001",
                question_id="Q.STR.01.01",
                level=2,
                item_index=1,
                text="Transition roadmap is documented.",
            )
        ],
    )

    class RepositoryStub:
        """Repository stub capturing route and extraction persistence calls."""

        def __init__(self) -> None:
            """Initialize repository capture lists."""

            self.attempts: list[dict[str, Any]] = []
            self.extractions: list[dict[str, Any]] = []

        def list_questions(self, dimension: str) -> list[AssessmentQuestion]:
            """Return the single test question for Strategy."""

            assert dimension == "Strategy"
            return [question]

        def save_attachment_extraction(self, **kwargs: Any) -> str:
            """Record one attachment extraction payload."""

            self.extractions.append(kwargs)
            return "extraction-1"

        def save_question_attempt(self, **kwargs: Any) -> str:
            """Record one routing or model attempt payload."""

            self.attempts.append(kwargs)
            return f"attempt-{len(self.attempts)}"

    repository = RepositoryStub()
    react_calls: list[dict[str, Any]] = []
    extracted_documents = [
        ExtractedDocument(
            file_name="large.pdf",
            document_type="pdf",
            blocks=[
                ExtractedBlock(
                    block_id="pdf-page-0001",
                    block_type="page",
                    text="A" * 2_100_000,
                    page=1,
                )
            ],
            metadata={"warnings": ["ocr-truncated"], "page_count": 1},
        )
    ]

    def fake_extract_evidence_files(paths: list[Path]) -> list[ExtractedDocument]:
        """Return one oversized extracted PDF document."""

        assert [path.name for path in paths] == ["large.pdf"]
        return extracted_documents

    def fake_run_manual_question_react_review(
        review_input: QuestionReviewInput,
        extracted_documents: list[ExtractedDocument],
        attachment_id_by_file: dict[str, str],
        repository: Any,
        worker_id: str = "",
        review_client: Any | None = None,
        embedding_client: Any | None = None,
        llm: Any | None = None,
        attachment_ids_by_document: list[str] | None = None,
    ) -> QuestionReviewResult:
        """Capture manual ReAct RAG call and return a deterministic result."""

        react_calls.append(
            {
                "review_input": review_input,
                "extracted_documents": extracted_documents,
                "attachment_id_by_file": attachment_id_by_file,
                "repository": repository,
                "worker_id": worker_id,
                "attachment_ids_by_document": attachment_ids_by_document,
                "review_client": review_client,
                "embedding_client": embedding_client,
                "llm": llm,
            }
        )
        return QuestionReviewResult(
            question_id=review_input.question.question_id,
            model="fake-gpt-5.4",
            overall_status="rag-complete",
            highest_supported_level=2,
            current_selected_answer_item_ids=[],
            verified_selected_answer_item_ids=[],
        )

    monkeypatch.setattr(
        "app.services.ai_review_graph.extract_evidence_files",
        fake_extract_evidence_files,
    )
    monkeypatch.setattr(
        "app.services.batch_question_react_graph.run_manual_question_react_review",
        fake_run_manual_question_react_review,
    )

    result = run_task_from_hana(
        repository=repository,
        task={
            "task_id": "task-rag-route",
            "dimension": "Strategy",
            "question_id": "Q.STR.01.01",
            "current_selected_answer_item_ids": [],
        },
        attachments=[
            {
                "attachment_id": "attachment-1",
                "file_name": "large.pdf",
                "content": b"%PDF-1.7 large evidence",
            }
        ],
    )

    assert result.overall_status == "rag-complete"
    assert len(react_calls) == 1
    assert react_calls[0]["attachment_id_by_file"] == {"large.pdf": "attachment-1"}
    assert react_calls[0]["attachment_ids_by_document"] == ["attachment-1"]
    assert repository.attempts == []
    assert repository.extractions[0]["attachment_id"] == "attachment-1"


def test_routed_review_falls_back_to_rag_on_direct_context_limit(
    monkeypatch: Any,
) -> None:
    """Verify routed direct review falls back to RAG on context limit errors.

    Inputs:
        monkeypatch: Pytest helper used to replace the RAG fallback callable.

    Outputs:
        None. Assertions validate route persistence, direct failure capture, and
        RAG fallback invocation after a context-limit exception.
    """

    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L2-001",
                question_id="Q.STR.01.01",
                level=2,
                item_index=1,
                text="Transition roadmap is documented.",
            )
        ],
    )

    class RepositoryStub:
        """Repository stub capturing routed attempt persistence."""

        def __init__(self) -> None:
            """Initialize empty attempt capture storage."""

            self.attempts: list[dict[str, Any]] = []

        def save_question_attempt(self, **kwargs: Any) -> str:
            """Record one question attempt payload."""

            self.attempts.append(kwargs)
            return f"attempt-{len(self.attempts)}"

    class ContextLimitReviewClient:
        """Review client stub that raises a direct context-limit error."""

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Raise the routed direct review context-limit error."""

            raise ReviewContextLimitError(
                "Direct review exceeded the model context window."
            )

    repository = RepositoryStub()
    rag_calls: list[dict[str, Any]] = []
    extracted_documents = [
        ExtractedDocument(
            file_name="small.docx",
            document_type="docx",
            blocks=[
                ExtractedBlock(
                    block_id="paragraph-0001",
                    block_type="paragraph",
                    text="The company has a documented strategy review cycle.",
                    section_label="Strategy",
                )
            ],
        )
    ]

    def fake_run_rag_question_review(
        review_input: QuestionReviewInput,
        extracted_documents: list[ExtractedDocument],
        attachment_id_by_file: dict[str, str],
        repository: Any,
        review_client: Any | None = None,
        embedding_client: Any | None = None,
        attachment_ids_by_document: list[str] | None = None,
    ) -> QuestionReviewResult:
        """Capture the fallback call and return a deterministic RAG result."""

        rag_calls.append(
            {
                "review_input": review_input,
                "extracted_documents": extracted_documents,
                "attachment_id_by_file": attachment_id_by_file,
                "repository": repository,
                "attachment_ids_by_document": attachment_ids_by_document,
                "review_client": review_client,
                "embedding_client": embedding_client,
            }
        )
        return QuestionReviewResult(
            question_id=review_input.question.question_id,
            model="fake-gpt-5.4",
            overall_status="rag-fallback",
            highest_supported_level=2,
            current_selected_answer_item_ids=[],
            verified_selected_answer_item_ids=[],
        )

    monkeypatch.setattr(
        "app.services.ai_review_graph.run_rag_question_review",
        fake_run_rag_question_review,
    )

    result = run_question_review_with_routing(
        review_input=QuestionReviewInput(
            task_id="task-direct-fallback",
            question=question,
            current_selected_answer_item_ids=[],
            attachment_paths=[Path("/tmp/small.docx")],
        ),
        extracted_documents=extracted_documents,
        attachment_id_by_file={"small.docx": "attachment-1"},
        repository=repository,
        review_client=ContextLimitReviewClient(),
    )

    assert result.overall_status == "rag-fallback"
    assert len(rag_calls) == 1
    assert repository.attempts[0]["mode"] == "direct"
    assert repository.attempts[0]["status"] == "routed"
    assert repository.attempts[1]["mode"] == "direct"
    assert repository.attempts[1]["status"] == "failed"
    assert repository.attempts[1]["error_code"] == "context_limit"


def test_visual_direct_context_limit_returns_extractable_text_limitation(
    tmp_path: Path,
) -> None:
    """Verify visual-only context overflow does not fall into empty RAG review.

    Inputs:
        tmp_path: Pytest temporary directory used to create an image attachment.

    Outputs:
        None. Assertions confirm direct overflow with image-only evidence returns
        an explicit limitation result.
    """

    image_path = tmp_path / "strategy-chart.png"
    image_path.write_bytes(b"\x89PNG\r\n\x1a\nvisual evidence")
    question = AssessmentQuestion(
        question_id="Q.STR.01.01",
        dimension="Strategy",
        section="Governance",
        question="Is there a transition strategy?",
        answer_items=[
            AnswerItem(
                answer_item_id="Q.STR.01.01-L1-001",
                question_id="Q.STR.01.01",
                level=1,
                item_index=1,
                text="Strategy is documented.",
            )
        ],
    )

    class RepositoryStub:
        """Repository stub capturing route and failure attempts."""

        def __init__(self) -> None:
            """Initialize attempt capture storage."""

            self.attempts: list[dict[str, Any]] = []

        def save_question_attempt(self, **kwargs: Any) -> str:
            """Record one question attempt."""

            self.attempts.append(kwargs)
            return f"attempt-{len(self.attempts)}"

    class ContextLimitReviewClient:
        """Review client that raises on direct visual review."""

        model = "fake-gpt-5.4"

        def review_question(
            self,
            input_messages: list[dict[str, Any]],
            language: str = "en",
        ) -> QuestionReviewResult:
            """Raise a direct context-limit error."""

            raise ReviewContextLimitError("visual evidence exceeded context")

    repository = RepositoryStub()
    result = run_question_review_with_routing(
        review_input=QuestionReviewInput(
            task_id="task-visual-overflow",
            question=question,
            attachment_paths=[image_path],
        ),
        extracted_documents=[],
        attachment_id_by_file={},
        repository=repository,
        review_client=ContextLimitReviewClient(),
    )

    assert result.overall_status == "insufficient_extractable_text"
    assert result.verified_selected_answer_item_ids == []
    assert "text-only RAG fallback cannot index" in result.warnings[0]
    assert [attempt["status"] for attempt in repository.attempts] == [
        "routed",
        "failed",
    ]
