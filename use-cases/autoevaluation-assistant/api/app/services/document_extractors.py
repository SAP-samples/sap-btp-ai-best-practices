"""Pure-Python extractors for Office evidence documents."""

from __future__ import annotations

from email import policy
from email.message import EmailMessage, Message
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from typing import cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field


class ExtractedBlock(BaseModel):
    """Traceable text block extracted from a source document.

    Inputs:
        Field values describing the extracted text and optional location
        metadata for document sections, sheets, tables, and rows.

    Outputs:
        A validated block payload that downstream review prompts can cite.

    Attributes:
        block_id: Stable identifier unique within the extracted document.
        block_type: Type of source content, such as heading, paragraph, table,
            or sheet.
        text: Extracted non-empty text for this block.
        section_label: Optional DOCX heading context for the block.
        sheet_name: Optional XLSX sheet name for the block.
        table_name: Optional table or range label for tabular content.
        page: Optional page number for paginated documents such as PDFs.
        row_start: Optional first source row represented by this block.
        row_end: Optional last source row represented by this block.
    """

    block_id: str
    block_type: str
    text: str
    section_label: str | None = None
    sheet_name: str | None = None
    table_name: str | None = None
    page: int | None = None
    row_start: int | None = None
    row_end: int | None = None


class ExtractedDocument(BaseModel):
    """Extracted text payload for one uploaded evidence document.

    Inputs:
        Field values describing the source file, document type, and extracted
        blocks.

    Outputs:
        A validated document payload with an empty block list when no text is
        extracted.

    Attributes:
        file_name: Base name of the source file.
        document_type: Source document type, such as docx or xlsx.
        blocks: Ordered extracted text blocks with traceable metadata.
        metadata: Optional extraction metadata such as page counts and quality
            flags.
    """

    file_name: str
    document_type: str
    blocks: list[ExtractedBlock] = Field(default_factory=list)
    metadata: dict[str, object] = Field(default_factory=dict)


def _docx_block_id(index: int) -> str:
    """Build a stable DOCX block identifier.

    Args:
        index: One-based block position in the extracted document.

    Returns:
        Stable block identifier for a DOCX extraction block.
    """

    return f"docx-block-{index:04d}"


def _xlsx_block_id(index: int) -> str:
    """Build a stable XLSX block identifier.

    Args:
        index: One-based block position in the extracted workbook.

    Returns:
        Stable block identifier for an XLSX extraction block.
    """

    return f"xlsx-block-{index:04d}"


def _pdf_block_id(page_number: int) -> str:
    """Build a stable PDF page block identifier.

    Args:
        page_number: One-based page number.

    Returns:
        Stable block identifier for a PDF page extraction block.
    """

    return f"pdf-page-{page_number:04d}"


def _eml_block_id(index: int) -> str:
    """Build a stable EML block identifier.

    Args:
        index: One-based block position in the extracted email.

    Returns:
        Stable block identifier for an EML extraction block.
    """

    return f"eml-block-{index:04d}"


class _EmailHtmlTextExtractor(HTMLParser):
    """Convert simple email HTML bodies into readable plain text.

    Inputs:
        HTML fragments from `text/html` email parts.

    Outputs:
        Accumulated text chunks exposed through `text()`.
    """

    _BLOCK_TAGS = {
        "br",
        "div",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "li",
        "p",
        "tr",
    }
    _SKIPPED_TAGS = {"script", "style"}

    def __init__(self) -> None:
        """Initialize parser state for one HTML email body.

        Inputs:
            None.

        Outputs:
            None. Parsed text chunks are stored on the instance.
        """

        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._skip_depth = 0

    def handle_starttag(
        self,
        tag: str,
        attrs: list[tuple[str, str | None]],
    ) -> None:
        """Handle opening tags that affect text extraction.

        Inputs:
            tag: Lowercase or mixed-case HTML tag name.
            attrs: HTML attributes, unused for extraction.

        Outputs:
            None. Block tags append line breaks and script/style starts are
            skipped until their matching end tag.
        """

        _ = attrs
        normalized_tag = tag.lower()
        if normalized_tag in self._SKIPPED_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth == 0 and normalized_tag in self._BLOCK_TAGS:
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        """Handle closing tags that affect text extraction.

        Inputs:
            tag: Lowercase or mixed-case HTML tag name.

        Outputs:
            None. Block tags append line breaks and script/style closes resume
            text extraction.
        """

        normalized_tag = tag.lower()
        if normalized_tag in self._SKIPPED_TAGS and self._skip_depth > 0:
            self._skip_depth -= 1
            return
        if self._skip_depth == 0 and normalized_tag in self._BLOCK_TAGS:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        """Collect visible HTML text.

        Inputs:
            data: Text content emitted by the HTML parser.

        Outputs:
            None. Non-empty text is appended when not inside skipped tags.
        """

        if self._skip_depth == 0 and data.strip():
            self._chunks.append(data)

    def text(self) -> str:
        """Return normalized text collected from the HTML body.

        Inputs:
            None.

        Outputs:
            str: Plain text with empty lines collapsed.
        """

        return _normalize_email_text("".join(self._chunks))


def _normalize_email_text(text: str) -> str:
    """Normalize email text while preserving meaningful line boundaries.

    Args:
        text: Raw email text extracted from a MIME part.

    Returns:
        Text with CRLF normalized, surrounding whitespace removed, and repeated
        blank lines collapsed.
    """

    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in normalized.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def _html_email_to_text(html: str) -> str:
    """Convert an HTML email body to normalized plain text.

    Args:
        html: Raw HTML body from a `text/html` MIME part.

    Returns:
        Plain text extracted from visible HTML nodes.
    """

    parser = _EmailHtmlTextExtractor()
    parser.feed(html)
    parser.close()
    return parser.text()


def _email_part_text(part: Message) -> str:
    """Extract readable text from one email body part.

    Args:
        part: MIME part chosen as the preferred body candidate.

    Returns:
        Plain text for `text/plain` parts, sanitized text for `text/html`
        parts, and an empty string for unsupported part types.
    """

    content_type = part.get_content_type()
    if content_type not in {"text/plain", "text/html"}:
        return ""
    content = part.get_content()
    text = content if isinstance(content, str) else str(content)
    if content_type == "text/html":
        return _html_email_to_text(text)
    return _normalize_email_text(text)


def _email_body_text(message: EmailMessage) -> str:
    """Return the preferred readable body text from an email message.

    Args:
        message: Parsed RFC 822 email message.

    Returns:
        Plain body text when available, otherwise sanitized HTML text, or an
        empty string when no extractable body exists.
    """

    if message.is_multipart():
        body = message.get_body(preferencelist=("plain", "html"))
        return _email_part_text(body) if body is not None else ""
    return _email_part_text(message)


def _email_header_block_text(message: EmailMessage) -> str:
    """Build a normalized header block for important email fields.

    Args:
        message: Parsed RFC 822 email message.

    Returns:
        Newline-delimited header fields, excluding missing or empty values.
    """

    header_names = ["Subject", "From", "To", "Cc", "Date"]
    lines: list[str] = []
    for header_name in header_names:
        value = str(message.get(header_name, "")).strip()
        if value:
            lines.append(f"{header_name}: {value}")
    return "\n".join(lines)


def _email_attachment_metadata(
    message: EmailMessage,
) -> tuple[list[dict[str, str]], list[str]]:
    """Return metadata for email attachments skipped by v1 extraction.

    Args:
        message: Parsed RFC 822 email message.

    Returns:
        Tuple containing attachment descriptors and extraction warnings.
    """

    attachments = [
        {
            "file_name": part.get_filename() or "unnamed attachment",
            "content_type": part.get_content_type(),
        }
        for part in message.iter_attachments()
    ]
    warnings = (
        ["Email attachments are recorded as metadata but not extracted."]
        if attachments
        else []
    )
    return attachments, warnings


def _iter_docx_body_blocks(document: DocxDocument) -> list[Paragraph | Table]:
    """Return paragraphs and tables in body order from a DOCX document.

    Args:
        document: Loaded python-docx document object.

    Returns:
        Ordered list of paragraph and table wrappers from the document body.
    """

    body_blocks: list[Paragraph | Table] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            body_blocks.append(Paragraph(child, document))
        elif isinstance(child, CT_Tbl):
            body_blocks.append(Table(child, document))
    return body_blocks


def _extract_docx_table_rows(table: Table) -> list[str]:
    """Extract non-empty text rows from a DOCX table.

    Args:
        table: Table object from python-docx.

    Returns:
        List of row strings with cell values joined by `` | ``.
    """

    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        row_text = " | ".join(cell for cell in cells if cell)
        if row_text:
            rows.append(row_text)
    return rows


def extract_docx(path: Path) -> ExtractedDocument:
    """Extract headings, paragraphs, and tables from a DOCX file.

    Args:
        path: Path to a DOCX evidence file.

    Returns:
        ExtractedDocument containing non-empty paragraph and table blocks with
        heading section labels when available.
    """

    document = Document(path)
    extracted = ExtractedDocument(file_name=path.name, document_type="docx")
    current_section: str | None = None
    table_index = 0

    for body_block in _iter_docx_body_blocks(document):
        if isinstance(body_block, Paragraph):
            text = body_block.text.strip()
            if not text:
                continue

            style_name = body_block.style.name if body_block.style else ""
            block_type = "heading" if style_name.startswith("Heading") else "paragraph"
            if block_type == "heading":
                current_section = text

            extracted.blocks.append(
                ExtractedBlock(
                    block_id=_docx_block_id(len(extracted.blocks) + 1),
                    block_type=block_type,
                    text=text,
                    section_label=current_section,
                )
            )
            continue

        table_index += 1
        rows = _extract_docx_table_rows(body_block)
        if not rows:
            continue

        extracted.blocks.append(
            ExtractedBlock(
                block_id=_docx_block_id(len(extracted.blocks) + 1),
                block_type="table",
                text="\n".join(rows),
                section_label=current_section,
                table_name=f"Table {table_index}",
                row_start=1,
                row_end=len(rows),
            )
        )

    return extracted


def extract_pdf(path: Path) -> ExtractedDocument:
    """Extract page-level text from a PDF file using pure Python parsing.

    Args:
        path: Path to a PDF evidence file.

    Returns:
        ExtractedDocument containing one text block per page with extractable
        text, plus extraction quality metadata.
    """

    from pypdf import PdfReader

    reader = PdfReader(str(path))
    extracted = ExtractedDocument(file_name=path.name, document_type="pdf")
    empty_page_count = 0
    extraction_warnings: list[str] = []

    for page_index, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as exc:
            text = ""
            extraction_warnings.append(
                f"Page {page_index} text extraction failed: {type(exc).__name__}"
            )

        if not text:
            empty_page_count += 1
            continue

        extracted.blocks.append(
            ExtractedBlock(
                block_id=_pdf_block_id(page_index),
                block_type="page",
                text=text,
                page=page_index,
            )
        )

    total_characters = sum(len(block.text) for block in extracted.blocks)
    page_count = len(reader.pages)
    extracted.metadata = {
        "page_count": page_count,
        "extracted_page_count": len(extracted.blocks),
        "empty_page_count": empty_page_count,
        "total_characters": total_characters,
        "low_text_density": page_count > 0 and total_characters / page_count < 50,
        "warnings": extraction_warnings,
    }
    return extracted


def extract_eml(path: Path) -> ExtractedDocument:
    """Extract headers and readable body text from an EML email file.

    Args:
        path: Path to an RFC 822 `.eml` evidence file.

    Returns:
        ExtractedDocument containing one header block when key headers are
        present and one body block when the message has readable text. Email
        attachments are recorded as metadata and intentionally not extracted.
    """

    message = cast(
        EmailMessage,
        BytesParser(policy=policy.default).parsebytes(path.read_bytes()),
    )
    extracted = ExtractedDocument(file_name=path.name, document_type="eml")

    header_text = _email_header_block_text(message)
    if header_text:
        extracted.blocks.append(
            ExtractedBlock(
                block_id=_eml_block_id(len(extracted.blocks) + 1),
                block_type="headers",
                text=header_text,
            )
        )

    body_text = _email_body_text(message)
    if body_text:
        extracted.blocks.append(
            ExtractedBlock(
                block_id=_eml_block_id(len(extracted.blocks) + 1),
                block_type="body",
                text=body_text,
            )
        )

    skipped_attachments, warnings = _email_attachment_metadata(message)
    if not body_text:
        warnings.append("Email body did not contain extractable text.")

    extracted.metadata = {
        "total_characters": sum(len(block.text) for block in extracted.blocks),
        "has_attachments": bool(skipped_attachments),
        "attachment_count": len(skipped_attachments),
        "skipped_attachments": skipped_attachments,
        "warnings": warnings,
    }
    return extracted


def _format_xlsx_cell(coordinate: str, value: object) -> str:
    """Format an XLSX cell value with coordinate metadata.

    Args:
        coordinate: Excel-style cell coordinate, such as ``A1``.
        value: Cell value loaded by openpyxl.

    Returns:
        String formatted as ``coordinate=value``.
    """

    return f"{coordinate}={value}"


def _format_xlsx_range(
    sheet_name: str,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
) -> str:
    """Build an Excel range label for extracted non-empty cells.

    Args:
        sheet_name: Worksheet title.
        min_row: First row containing extracted content.
        max_row: Last row containing extracted content.
        min_column: First column containing extracted content.
        max_column: Last column containing extracted content.

    Returns:
        Excel-style sheet range, for example ``Objectives!A2:C5``.
    """

    start_cell = f"{get_column_letter(min_column)}{min_row}"
    end_cell = f"{get_column_letter(max_column)}{max_row}"
    return f"{sheet_name}!{start_cell}:{end_cell}"


def extract_xlsx(path: Path) -> ExtractedDocument:
    """Extract non-empty sheet text from an XLSX workbook.

    Args:
        path: Path to an XLSX evidence file.

    Returns:
        ExtractedDocument containing one block per non-empty sheet with cell
        coordinates, formulas, range labels, and row bounds.
    """

    workbook = load_workbook(path, data_only=False, read_only=True)
    extracted = ExtractedDocument(file_name=path.name, document_type="xlsx")

    try:
        for sheet in workbook.worksheets:
            cell_texts: list[str] = []
            non_empty_rows: list[int] = []
            non_empty_columns: list[int] = []
            for row in sheet.iter_rows():
                row_values: list[str] = []
                for cell in row:
                    if cell.value in (None, ""):
                        continue
                    row_values.append(_format_xlsx_cell(cell.coordinate, cell.value))
                    non_empty_rows.append(cell.row)
                    non_empty_columns.append(cell.column)
                if row_values:
                    cell_texts.append(" | ".join(row_values))

            if not cell_texts:
                continue

            first_row = min(non_empty_rows)
            last_row = max(non_empty_rows)
            first_column = min(non_empty_columns)
            last_column = max(non_empty_columns)

            extracted.blocks.append(
                ExtractedBlock(
                    block_id=_xlsx_block_id(len(extracted.blocks) + 1),
                    block_type="sheet",
                    text="\n".join(cell_texts),
                    sheet_name=sheet.title,
                    table_name=_format_xlsx_range(
                        sheet_name=sheet.title,
                        min_row=first_row,
                        max_row=last_row,
                        min_column=first_column,
                        max_column=last_column,
                    ),
                    row_start=first_row,
                    row_end=last_row,
                )
            )
    finally:
        workbook.close()

    return extracted
