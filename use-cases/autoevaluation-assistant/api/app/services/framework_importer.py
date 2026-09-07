"""Import assessment framework seed files into normalized objects."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.models.assessment import AnswerItem, AssessmentQuestion
from app.services.framework_topics import topic_title_for, topic_translation_rows


ITALIAN_DIMENSION_FILES: dict[str, tuple[str, str]] = {
    "Strategy": ("Autoevaluation_Strategia.csv", "Strategia"),
    "Risk & Control Governance": (
        "Autoevaluation_Risk_Control_Governance.csv",
        "Risk & Control Governance",
    ),
    "Organization & Internal Regulatory System": (
        "Autoevaluation_Organizzazione_Sistema_Normativo.csv",
        "Organizzazione & Sistema Normativo Interno",
    ),
    "People & Culture": (
        "Autoevaluation_People_Culture.csv",
        "Persone & Cultura",
    ),
    "Combined Assurance & Management Oversight": (
        "Autoevaluation_Combined_Assurance_Management_Oversight.csv",
        "Combined Assurance & Management Oversight",
    ),
    "Information Systems & Digital": (
        "Autoevaluation_Sistemi_Informativi_Digital.csv",
        "Sistemi Informativi & Digital",
    ),
    "Resilience Factors": (
        "Autoevaluation_Fattori_di_Resilienza.csv",
        "Fattori di Resilienza",
    ),
}
"""Italian CSV filenames and display labels keyed by canonical dimensions."""


@dataclass(frozen=True)
class FrameworkSeed:
    """Normalized framework seed loaded from local import files.

    Inputs:
        questions: Normalized assessment questions parsed from the workbook.
        explanation_count: Number of question explanation rows parsed from CSV.

    Outputs:
        Immutable seed payload used by import scripts and repository writers.

    Attributes:
        questions: Normalized assessment questions with answer items.
        explanation_count: Number of imported question explanation rows.
    """

    questions: list[AssessmentQuestion]
    explanation_count: int


@dataclass(frozen=True)
class QuestionTranslation:
    """Localized text for one assessment question.

    Inputs:
        section: Localized section or question title from the translation CSV.
        question: Localized question text from the translation CSV.

    Outputs:
        Immutable question translation used by the importer and HANA writer.
    """

    section: str
    topic_title: str
    question: str


@dataclass(frozen=True)
class FrameworkTranslationSeed:
    """Localized framework text keyed by canonical English framework IDs.

    Inputs:
        language: Language code for all translations in the seed.
        dimension_translations: Localized dimension labels keyed by canonical
            dimension name.
        question_translations: Localized sections and question text keyed by
            question ID.
        answer_item_translations: Localized answer text keyed by answer item ID.

    Outputs:
        Immutable translation seed ready to persist into HANA companion tables.
    """

    language: str
    dimension_translations: dict[str, str]
    question_translations: dict[str, QuestionTranslation]
    answer_item_translations: dict[str, str]


def _load_explanations(path: Path) -> dict[str, str]:
    """Load question explanations keyed by question ID.

    Args:
        path: CSV path containing question explanations.

    Returns:
        Mapping from question ID to explanation text.
    """

    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        return {
            str(row["question_id"]).strip(): str(row.get("explanation") or "").strip()
            for row in reader
            if row.get("question_id")
        }


def load_framework_seed(workbook_path: Path, explanations_path: Path) -> FrameworkSeed:
    """Load and normalize the assessment framework seed files.

    Args:
        workbook_path: Path to ``assessment_framework.xlsx``.
        explanations_path: Path to the question explanations CSV.

    Returns:
        FrameworkSeed containing normalized questions and answer items.
    """

    explanations = _load_explanations(explanations_path)
    level_columns = {
        1: 4,
        2: 7,
        3: 10,
        4: 13,
        5: 16,
    }
    grouped: dict[str, AssessmentQuestion] = {}
    item_counts: dict[tuple[str, int], int] = {}

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Ecomarket"]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            question_id = str(row[0] or "").strip()
            if not question_id:
                continue

            if question_id not in grouped:
                grouped[question_id] = AssessmentQuestion(
                    question_id=question_id,
                    dimension=str(row[1] or "").strip(),
                    section=str(row[2] or "").strip(),
                    topic_title=topic_title_for(question_id, "en"),
                    question=str(row[3] or "").strip(),
                    explanation=explanations.get(question_id),
                    answer_items=[],
                )

            question = grouped[question_id]
            for level, text_col in level_columns.items():
                text = str(row[text_col] or "").strip()
                if not text:
                    continue

                key = (question_id, level)
                item_counts[key] = item_counts.get(key, 0) + 1
                item_index = item_counts[key]
                question.answer_items.append(
                    AnswerItem(
                        answer_item_id=f"{question_id}-L{level}-{item_index:03d}",
                        question_id=question_id,
                        level=level,
                        item_index=item_index,
                        text=text,
                    )
                )
    finally:
        workbook.close()

    return FrameworkSeed(
        questions=list(grouped.values()),
        explanation_count=len(explanations),
    )


def _validate_italian_csv_columns(path: Path, fieldnames: list[str] | None) -> None:
    """Validate that an Italian framework CSV has the expected columns.

    Args:
        path: CSV file being loaded.
        fieldnames: Header row parsed by ``csv.DictReader``.

    Returns:
        None. A ``ValueError`` is raised when required columns are missing.
    """

    required_columns = {
        "question_code",
        "question_title",
        "question_text",
        "level",
        "item_index",
        "item_text",
    }
    available_columns = set(fieldnames or [])
    missing_columns = sorted(required_columns - available_columns)
    if missing_columns:
        raise ValueError(
            f"Italian framework CSV {path} is missing columns: {missing_columns}"
        )


def _canonical_answer_item_ids(
    questions: list[AssessmentQuestion],
) -> dict[tuple[str, int, int], str]:
    """Return answer item IDs keyed by their question, level, and item index.

    Args:
        questions: Canonical English framework questions.

    Returns:
        Mapping from ``(question_id, level, item_index)`` to answer item ID.
    """

    return {
        (item.question_id, item.level, item.item_index): item.answer_item_id
        for question in questions
        for item in question.answer_items
    }


def load_italian_framework_translations(
    italian_dir: Path,
    questions: list[AssessmentQuestion],
) -> FrameworkTranslationSeed:
    """Load Italian framework CSVs and validate coverage against English seed.

    Args:
        italian_dir: Directory containing one Italian CSV file per dimension.
        questions: Canonical English framework questions already loaded from
            ``assessment_framework.xlsx``.

    Returns:
        FrameworkTranslationSeed: Complete Italian translation payload keyed by
        canonical dimension, question, and answer item identifiers.

    Raises:
        ValueError: Raised when files, columns, question IDs, answer item IDs,
            or repeated question text values do not match the English seed.
    """

    questions_by_id = {question.question_id: question for question in questions}
    answer_item_ids = _canonical_answer_item_ids(questions)
    question_translations: dict[str, QuestionTranslation] = {}
    answer_item_translations: dict[str, str] = {}

    for dimension, (file_name, _display_name) in ITALIAN_DIMENSION_FILES.items():
        exact_path = italian_dir / file_name
        matches = [exact_path] if exact_path.is_file() else list(italian_dir.glob(f"*{file_name}"))
        if len(matches) != 1:
            raise ValueError(
                f"Expected one Italian framework CSV ending in {file_name!r}; "
                f"found {len(matches)} in {italian_dir}"
            )
        path = matches[0]

        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            _validate_italian_csv_columns(path, reader.fieldnames)
            for row in reader:
                question_id = str(row["question_code"]).strip()
                if question_id not in questions_by_id:
                    raise ValueError(
                        f"Italian CSV {path} contains unknown question {question_id}"
                    )
                if questions_by_id[question_id].dimension != dimension:
                    raise ValueError(
                        f"Italian CSV {path} contains question {question_id} "
                        f"from dimension {questions_by_id[question_id].dimension}"
                    )

                translation = QuestionTranslation(
                    section=str(row["question_title"]).strip(),
                    topic_title=topic_title_for(question_id, "it"),
                    question=str(row["question_text"]).strip(),
                )
                existing_translation = question_translations.get(question_id)
                if (
                    existing_translation is not None
                    and existing_translation != translation
                ):
                    raise ValueError(
                        f"Italian CSV {path} has inconsistent repeated text "
                        f"for question {question_id}"
                    )
                question_translations[question_id] = translation

                key = (
                    question_id,
                    int(str(row["level"]).strip()),
                    int(str(row["item_index"]).strip()),
                )
                answer_item_id = answer_item_ids.get(key)
                if answer_item_id is None:
                    raise ValueError(
                        f"Italian CSV {path} contains unknown answer item key {key}"
                    )
                if answer_item_id in answer_item_translations:
                    raise ValueError(
                        f"Italian CSV {path} duplicates answer item {answer_item_id}"
                    )
                answer_item_translations[answer_item_id] = str(
                    row["item_text"]
                ).strip()

    missing_questions = sorted(set(questions_by_id) - set(question_translations))
    missing_items = sorted(set(answer_item_ids.values()) - set(answer_item_translations))
    if missing_questions or missing_items:
        raise ValueError(
            "Italian framework CSVs do not cover the English seed. "
            f"Missing questions: {missing_questions}; missing answer items: {missing_items}"
        )

    return FrameworkTranslationSeed(
        language="it",
        dimension_translations={
            dimension: display_name
            for dimension, (_file_name, display_name) in ITALIAN_DIMENSION_FILES.items()
        },
        question_translations=question_translations,
        answer_item_translations=answer_item_translations,
    )


def build_topic_translation_rows(
    questions: list[AssessmentQuestion],
) -> list[dict[str, str]]:
    """Return English/Italian HANA topic rows in framework question order.

    Inputs:
        questions: Canonical imported framework questions.

    Outputs:
        list[dict[str, str]]: Two approved localized rows per question.
    """

    return topic_translation_rows([question.question_id for question in questions])
