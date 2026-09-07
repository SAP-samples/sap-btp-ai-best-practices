"""Required-sheet and exact-header extraction from validated XLSX bytes."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import load_workbook

from .constants import (
    ESTRAZIONE_SHEET,
    OPTIONAL_COMPANY_NAME_COLUMN,
    REQUIRED_COLUMNS,
    SCORING_VERSION,
)
from .models import (
    BenchmarkRowError,
    BenchmarkValidationError,
    BenchmarkValidationSummary,
)


def file_failure(
    filename: str,
    source_sha256: str,
    code: str,
    message: str,
) -> BenchmarkValidationError:
    """Build a structured validation exception for a workbook-level failure.

    Inputs:
        filename: Original source filename.
        source_sha256: Hex-encoded source content hash.
        code: Stable file validation category.
        message: Human-readable failure description.

    Outputs:
        BenchmarkValidationError: Exception ready to raise to the caller.
    """

    return BenchmarkValidationError(
        BenchmarkValidationSummary(
            source_filename=filename,
            source_sha256=source_sha256,
            scoring_version=SCORING_VERSION,
            sampled_errors=[
                BenchmarkRowError(row_number=0, code=code, message=message)
            ],
            success=False,
            status="rejected",
        )
    )


def read_source_rows(
    content: bytes,
    filename: str,
    source_sha256: str,
) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    """Open the required sheet and return exact-header row mappings.

    Inputs:
        content: Validated XLSX ZIP bytes.
        filename: Original source filename for structured errors.
        source_sha256: Hex-encoded content hash for structured errors.

    Outputs:
        tuple: Actual headers and non-empty worksheet rows with row numbers.

    Raises:
        BenchmarkValidationError: If the workbook, sheet, or headers are invalid.
    """

    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if ESTRAZIONE_SHEET not in workbook.sheetnames:
            raise file_failure(
                filename,
                source_sha256,
                "missing_sheet",
                f"Workbook must contain the {ESTRAZIONE_SHEET!r} sheet",
            )
        sheet = workbook[ESTRAZIONE_SHEET]
        header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [str(value) if value is not None else "" for value in header_values]
        required_only = [
            header for header in headers if header != OPTIONAL_COMPANY_NAME_COLUMN
        ]
        optional_count = headers.count(OPTIONAL_COMPANY_NAME_COLUMN)
        if tuple(required_only) != REQUIRED_COLUMNS or optional_count > 1:
            raise file_failure(
                filename,
                source_sha256,
                "invalid_headers",
                "Workbook headers must contain the exact 30-column contract "
                f"plus optional {OPTIONAL_COMPANY_NAME_COLUMN!r}",
            )
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if all(value is None for value in values):
                continue
            rows.append((row_number, dict(zip(headers, values, strict=False))))
        return headers, rows
    except BenchmarkValidationError:
        raise
    except Exception as exc:
        raise file_failure(
            filename,
            source_sha256,
            "malformed_xlsx",
            "Workbook could not be opened as a valid XLSX file",
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()
