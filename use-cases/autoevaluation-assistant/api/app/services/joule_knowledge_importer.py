"""Workbook parsing and embedding helpers for Joule knowledge resources."""

from __future__ import annotations

import hashlib
import json
import os
import time
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from app.observability.llm_usage_logging import (
    LlmUsageContext,
    TokenUsage,
    emit_llm_usage_event,
    extract_token_usage,
)
from app.services.joule_knowledge import (
    DIMENSION_ALIASES,
    DimensionSeed,
    GlossaryTermSeed,
    JouleKnowledgeSeed,
    QuestionExplanationSeed,
    canonical_dimension_key,
    normalize_question_id,
    normalize_search_text,
)

DEFAULT_EMBEDDING_MODEL = "text-embedding-3-large"
"""Default SAP Gen AI Hub embedding deployment for semantic retrieval."""


def file_sha256(path: Path) -> str:
    """Return a SHA-256 hash for a source workbook.

    Inputs:
        path: File path to hash.

    Outputs:
        str: Hex-encoded SHA-256 digest.
    """

    digest = hashlib.sha256()
    with path.open("rb") as source_file:
        for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def require_headers(
    actual_headers: Iterable[object],
    expected_headers: tuple[str, ...],
    sheet_name: str,
) -> None:
    """Validate that a worksheet starts with the expected headers.

    Inputs:
        actual_headers: Header values read from the first worksheet row.
        expected_headers: Exact header labels required by the parser.
        sheet_name: Worksheet name used in error messages.

    Outputs:
        None.

    Raises:
        ValueError: Raised when any required header is missing or different.
    """

    actual = tuple("" if value is None else str(value).strip() for value in actual_headers)
    if actual[: len(expected_headers)] != expected_headers:
        raise ValueError(
            f"Sheet {sheet_name!r} has unexpected headers. "
            f"Expected {expected_headers!r}, received {actual[:len(expected_headers)]!r}."
        )


def load_glossary_terms(glossary_workbook: Path) -> list[GlossaryTermSeed]:
    """Load English and Italian glossary terms from the source workbook.

    Inputs:
        glossary_workbook: Path to ``assessment_glossary.xlsx``.

    Outputs:
        list[GlossaryTermSeed]: One term for every populated glossary row.
    """

    workbook = load_workbook(glossary_workbook, read_only=True, data_only=True)
    sheet_specs = {
        "Glossario ENG": ("en", ("Word", "Definition")),
        "Glossario ITA": ("it", ("Parola", "Definizione")),
    }
    terms: list[GlossaryTermSeed] = []
    for sheet_name, (language, expected_headers) in sheet_specs.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Missing required glossary sheet {sheet_name!r}.")
        worksheet = workbook[sheet_name]
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        require_headers(header_row, expected_headers, sheet_name)
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            raw_term, raw_definition = row[:2]
            term = "" if raw_term is None else str(raw_term).strip()
            definition = "" if raw_definition is None else str(raw_definition).strip()
            if not term and not definition:
                continue
            if not term or not definition:
                raise ValueError(
                    f"Glossary sheet {sheet_name!r} row {row_number} must contain "
                    "both term and definition."
                )
            terms.append(
                GlossaryTermSeed(
                    language=language,
                    term=term,
                    normalized_term=normalize_search_text(term),
                    definition=definition,
                    source_workbook=glossary_workbook.name,
                    source_sheet=sheet_name,
                    source_row=row_number,
                )
            )
    return terms


def load_question_explanations(
    explanations_workbook: Path,
) -> list[QuestionExplanationSeed]:
    """Load question explanation rows from the source workbook.

    Inputs:
        explanations_workbook: Path to the ``assessment explanation`` explanation workbook.

    Outputs:
        list[QuestionExplanationSeed]: Normalized question explanation rows.
    """

    workbook = load_workbook(explanations_workbook, read_only=True, data_only=True)
    sheet_name = "Recap-eng"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing required explanation sheet {sheet_name!r}.")
    worksheet = workbook[sheet_name]
    header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    require_headers(
        header_row,
        ("ID", "Size", "Question", "Review-Why is it important to talk about the topic?"),
        sheet_name,
    )

    explanations: list[QuestionExplanationSeed] = []
    seen_question_ids: set[str] = set()
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        raw_question_id, raw_dimension, raw_question, raw_explanation = row[:4]
        if not any(cell is not None and str(cell).strip() for cell in row[:4]):
            continue
        question_id = normalize_question_id(str(raw_question_id or ""))
        if question_id in seen_question_ids:
            raise ValueError(f"Duplicate question ID {question_id!r} at row {row_number}.")
        seen_question_ids.add(question_id)
        dimension_text = str(raw_dimension or "").strip()
        question_text = str(raw_question or "").strip()
        explanation_text = str(raw_explanation or "").strip()
        if not dimension_text or not question_text or not explanation_text:
            raise ValueError(f"Explanation sheet row {row_number} is incomplete.")
        explanations.append(
            QuestionExplanationSeed(
                question_id=question_id,
                dimension_key=canonical_dimension_key(dimension_text),
                source_dimension=dimension_text,
                question=question_text,
                explanation=explanation_text,
                source_workbook=explanations_workbook.name,
                source_sheet=sheet_name,
                source_row=row_number,
            )
        )
    return explanations


def load_dimensions(explanations_workbook: Path) -> list[DimensionSeed]:
    """Load dimension explanations from the source workbook.

    Inputs:
        explanations_workbook: Path to the ``assessment explanation`` explanation workbook.

    Outputs:
        list[DimensionSeed]: One canonical dimension explanation per row.
    """

    workbook = load_workbook(explanations_workbook, read_only=True, data_only=True)
    sheet_name = "Dimensioni-ing"
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing required dimension sheet {sheet_name!r}.")
    worksheet = workbook[sheet_name]

    dimensions: list[DimensionSeed] = []
    seen_dimensions: set[str] = set()
    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        raw_text = row[0] if row else None
        text = "" if raw_text is None else str(raw_text).strip()
        if not text:
            continue
        if ":" not in text:
            raise ValueError(
                f"Dimension sheet row {row_number} must use '<dimension>: <explanation>'."
            )
        raw_name, explanation = text.split(":", maxsplit=1)
        dimension_key = canonical_dimension_key(raw_name)
        if dimension_key in seen_dimensions:
            raise ValueError(f"Duplicate dimension {dimension_key!r} at row {row_number}.")
        seen_dimensions.add(dimension_key)
        aliases = tuple(dict.fromkeys((*DIMENSION_ALIASES[dimension_key], raw_name.strip())))
        dimensions.append(
            DimensionSeed(
                dimension_key=dimension_key,
                name=dimension_key,
                aliases=aliases,
                explanation=explanation.strip(),
                source_workbook=explanations_workbook.name,
                source_sheet=sheet_name,
                source_row=row_number,
            )
        )
    return dimensions


def load_joule_knowledge_seed(
    glossary_workbook: Path,
    explanations_workbook: Path,
) -> JouleKnowledgeSeed:
    """Load all source resources used by the Joule knowledge agent.

    Inputs:
        glossary_workbook: Path to the bilingual glossary workbook.
        explanations_workbook: Path to the question and dimension explanation
            workbook.

    Outputs:
        JouleKnowledgeSeed: Normalized glossary, question, and dimension data.
    """

    return JouleKnowledgeSeed(
        glossary_terms=load_glossary_terms(glossary_workbook),
        question_explanations=load_question_explanations(explanations_workbook),
        dimensions=load_dimensions(explanations_workbook),
        glossary_workbook=str(glossary_workbook),
        explanations_workbook=str(explanations_workbook),
    )


class GenAiHubEmbeddingClient:
    """Generate embeddings through SAP Gen AI Hub's OpenAI-compatible client.

    Inputs:
        model_name: Optional embedding deployment name. Defaults to the
            ``GENAI_EMBEDDING_MODEL`` environment variable, then
            ``text-embedding-3-large``.

    Outputs:
        Client object exposing single-query and batch embedding helpers.
    """

    def __init__(self, model_name: str | None = None) -> None:
        """Initialize the embedding client configuration.

        Inputs:
            model_name: Optional SAP Gen AI Hub embedding model name.

        Outputs:
            None.
        """

        self.model_name = (
            model_name
            or os.getenv("GENAI_EMBEDDING_MODEL")
            or DEFAULT_EMBEDDING_MODEL
        )

    def embed_query(self, text: str) -> list[float]:
        """Embed one query string.

        Inputs:
            text: Text to embed.

        Outputs:
            list[float]: Embedding vector returned by SAP Gen AI Hub.
        """

        return self.embed_batch([text])[0]

    def embed_batch(self, texts: list[str]) -> list[list[float]]:
        """Embed a batch of text strings.

        Inputs:
            texts: Ordered input strings.

        Outputs:
            list[list[float]]: Embedding vectors in the same order as input.

        Raises:
            RuntimeError: Raised when the SAP Gen AI Hub SDK is unavailable.
        """

        started_at = time.perf_counter()
        usage = TokenUsage()
        context = LlmUsageContext(
            route="worker:document-processing:embeddings",
            actor_type="batch",
        )
        try:
            from gen_ai_hub.proxy.native.openai import embeddings
        except ImportError as exc:
            emit_llm_usage_event(
                route=context.route,
                method=context.method,
                actor_type=context.actor_type,
                model=self.model_name,
                llm_endpoint="embeddings",
                input_tokens=0,
                output_tokens=0,
                outcome="error",
                latency_ms=int((time.perf_counter() - started_at) * 1000),
                correlation_id=context.correlation_id,
            )
            raise RuntimeError(
                "SAP Gen AI Hub embeddings are unavailable. Install "
                "sap-ai-sdk-gen and configure SAP AI Core credentials."
            ) from exc

        try:
            response = embeddings.create(model_name=self.model_name, input=texts)
            usage = extract_token_usage(response)
            vectors = [list(item.embedding) for item in response.data]
            emit_llm_usage_event(
                route=context.route,
                method=context.method,
                actor_type=context.actor_type,
                model=self.model_name,
                llm_endpoint="embeddings",
                input_tokens=usage.input_tokens,
                cached_input_tokens=usage.cached_input_tokens,
                cache_write_input_tokens=usage.cache_write_input_tokens,
                input_total_tokens=usage.input_total_tokens,
                total_tokens=usage.total_tokens,
                output_tokens=0,
                outcome="success",
                latency_ms=int((time.perf_counter() - started_at) * 1000),
                correlation_id=context.correlation_id,
            )
            return vectors
        except Exception:
            emit_llm_usage_event(
                route=context.route,
                method=context.method,
                actor_type=context.actor_type,
                model=self.model_name,
                llm_endpoint="embeddings",
                input_tokens=usage.input_tokens,
                cached_input_tokens=usage.cached_input_tokens,
                cache_write_input_tokens=usage.cache_write_input_tokens,
                input_total_tokens=usage.input_total_tokens,
                total_tokens=usage.total_tokens,
                output_tokens=0,
                outcome="error",
                latency_ms=int((time.perf_counter() - started_at) * 1000),
                correlation_id=context.correlation_id,
            )
            raise


def build_question_embedding_rows(
    seed: JouleKnowledgeSeed,
    embedding_client: GenAiHubEmbeddingClient,
    batch_size: int = 32,
    show_progress: bool = True,
) -> dict[str, tuple[list[float], list[float]]]:
    """Generate question and explanation embeddings for each question row.

    Inputs:
        seed: Normalized Joule knowledge seed.
        embedding_client: SAP Gen AI Hub embedding client.
        batch_size: Number of texts sent per embedding API call.
        show_progress: Whether to render a CLI progress bar.

    Outputs:
        dict[str, tuple[list[float], list[float]]]: Mapping from question ID to
        ``(question_embedding, explanation_embedding)``.
    """

    from tqdm import tqdm

    texts: list[str] = []
    keys: list[tuple[str, str]] = []
    for question in seed.question_explanations:
        texts.append(question.question)
        keys.append((question.question_id, "question"))
        texts.append(question.explanation)
        keys.append((question.question_id, "explanation"))

    embeddings_by_key: dict[tuple[str, str], list[float]] = {}
    iterator = range(0, len(texts), batch_size)
    if show_progress:
        iterator = tqdm(iterator, desc="Embedding Joule knowledge", unit="batch")
    for start in iterator:
        batch_texts = texts[start : start + batch_size]
        batch_embeddings = embedding_client.embed_batch(batch_texts)
        for key, embedding in zip(keys[start : start + batch_size], batch_embeddings):
            embeddings_by_key[key] = embedding

    return {
        question.question_id: (
            embeddings_by_key[(question.question_id, "question")],
            embeddings_by_key[(question.question_id, "explanation")],
        )
        for question in seed.question_explanations
    }


def source_metadata(seed: JouleKnowledgeSeed) -> list[dict[str, object]]:
    """Build source metadata rows for the import tracking table.

    Inputs:
        seed: Normalized Joule knowledge seed.

    Outputs:
        list[dict[str, object]]: Source workbook metadata including hashes and
        row counts.
    """

    glossary_path = Path(seed.glossary_workbook)
    explanations_path = Path(seed.explanations_workbook)
    return [
        {
            "source_name": "glossary",
            "source_path": glossary_path.name,
            "source_sha256": file_sha256(glossary_path),
            "source_sheet": "Glossario ENG, Glossario ITA",
            "row_count": len(seed.glossary_terms),
            "status": "completed",
        },
        {
            "source_name": "question_explanations",
            "source_path": explanations_path.name,
            "source_sha256": file_sha256(explanations_path),
            "source_sheet": "Recap-eng, Dimensioni-ing",
            "row_count": len(seed.question_explanations) + len(seed.dimensions),
            "status": "completed",
        },
    ]


def vector_to_json(vector: list[float] | None) -> str | None:
    """Serialize an embedding vector for HANA ``TO_REAL_VECTOR``.

    Inputs:
        vector: Optional embedding vector.

    Outputs:
        str | None: JSON array string accepted by ``TO_REAL_VECTOR``.
    """

    if vector is None:
        return None
    return json.dumps(vector, separators=(",", ":"))
