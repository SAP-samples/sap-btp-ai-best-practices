"""Generate synthetic Strategy evidence files for the assessment document review PoC.

Example commands:
    cd api
    python scripts/generate_synthetic_strategy_fixtures.py --output-dir ../data/synthetic_strategy
    python -m scripts.generate_synthetic_strategy_fixtures --output-dir /tmp/assessment_strategy_fixtures
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


SYNTHETIC_NOTICE = "SYNTHETIC PoC EVIDENCE - not real customer evidence."


def _write_strategy_plan_pdf(pdf_path: Path) -> None:
    """Write a synthetic Strategy plan PDF fixture.

    Inputs:
        pdf_path: Destination path for the generated PDF file.

    Outputs:
        None. The PDF is written to ``pdf_path``.
    """

    pdf = FPDF()
    pdf.set_compression(False)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(0, 8, SYNTHETIC_NOTICE)
    pdf.ln(4)
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(
        0,
        7,
        "Strategy governance excerpt for Q.STR.01.01: the Board Strategy "
        "Committee approves the five-year strategic plan, annual budget, "
        "capital allocation priorities, and climate transition objectives. "
        "Minutes record approval dates, owners, and required follow-up actions.",
    )
    pdf.ln(2)
    pdf.multi_cell(
        0,
        7,
        "Strategy monitoring excerpt for Q.STR.03.01: management reviews "
        "strategic KPIs quarterly, including project delivery milestones, "
        "cash-flow discipline, emissions intensity, and corrective actions "
        "for underperforming initiatives.",
    )
    pdf.output(str(pdf_path))


def _write_objectives_policy_docx(docx_path: Path) -> None:
    """Write a synthetic objectives and governance DOCX fixture.

    Inputs:
        docx_path: Destination path for the generated Word document.

    Outputs:
        None. The DOCX document is written to ``docx_path``.
    """

    document = Document()
    document.add_heading(SYNTHETIC_NOTICE, level=1)
    document.add_heading("Strategy Objective Governance Policy", level=2)
    document.add_paragraph(
        "Q.STR.01.01 evidence: the company defines medium-term and long-term "
        "strategic objectives through an annual planning cycle sponsored by "
        "the Chief Strategy Officer and approved by the Board."
    )
    document.add_paragraph(
        "Q.STR.03.01 evidence: each objective is assigned an owner, a target "
        "metric, a review frequency, and a corrective-action protocol when "
        "quarterly performance falls below tolerance."
    )
    table = document.add_table(rows=1, cols=4)
    header_cells = table.rows[0].cells
    header_cells[0].text = "Question ID"
    header_cells[1].text = "Control"
    header_cells[2].text = "Owner"
    header_cells[3].text = "Evidence"
    rows = [
        (
            "Q.STR.01.01",
            "Board-approved strategic plan",
            "Board Strategy Committee",
            "Annual plan approval pack and minutes",
        ),
        (
            "Q.STR.03.01",
            "Quarterly KPI and action review",
            "Executive Management Committee",
            "Quarterly strategy performance dashboard",
        ),
    ]
    for question_id, control, owner, evidence in rows:
        cells = table.add_row().cells
        cells[0].text = question_id
        cells[1].text = control
        cells[2].text = owner
        cells[3].text = evidence
    document.save(docx_path)


def _write_objectives_register_xlsx(xlsx_path: Path) -> None:
    """Write a synthetic Strategy objectives register XLSX fixture.

    Inputs:
        xlsx_path: Destination path for the generated Excel workbook.

    Outputs:
        None. The workbook is written to ``xlsx_path``.
    """

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Objectives"
    sheet.append(["Question ID", "Evidence Area", "Metric", "Frequency", "Owner"])
    sheet.append(
        [
            "Q.STR.01.01",
            "Board-approved strategic plan and annual budget",
            "Approval status and documented strategic priorities",
            "Annual",
            "Board Strategy Committee",
        ]
    )
    sheet.append(
        [
            "Q.STR.03.01",
            "Strategic objective performance monitoring",
            "KPI status, variance, and corrective action completion",
            "Quarterly",
            "Executive Management Committee",
        ]
    )
    sheet.append(
        [
            "Synthetic Notice",
            SYNTHETIC_NOTICE,
            "Used only for PoC document review testing",
            "N/A",
            "AI Review Team",
        ]
    )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E4"
    sheet.row_dimensions[1].height = 28

    header_fill = PatternFill(fill_type="solid", fgColor="0A6ED1")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in sheet.iter_rows(min_row=2, max_row=4, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column, width in {
        "A": 20,
        "B": 42,
        "C": 44,
        "D": 15,
        "E": 32,
    }.items():
        sheet.column_dimensions[column].width = width

    workbook.save(xlsx_path)


def generate_fixtures(output_dir: Path) -> list[Path]:
    """Generate synthetic PDF, DOCX, and XLSX Strategy evidence files.

    Inputs:
        output_dir: Directory where generated fixture files are written.

    Outputs:
        list[Path]: Paths to the generated PDF, DOCX, and XLSX files.
    """

    output_dir.mkdir(parents=True, exist_ok=True)

    pdf_path = output_dir / "synthetic_strategy_plan.pdf"
    docx_path = output_dir / "synthetic_objectives_policy.docx"
    xlsx_path = output_dir / "synthetic_objectives_register.xlsx"

    _write_strategy_plan_pdf(pdf_path)
    _write_objectives_policy_docx(docx_path)
    _write_objectives_register_xlsx(xlsx_path)

    return [pdf_path, docx_path, xlsx_path]


def parse_args() -> argparse.Namespace:
    """Parse fixture generator CLI arguments.

    Inputs:
        None. Arguments are read from the command line.

    Outputs:
        argparse.Namespace: Parsed command arguments with ``output_dir``.
    """

    parser = argparse.ArgumentParser(description="Generate synthetic Strategy evidence fixtures.")
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    """Run the fixture generator CLI.

    Inputs:
        None. Command-line arguments are parsed from the current process.

    Outputs:
        None. Generated file paths are printed to standard output.
    """

    args = parse_args()
    outputs = generate_fixtures(args.output_dir)
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
