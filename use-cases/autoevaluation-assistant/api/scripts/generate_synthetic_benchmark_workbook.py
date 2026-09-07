"""Generate a deterministic, synthetic assessment benchmark workbook.

Examples:
    python scripts/generate_synthetic_benchmark_workbook.py --template path/to/source.xlsx --output path/to/synthetic.xlsx
    python scripts/generate_synthetic_benchmark_workbook.py --template path/to/source.xlsx --output path/to/synthetic.xlsx --companies 24 --seed 20260713
"""

from __future__ import annotations

import argparse
import hashlib
import math
import re
import sys
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

API_ROOT = Path(__file__).resolve().parents[1]
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))

from app.services.benchmark_import.constants import REQUIRED_COLUMNS
from app.services.customer_class_scope import max_allowed_level
from app.services.question_scoring import calculate_maturity_score


DIMENSION_LABEL_BY_PREFIX = {
    "STR": "Strategia",
    "RCG": "Risk & Control Governance",
    "ORG": "Organizzazione & Sistema Normativo Interno",
    "PCU": "Persone & Cultura",
    "CAM": "Combined Assurance & Management Oversight",
    "SID": "Sistemi Informativi & Digital",
    "FDR": "Fattori di Resilienza",
}
"""Correct Italian source dimension labels keyed by canonical question prefix."""

FIXED_WORKBOOK_TIME = datetime(2026, 1, 1, 0, 0, 0)
"""Stable workbook metadata timestamp used for byte-deterministic output."""

FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
"""Stable ZIP member timestamp used after openpyxl serialization."""

LONG_TEXT_COLUMNS = frozenset({"Domanda", "Testo risposta"})
"""Columns whose cell values should wrap instead of being visually clipped."""

COLUMN_WIDTH_BY_HEADER = {
    "Data Estrazione": 16,
    "Data Sottomissione": 18,
    "ID Impresa": 20,
    "Fatturato": 14,
    "Nr Dipendenti": 14,
    "Settore Operativo (NACE) 1": 34,
    "Settore Operativo (NACE) 2": 34,
    "Settore Operativo (NACE) 3": 34,
    "Dimensione Azienda": 26,
    "Classe": 10,
    "Forma Giuridica": 26,
    "Presenza Geografica": 24,
    "Quotata": 12,
    "Committente Contratti Pubblici": 30,
    "Adesione Codice di  Autodisciplina": 32,
    "ID Questionario": 27,
    "Stato Questionario": 20,
    "Assessment Score": 22,
    "Dimensione": 42,
    "Dimensione Score": 20,
    "ID Domanda": 18,
    "Sezione": 36,
    "Domanda": 60,
    "Domanda Gestita": 18,
    "Stato Validazione": 20,
    "Livello": 10,
    "ID Risposta": 28,
    "Testo risposta": 60,
    "Valore Risposta": 18,
    "Opzionale": 12,
}
"""Bounded widths selected for each field's identifier or text semantics."""


def _stable_number(seed: int, company_index: int, label: str, modulus: int) -> int:
    """Return a deterministic bounded integer without process-randomized hashes.

    Inputs:
        seed: Caller-provided generator seed.
        company_index: Zero-based synthetic company index.
        label: Stable field or question identity.
        modulus: Positive exclusive upper bound.

    Outputs:
        int: Deterministic value from zero through ``modulus - 1``.
    """

    if modulus <= 0:
        raise ValueError("modulus must be positive")
    digest = hashlib.sha256(
        f"{seed}|{company_index}|{label}".encode("utf-8")
    ).digest()
    return int.from_bytes(digest[:8], "big") % modulus


def _dimension_label(question_id: str) -> str:
    """Derive the correct source dimension label from a canonical question ID.

    Inputs:
        question_id: Canonical ID such as ``Q.STR.01.01``.

    Outputs:
        str: Correct Italian workbook dimension label.

    Raises:
        ValueError: If the question prefix is unsupported.
    """

    parts = question_id.split(".")
    prefix = parts[1] if len(parts) > 1 else ""
    try:
        return DIMENSION_LABEL_BY_PREFIX[prefix]
    except KeyError as exc:
        raise ValueError(f"Unsupported question dimension prefix: {question_id}") from exc


def _load_response_catalog(template: Path) -> list[dict[str, Any]]:
    """Load one complete questionnaire's response catalog from a template.

    Inputs:
        template: Source XLSX containing the exact ``Estrazione`` shape.

    Outputs:
        list[dict[str, Any]]: First questionnaire's 551 ordered response rows.

    Raises:
        ValueError: If sheet, headers, or expected catalog counts are invalid.
    """

    workbook = load_workbook(template, read_only=True, data_only=True)
    try:
        if "Estrazione" not in workbook.sheetnames:
            raise ValueError("Template must contain the 'Estrazione' sheet")
        sheet = workbook["Estrazione"]
        headers = [
            str(value)
            for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        ]
        if headers != list(REQUIRED_COLUMNS):
            raise ValueError("Template must use the exact 30-column benchmark contract")
        rows: list[dict[str, Any]] = []
        first_questionnaire_id: str | None = None
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in values):
                continue
            row = dict(zip(headers, values, strict=True))
            questionnaire_id = str(row["ID Questionario"])
            if first_questionnaire_id is None:
                first_questionnaire_id = questionnaire_id
            if questionnaire_id != first_questionnaire_id:
                break
            rows.append(row)
    finally:
        workbook.close()

    if len(rows) != 551:
        raise ValueError(
            f"Template questionnaire must contain 551 response rows; found {len(rows)}"
        )
    if len({str(row["ID Domanda"]) for row in rows}) != 38:
        raise ValueError("Template questionnaire must contain 38 class-3 questions")
    return rows


def _synthetic_profile(seed: int, company_index: int) -> dict[str, Any]:
    """Create deterministic profile metadata unrelated to source companies.

    Inputs:
        seed: Caller-provided generator seed.
        company_index: Zero-based synthetic company index.

    Outputs:
        dict[str, Any]: Synthetic company, questionnaire, date, and profile fields.
    """

    company_number = company_index + 1
    submitted = datetime(2025, 1, 15) + timedelta(
        days=_stable_number(seed, company_index, "submission-date", 400)
    )
    extracted = submitted + timedelta(days=30)
    return {
        "Data Estrazione": extracted.date().isoformat(),
        "Data Sottomissione": submitted.date().isoformat(),
        "ID Impresa": f"SYN-COMPANY-{company_number:03d}",
        "Fatturato": 5_000_000
        + _stable_number(seed, company_index, "revenue", 35_000_000),
        "Nr Dipendenti": 50
        + _stable_number(seed, company_index, "employees", 175),
        "Settore Operativo (NACE) 1": (
            "Synthetic professional, scientific and technical activities"
        ),
        "Settore Operativo (NACE) 2": (
            f"Synthetic professional services cohort {1 + company_index % 4}"
        ),
        "Settore Operativo (NACE) 3": (
            f"Synthetic technical services cohort {1 + company_index % 8}"
        ),
        "Dimensione Azienda": "Synthetic medium enterprise",
        "Classe": "3",
        "Forma Giuridica": "Synthetic limited company",
        "Presenza Geografica": (
            "Synthetic national" if company_index % 2 == 0 else "Synthetic international"
        ),
        "Quotata": "No",
        "Committente Contratti Pubblici": "Si" if company_index % 3 == 0 else "No",
        "Adesione Codice di  Autodisciplina": (
            "Si" if company_index % 5 == 0 else "No"
        ),
        "ID Questionario": f"SYN-QUESTIONNAIRE-{company_number:03d}",
        "Stato Questionario": "REL",
        "Assessment Score": "Synthetic benchmark",
    }


def _source_answer_ids(catalog: list[dict[str, Any]]) -> list[str]:
    """Create stable synthetic source response IDs in catalog row order.

    Inputs:
        catalog: Ordered source response catalog for one questionnaire.

    Outputs:
        list[str]: One synthetic response ID per catalog row.
    """

    ordinals: dict[str, int] = defaultdict(int)
    answer_ids: list[str] = []
    for row in catalog:
        question_id = str(row["ID Domanda"])
        ordinals[question_id] += 1
        answer_ids.append(f"SYN-R-{question_id}-{ordinals[question_id]:03d}")
    return answer_ids


def _company_rows(
    catalog: list[dict[str, Any]],
    synthetic_answer_ids: list[str],
    *,
    seed: int,
    company_index: int,
) -> list[dict[str, Any]]:
    """Build one complete synthetic questionnaire and reconciled scores.

    Inputs:
        catalog: Ordered 551-row response catalog.
        synthetic_answer_ids: Stable synthetic IDs aligned to catalog rows.
        seed: Caller-provided generator seed.
        company_index: Zero-based synthetic company index.

    Outputs:
        list[dict[str, Any]]: Exact 30-column rows for one synthetic company.
    """

    profile = _synthetic_profile(seed, company_index)
    target_level_by_question: dict[str, int] = {}
    for question_id in {str(row["ID Domanda"]) for row in catalog}:
        maximum_level = max_allowed_level(question_id, "class_3")
        if maximum_level <= 0:
            raise ValueError(f"Template contains a non-class-3 question: {question_id}")
        target_level_by_question[question_id] = 1 + _stable_number(
            seed,
            company_index,
            question_id,
            maximum_level,
        )

    rows: list[dict[str, Any]] = []
    for catalog_row, synthetic_answer_id in zip(
        catalog,
        synthetic_answer_ids,
        strict=True,
    ):
        question_id = str(catalog_row["ID Domanda"])
        maturity_level = int(catalog_row["Livello"])
        managed = bool(catalog_row["Domanda Gestita"])
        selected = managed and maturity_level <= target_level_by_question[question_id]
        row = dict(catalog_row)
        row.update(profile)
        row.update(
            {
                "Dimensione": _dimension_label(question_id),
                "ID Risposta": synthetic_answer_id,
                "Valore Risposta": "Si" if selected else "No",
            }
        )
        rows.append(row)

    # Calculate exactly the same equal-share-per-level question formula used by
    # the application and importer, then repeat truncated dimension audit scores.
    rows_by_question: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        rows_by_question[str(row["ID Domanda"])].append(row)
    question_scores: dict[str, float] = {}
    for question_id, question_rows in rows_by_question.items():
        maximum_level = max_allowed_level(question_id, "class_3")
        calculation = calculate_maturity_score(
            [
                (str(row["ID Risposta"]), int(row["Livello"]))
                for row in question_rows
            ],
            [
                str(row["ID Risposta"])
                for row in question_rows
                if row["Valore Risposta"] == "Si"
            ],
            maximum_level,
        )
        question_scores[question_id] = calculation.score
    questions_by_dimension: dict[str, list[str]] = defaultdict(list)
    for question_id in question_scores:
        questions_by_dimension[_dimension_label(question_id)].append(question_id)
    dimension_scores = {
        dimension: math.floor(
            sum(question_scores[question_id] for question_id in question_ids)
            / len(question_ids)
            + 1e-9
        )
        for dimension, question_ids in questions_by_dimension.items()
    }
    for row in rows:
        row["Dimensione Score"] = dimension_scores[str(row["Dimensione"])]
    return rows


def _normalize_zip_metadata(path: Path) -> None:
    """Rewrite XLSX members with stable timestamps and deterministic ordering.

    Inputs:
        path: Generated XLSX path to normalize in place.

    Outputs:
        None. A deterministic temporary ZIP atomically replaces the source path.
    """

    temporary = path.with_suffix(path.suffix + ".deterministic")
    with zipfile.ZipFile(path, mode="r") as source, zipfile.ZipFile(
        temporary,
        mode="w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as destination:
        for original in sorted(source.infolist(), key=lambda item: item.filename):
            payload = source.read(original.filename)
            if original.filename == "docProps/core.xml":
                # openpyxl overwrites ``modified`` at save time, so normalize
                # the embedded property as well as the surrounding ZIP metadata.
                payload = re.sub(
                    rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-01-01T00:00:00Z\g<2>",
                    payload,
                    count=1,
                )
            normalized = zipfile.ZipInfo(original.filename, date_time=FIXED_ZIP_TIME)
            normalized.compress_type = zipfile.ZIP_DEFLATED
            normalized.external_attr = original.external_attr
            normalized.create_system = 0
            destination.writestr(normalized, payload)
    temporary.replace(path)


def _configure_sheet(sheet: Any, *, data_row_count: int) -> None:
    """Apply stable navigation aids and bounded semantic column widths.

    Inputs:
        sheet: Write-only ``Estrazione`` worksheet being generated.
        data_row_count: Number of response rows that will follow the header.

    Outputs:
        None. Sheet view, filter range, and column dimensions are updated.
    """

    sheet.freeze_panes = "A2"
    final_column = get_column_letter(len(REQUIRED_COLUMNS))
    sheet.auto_filter.ref = f"A1:{final_column}{data_row_count + 1}"
    for column_index, header in enumerate(REQUIRED_COLUMNS, start=1):
        column_letter = get_column_letter(column_index)
        sheet.column_dimensions[column_letter].width = COLUMN_WIDTH_BY_HEADER[header]


def _append_header(sheet: Any) -> None:
    """Append the styled exact-contract header to a write-only worksheet.

    Inputs:
        sheet: Write-only ``Estrazione`` worksheet being generated.

    Outputs:
        None. One bold, filled, wrapped header row is appended.
    """

    header_cells: list[WriteOnlyCell] = []
    for header in REQUIRED_COLUMNS:
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header_cells.append(cell)
    sheet.append(header_cells)


def _append_data_row(sheet: Any, row: dict[str, Any]) -> None:
    """Append one exact-contract row with wrapping on long narrative fields.

    Inputs:
        sheet: Write-only ``Estrazione`` worksheet being generated.
        row: Synthetic response values keyed by required column name.

    Outputs:
        None. One ordered data row is appended to the worksheet.
    """

    cells: list[Any] = []
    for header in REQUIRED_COLUMNS:
        value = row.get(header)
        if header in LONG_TEXT_COLUMNS:
            cell = WriteOnlyCell(sheet, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cells.append(cell)
        else:
            cells.append(value)
    sheet.append(cells)


def generate_workbook(
    template: Path,
    output: Path,
    companies: int = 24,
    seed: int = 20260713,
) -> Path:
    """Generate deterministic synthetic questionnaires from a response template.

    Inputs:
        template: Source XLSX used only for question and answer catalog shape.
        output: Destination path for the generated workbook.
        companies: Number of complete synthetic class-3 companies to generate.
        seed: Stable pseudo-random seed controlling answer selections.

    Outputs:
        Path: Destination workbook path after generation.
    """

    if companies <= 0:
        raise ValueError("companies must be positive")
    catalog = _load_response_catalog(template)
    synthetic_answer_ids = _source_answer_ids(catalog)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook(write_only=True)
    workbook.properties.creator = "Evaluation Assessment Assistant synthetic benchmark generator"
    workbook.properties.lastModifiedBy = (
        "Evaluation Assessment Assistant synthetic benchmark generator"
    )
    workbook.properties.title = "Synthetic assessment benchmark fixture"
    workbook.properties.created = FIXED_WORKBOOK_TIME
    workbook.properties.modified = FIXED_WORKBOOK_TIME
    sheet = workbook.create_sheet("Estrazione")
    _configure_sheet(sheet, data_row_count=companies * len(catalog))
    _append_header(sheet)

    # The company loop is the long-running unit and always reports progress in
    # CLI use. ``disable=None`` keeps non-interactive test output quiet.
    for company_index in tqdm(
        range(companies),
        desc="Generating synthetic benchmark companies",
        unit="company",
        disable=None,
    ):
        for row in _company_rows(
            catalog,
            synthetic_answer_ids,
            seed=seed,
            company_index=company_index,
        ):
            _append_data_row(sheet, row)
    workbook.save(output)
    workbook.close()
    _normalize_zip_metadata(output)
    return output


def build_parser() -> argparse.ArgumentParser:
    """Build the synthetic benchmark generator command-line parser.

    Inputs:
        None.

    Outputs:
        argparse.ArgumentParser: Parser exposing template, output, count, and seed.
    """

    parser = argparse.ArgumentParser(
        description="Generate a deterministic synthetic assessment benchmark XLSX"
    )
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--companies", type=int, default=24)
    parser.add_argument("--seed", type=int, default=20260713)
    return parser


def main() -> None:
    """Generate a workbook using command-line arguments and print its path.

    Inputs:
        Command-line arguments parsed from ``sys.argv``.

    Outputs:
        None. Progress and the final destination are written to the terminal.
    """

    arguments = build_parser().parse_args()
    destination = generate_workbook(
        arguments.template,
        arguments.output,
        companies=arguments.companies,
        seed=arguments.seed,
    )
    print(f"Synthetic benchmark workbook written to {destination}")


if __name__ == "__main__":
    main()
