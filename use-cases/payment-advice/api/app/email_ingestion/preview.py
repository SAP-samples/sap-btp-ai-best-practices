"""Bounded inert attachment previews; original bytes remain authenticated downloads."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree


def preview(filename, content):
    """Return at most 100 rows/30 columns or 20000 text characters; reject zip bombs."""
    ext = Path(filename).suffix.lower()
    if ext in {'.txt', '.csv', '.tsv'}:
        return content[:100000].decode('utf-8-sig', errors='replace')[:20000]
    if ext in {'.xlsx', '.docx'}:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Office file expands beyond the preview limit')
            if ext == '.docx':
                root = ElementTree.fromstring(archive.read('word/document.xml'))
                return '\n'.join(node.text or '' for node in root.iter() if node.tag.endswith('}t'))[:20000]
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            return '\n'.join('\t'.join('' if v is None else str(v)[:200] for v in row)
                for row in workbook.active.iter_rows(max_row=100, max_col=30, values_only=True))[:20000]
        finally:
            workbook.close()
    return None
