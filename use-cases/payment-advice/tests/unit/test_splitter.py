"""
Unit tests for the format-aware splitter (step 2). Offline; no SAP calls.

Covers tabular row-chunking (boundaries, header repetition, row-order integrity),
the >49-column refusal, PDF page-chunking, and format probing.

Run from the repo root:
    PYTHONPATH=api python -m unittest discover -s tests/unit -q
"""

from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from pathlib import Path

_API_DIR = Path(__file__).resolve().parents[2] / "api"
if str(_API_DIR) not in sys.path:
    sys.path.insert(0, str(_API_DIR))

from openpyxl import Workbook, load_workbook  # noqa: E402
from pypdf import PdfReader, PdfWriter  # noqa: E402

from app.payment_advice import splitter  # noqa: E402
from app.payment_advice.splitter import (  # noqa: E402
    ColumnLimitExceeded,
    FORMAT_OPAQUE,
    FORMAT_PDF,
    FORMAT_TABULAR,
    probe,
    split_document,
)


def _write_csv(path: Path, header: list[str], data_rows: list[list]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(data_rows)


def _make_pdf(path: Path, pages: int) -> None:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    with path.open("wb") as handle:
        writer.write(handle)


class TabularRowChunking(unittest.TestCase):
    """Row-based splitting keeps the header on every chunk and preserves order."""

    def _rows(self, n: int) -> list[list]:
        # Each row is uniquely identifiable by its first cell to check ordering.
        return [[i, f"inv{i}", i * 1.5] for i in range(n)]

    def test_exactly_at_limit_is_not_split(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "advice.csv"
            _write_csv(src, ["id", "invoice", "amount"], self._rows(100))
            parts = split_document(src, Path(tmp) / "out", max_line_items=100)
            self.assertEqual(parts, [src], "at-limit file must not be split")

    def test_over_limit_splits_with_header_and_order(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "advice.csv"
            original = self._rows(205)
            _write_csv(src, ["id", "invoice", "amount"], original)
            parts = split_document(src, Path(tmp) / "out", max_line_items=100)
            self.assertEqual(len(parts), 3)  # 100 + 100 + 5

            reconstructed: list[list] = []
            for part in parts:
                with part.open(newline="", encoding="utf-8") as handle:
                    rows = list(csv.reader(handle))
                self.assertEqual(rows[0], ["id", "invoice", "amount"], "header on every chunk")
                reconstructed.extend(rows[1:])

            # Chunk sizes 100/100/5.
            sizes = []
            for part in parts:
                with part.open(newline="", encoding="utf-8") as handle:
                    sizes.append(sum(1 for _ in handle) - 1)  # minus header
            self.assertEqual(sizes, [100, 100, 5])

            # First cell of each reconstructed row must match original order 0..204.
            self.assertEqual([int(r[0]) for r in reconstructed], list(range(205)))

    def test_xlsx_split_roundtrip(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "advice.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["id", "invoice", "amount"])
            for row in self._rows(150):
                sheet.append(row)
            workbook.save(src)

            parts = split_document(src, Path(tmp) / "out", max_line_items=100)
            self.assertEqual(len(parts), 2)  # 100 + 50
            total_data = 0
            for part in parts:
                book = load_workbook(part, read_only=True)
                values = [list(r) for r in book.active.iter_rows(values_only=True)]
                book.close()
                self.assertEqual(values[0], ["id", "invoice", "amount"])
                total_data += len(values) - 1
            self.assertEqual(total_data, 150)

    def test_too_many_columns_refused(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "wide.csv"
            header = [f"c{i}" for i in range(50)]  # 50 > 49
            _write_csv(src, header, [[0] * 50])
            with self.assertRaises(ColumnLimitExceeded):
                split_document(src, Path(tmp) / "out", max_columns=49)


class PdfPageChunking(unittest.TestCase):
    """PDF splitting produces page-bounded parts that sum to the original."""

    def test_pages_within_limit_not_split(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "doc.pdf"
            _make_pdf(src, pages=2)
            parts = split_document(src, Path(tmp) / "out", max_pages=2)
            self.assertEqual(parts, [src])

    def test_pages_over_limit_split(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "doc.pdf"
            _make_pdf(src, pages=5)
            parts = split_document(src, Path(tmp) / "out", max_pages=2)
            self.assertEqual(len(parts), 3)  # 2 + 2 + 1
            page_counts = [len(PdfReader(str(p)).pages) for p in parts]
            self.assertEqual(page_counts, [2, 2, 1])
            self.assertEqual(sum(page_counts), 5)


class Probing(unittest.TestCase):
    """probe() classifies formats and flags oversized inputs."""

    def test_opaque_txt_never_oversized(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "advice.txt"
            src.write_text("line\n" * 5000, encoding="utf-8")
            report = probe(src)
            self.assertEqual(report.fmt, FORMAT_OPAQUE)
            self.assertFalse(report.oversized)

    def test_pdf_probe_reports_pages(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "doc.pdf"
            _make_pdf(src, pages=3)
            report = probe(src, max_pages=100)
            self.assertEqual(report.fmt, FORMAT_PDF)
            self.assertEqual(report.pages, 3)
            self.assertFalse(report.oversized)

    def test_tabular_probe_counts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "t.csv"
            _write_csv(src, ["a", "b"], [[1, 2], [3, 4], [5, 6]])
            report = probe(src)
            self.assertEqual(report.fmt, FORMAT_TABULAR)
            self.assertEqual(report.rows, 3)
            self.assertEqual(report.columns, 2)


if __name__ == "__main__":
    unittest.main()
