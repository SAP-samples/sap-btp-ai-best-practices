"""Deterministic eligibility/insights files stored by immutable scope token in HANA."""

import hashlib
import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook
from ...optimizer.report.markdown_to_pdf import markdown_to_pdf
from .schema import transaction


def workbook_bytes(book):
    """Serialize an in-memory workbook into download bytes without application files."""
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def safe_cell(value):
    """Keep user-controlled spreadsheet strings from becoming executable formulas."""
    return "'" + value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value


def eligibility_workbook(rows):
    """Export source IDs, original amounts and actual eligibility diagnostics."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Eligibility'
    sheet.append(['Source row ID','Source row','Eligibility','Customer','Seller','Reference','Amount','Currency','Due date','Rule failures'])
    for row in rows:
        invoice = row['invoice']
        values = [row['row_id'],row['source_row_number'],'Approved upstream' if row.get('eligibility_source') == 'upstream' else 'Eligible' if row['eligible'] else 'Not eligible',
                  invoice['debtor_name'],invoice['seller_name'],invoice['invoice_ref'],invoice.get('amount_original'),
                  invoice['original_currency'],invoice['due_date'],
                  '; '.join(f"{item['rule_code']}: {item['description']}" for item in row['diagnostics']['failed_rules'])]
        sheet.append([safe_cell(value) for value in values])
    sheet.freeze_panes = 'A2'
    return workbook_bytes(book)


def insights_workbook(snapshot):
    """Export the same current/history denominators, evidence and assumptions as the dialog."""
    book = Workbook()
    sheet = book.active
    sheet.title = 'Metrics'
    sheet.append(['Metric','Current','Historical'])
    for key in ('total','eligible','not_eligible','not_eligible_rate','missing_amount_count'):
        sheet.append([key,snapshot['current_metrics'][key],snapshot['historical_metrics'][key]])
    evidence = book.create_sheet('Rule evidence')
    evidence.append(['Rule','Current count','Historical count'])
    for item in snapshot['evidence']:
        evidence.append([item['rule_code'],item['current_count'],item['historical_count']])
    scope = book.create_sheet('Comparison scope')
    for key in ('scope','comparison_as_of','comparison_exclusions'):
        scope.append([key,safe_cell(json.dumps(snapshot[key]))])
    return workbook_bytes(book)


def insights_pdf(snapshot):
    """Render the saved scope through the application's existing PDF renderer."""
    lines = ['# Eligibility Insights', '', 'Comparison as of '+snapshot['comparison_as_of'], '',
             '| Population | Invoices | Not eligible | Rate |', '|---|---:|---:|---:|']
    for label,key in (('Current scope','current_metrics'),('Historical comparison','historical_metrics')):
        value = snapshot[key]
        rate = 'Unavailable' if value['not_eligible_rate'] is None else str(value['not_eligible_rate'])+'%'
        lines.append(f"| {label} | {value['total']} | {value['not_eligible']} | {rate} |")
    for alert in snapshot['alerts']:
        lines.extend(['', '## '+alert['title'], '', alert['description']])
    if not snapshot['alerts']:
        lines.extend(['', 'No supported alerts for this scope. An empty baseline is not evidence of zero risk.'])
    lines.extend(['', '## Comparison exclusions', '', json.dumps(snapshot['comparison_exclusions'])])
    with TemporaryDirectory(prefix='workspace-report-') as directory:
        output = Path(directory) / 'insights.pdf'
        markdown_to_pdf('\n'.join(lines), output)
        return output.read_bytes()


def save_analysis_exports(analyses, analysis_id, snapshot):
    """Persist an immutable workbook/PDF bundle and return its content-addressed token."""
    token = hashlib.sha256(json.dumps(snapshot,sort_keys=True).encode()).hexdigest()
    with transaction(analyses.backend, analyses.db_path) as cursor:
        cursor.execute('SELECT scope_token FROM RECEIVABLES_ANALYSIS_EXPORTS WHERE scope_token = ?', (token,))
        if cursor.fetchone():
            return token
        cursor.execute('INSERT INTO RECEIVABLES_ANALYSIS_EXPORTS '
                       '(scope_token,analysis_id,eligibility,insights_excel,insights_pdf) VALUES (?, ?, ?, ?, ?)',
                       (token,analysis_id,eligibility_workbook(snapshot['current_rows']),insights_workbook(snapshot),insights_pdf(snapshot)))
    return token


def read_analysis_export(analyses, analysis_id, token, kind):
    """Return only allowlisted artifacts belonging to the requested analysis and scope."""
    columns = {'eligibility':'eligibility','insights-excel':'insights_excel','insights-pdf':'insights_pdf'}
    if kind not in columns:
        raise LookupError('Unknown export')
    with transaction(analyses.backend, analyses.db_path) as cursor:
        cursor.execute(f'SELECT {columns[kind]} FROM RECEIVABLES_ANALYSIS_EXPORTS WHERE scope_token = ? AND analysis_id = ?', (token,analysis_id))
        row = cursor.fetchone()
        if row is None:
            raise LookupError('Export not found')
        return bytes(row[0].read() if hasattr(row[0],'read') else row[0])
