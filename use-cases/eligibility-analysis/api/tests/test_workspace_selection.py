"""Direct extraction import checks: no eligibility calls, exact scope, preserved dates."""
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

import test_workspace_analysis as analysis_tests
from app.models.workspace import CandidateScope
from app.services.workspace.candidates import canonical_candidates
from app.services.workspace.insights import WorkspaceInsights
from app.services.workspace.settings_import import import_settings
from app.services.workspace.settings import validate_settings


class SelectionImportTests(unittest.TestCase):
    """Reuse isolated stores to exercise the real extraction upload path."""

    setUp = analysis_tests.WorkspaceAnalysisTests.setUp

    def test_original_packages(self):
        """All original workbooks import directly without history or date flattening."""
        root = Path(__file__).resolve().parents[2] / 'data/synthetic'
        for name, count, history in [('stress_balanced_limits',1440,1440),
                                     ('stress_tight_capacity',1440,1440),
                                     ('synthetic_2025-01-28_w8_n40_seed42',320,52)]:
            with self.subTest(package=name), patch('app.services.workspace.service.evaluate_offer', side_effect=AssertionError('Eligibility must not run')):
                content=(root/name/'synthetic_extraction.xlsx').read_bytes()
                saved=self.service.import_selection(content,'synthetic_extraction.xlsx',date(2025,1,28),name)
                self.assertEqual(saved['total_invoices'],count)
                self.assertEqual(saved['settings']['historical_rows_excluded'],history)
                self.assertEqual(self.service.import_selection(content,'synthetic_extraction.xlsx',date(2025,1,28),name)['analysis_id'],saved['analysis_id'])
                rows=self.analyses.all_rows(saved['analysis_id'])
                self.assertTrue(all(row['eligibility_source']=='upstream' for row in rows))
                run=self.service.create_run(saved['analysis_id'],CandidateScope(mode='all_eligible'))
                settings=import_settings((root/name/'limits.yaml').read_bytes(),'limits.yaml',run['settings'])
                preview=validate_settings(settings,rows)
                self.assertEqual(preview['readiness_issues'],[])
                frame=canonical_candidates(saved,rows,settings)
                self.assertGreater(frame['Offer File Date (UTC)'].nunique(),1)
                self.assertEqual(len(frame),count)
                self.assertNotIn('Reconciliation File Date (UTC)',frame)
                with self.assertRaises(ValueError): WorkspaceInsights(self.analyses).analyze(saved['analysis_id'])

    def test_validation_and_history_boundary(self):
        """Reject malformed candidates atomically and exclude upstream approvals from diagnostics."""
        from io import BytesIO
        from openpyxl import Workbook
        book=Workbook(); sheet=book.active
        sheet.append(['Company Code','Customer','Invoice Reference','Offer File Date (UTC)',
                      'Due Date','Purchase Price','Currency','Issuance Date','Synthetic Row Type'])
        sheet.append(['F1','C1','R1','2025-01-28','2025-03-01',100,'EUR','2025-01-01','candidate'])
        data=BytesIO(); book.save(data)
        self.service.import_selection(data.getvalue(),'input.xlsx',date(2025,1,28),'upstream')
        local=self.service.analyze(self.content,'offer.xlsx',date(2026,2,2),{},'local')
        insights=WorkspaceInsights(self.analyses).analyze(local['analysis_id'])
        self.assertEqual(insights['historical_metrics']['total'],0)
        before=self.analyses.list_analyses()['total']
        for column, value in [(2,None),(5,'NaT'),(6,'Infinity'),(9,'unknown')]:
            original=sheet.cell(2,column).value; sheet.cell(2,column).value=value
            data=BytesIO(); book.save(data)
            with self.assertRaises(ValueError):
                self.service.import_selection(data.getvalue(),'bad.xlsx',date(2025,1,28),f'bad-{column}')
            self.assertEqual(self.analyses.list_analyses()['total'],before)
            sheet.cell(2,column).value=original
